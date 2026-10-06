"""After apply_truth.py: settle the register rows that register_drift.py flagged, using what the code now shows.

usage: python qos/evidence/apply_stale_claims.py <workbook.xlsx>
Rows whose mechanism now exists stop saying 'Mechanism not found'. Rows where a claim pattern was only matching an
unrelated word keep their status and get a narrower pattern in qos/register_claims.yaml.
"""
import re
import sys
from datetime import date

import openpyxl
import yaml

wb = openpyxl.load_workbook(sys.argv[1])
tr = wb["Task Register"]
h = {c.value: c.column for c in tr[1]}
rows = {tr.cell(r, 1).value: r for r in range(2, tr.max_row + 1)}
NOTE = "Drift check 2026-10-05: "

# task -> (Validation Status, Work status or None, why)
REAL = {
    "C-0614": ("Partial", "Built - partial", "Supplier.income_tax_specified_person flag and apply_specified_person_tds exist; no PAN utility lookup (dropped from the plan pending the CA)"),
    "X-0522": ("Partial", "Built - partial", "higher_withholding_rate exists in accounting/services.py; no PAN specified-person check (dropped from the plan pending the CA)"),
    "UX-0586": ("Partial", None, "CommandPalette.tsx exists (Ctrl+K over navigation); no customer or invoice search, no arrow-key selection"),
    "UX-0591": ("Partial", None, "PrivacyMask provider and toggle exist; blur covers papers and table cells only"),
    "C-0607": ("Partial", None, "Section 50 worksheet and CSV exist in planwave; rates are a single dated row pending the CA"),
    "C-0609": ("Partial", None, "Schedule H/H1/X check and register exist behind a per-tenant flag; pharmacist review of layout and retention pending"),
    "X-0483": ("Partial", "Built - partial", "Cash shift register with denominations and close exists in accounting/cash_shifts.py"),
    "P-0006": ("Partial", "Built - partial", "Cash shift register compares expected and counted cash"),
    "X-0347": ("Partial", "Built - partial", "Daily cash close compares expected and counted cash (accounting/cash_shifts.py)"),
    "X-0515": ("Partial", None, "Validity extension exists as an action; live provider extension is refused, stub only"),
    "X-0526": ("Partial", "Built - partial", "interest_exposure and the overdue report exist; report only, no posting"),
    "X-0220": ("Partial", "Built - partial", "match_bill_to_po 3-way match exists in purchases/services.py"),
    "X-0222": ("Partial", "Built - partial", "Price variance breach is raised by the 3-way match and the WAVG variance journal"),
    "X-0704": ("Partial", "Built - partial", "Bill rate against PO rate is checked by the 3-way match; manager override exists"),
    "UX-0592": ("Partial", "Built - partial", "scanSounds.ts exists (web audio cues)"),
}
for tid, (val, work, why) in REAL.items():
    r = rows[tid]
    tr.cell(r, h["Validation Status"]).value = val
    if work and tr.cell(r, h["Work status"]).value in ("Gap", None, ""):
        tr.cell(r, h["Work status"]).value = work
    old = tr.cell(r, h["Notes"]).value
    tr.cell(r, h["Notes"]).value = (str(old) + " | " if old else "") + NOTE + why

# claims whose pattern matched an unrelated word: narrow them
claims_path = r"E:\Bizboard\qos\register_claims.yaml"
claims = yaml.safe_load(open(claims_path, encoding="utf-8"))
NARROW = {
    "X-0529": "cheque.{0,30}(deposit_slip|clearing_account)|ChequeStatus.{0,20}DEPOSITED",
    "X-0372": "return_receipt|ReturnReceipt|goods_return_receipt|receive_returned",
    "P-0138": "salesperson.{0,40}(performance|leaderboard|target)|my_sales_performance",
    "X-0394": "salesperson.{0,40}(performance|leaderboard|target)|my_sales_performance",
    "X-0395": "team_performance|manager.{0,40}(team|leaderboard)",
}
for c in claims:
    if c["task_id"] in NARROW:
        c["pattern"] = NARROW[c["task_id"]]
open(claims_path, "w", encoding="utf-8").write(yaml.safe_dump(claims, sort_keys=False, allow_unicode=True, width=200))

# test references: drop trailing prose in parentheses and normalise separators
col = h["Test case ID"]
for r in range(2, tr.max_row + 1):
    v = tr.cell(r, col).value
    if not v:
        continue
    parts = []
    for x in str(v).split(";"):
        x = re.sub(r"\s*\([^)]*\)\s*$", "", x.strip())
        x = re.sub(r"\s*::\s*", "::", x)
        if x and x not in parts:
            parts.append(x)
    tr.cell(r, col).value = "; ".join(parts)

wb.save(sys.argv[1])
print("settled", len(REAL), "rows; narrowed", len(NARROW), "claims")
