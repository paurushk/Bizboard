"""Export the Implementation plan sheet to markdown.  usage: python qos/evidence/plan_to_md.py <workbook.xlsx> <out.md>"""
import math
import sys
from collections import defaultdict
from datetime import timedelta
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1])
ws = wb["Implementation plan"]
h = [c.value for c in ws[4]]
rows = [dict(zip(h, r)) for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
d = lambda v: v.strftime("%d %b %Y") if hasattr(v, "strftime") else ("" if isinstance(v, str) and v.startswith("=") else (v or ""))

def planned_days(row):
    v = row["Planned days (with buffer)"]
    if isinstance(v, (int, float)):
        return v
    est = row["Estimate (days)"] or 0
    return math.ceil(float(est) * 1.2) if est else 0
g = wb["Plan guide"]
guide = {r[0]: r[1] for r in g.iter_rows(min_row=10, values_only=True) if r[0]}
out = ["# Bizboard implementation plan (quality-first), revised", "",
       "Generated from the workbook sheets *Implementation plan* and *Plan decisions* on 2026-10-04, after the 63 pre-implementation questions. The workbook is the source of truth for status; this file is the readable copy.", "",
       "## Assumptions", "",
       f"- Start: {d(g['B4'].value)}; agent lanes: {g['B5'].value}; contingency buffer: {int(g['B6'].value * 100)}% per item plus one release-buffer week per milestone.",
       *[f"- **{k}.** {guide[k]}" for k in ("Team assumption", "Dates", "Waves", "Freeze", "Reviewer rule", "When tests run", "Source of truth", "Other plans", "Slips")], "",
       f"**Definition of done.** {guide['Definition of done']}", "",
       "## Milestones", "", "| Milestone | Planned date | Exit criteria |", "|---|---|---|"]
by_wave = {}
for r in rows:
    by_wave.setdefault(r["Wave"], []).append(r["Planned due"])
for r in wb["Plan milestones"].iter_rows(min_row=5, max_row=7, values_only=True):
    planned = r[2]
    if (planned is None or (isinstance(planned, str) and planned.startswith("="))) and r[1] in by_wave and by_wave[r[1]]:
        dues = [x for x in by_wave[r[1]] if hasattr(x, "strftime")]
        planned = max(dues) + timedelta(days=7) if dues else planned
    out.append(f"| {r[0]} | {d(planned)} | {r[5]} |")
out += ["", "## Summary by wave", "", "| Wave | Gate | Items | Estimate days | Planned days | First start | Last due |", "|---|---|---|---|---|---|---|"]
for w in (1, 2, 3):
    rs = [r for r in rows if r["Wave"] == w]
    out.append(f"| {w} | {rs[0]['Release gate']} | {len(rs)} | {sum(r['Estimate (days)'] for r in rs):g} | {sum(planned_days(r) for r in rs):g} | {d(min(r['Planned start'] for r in rs))} | {d(max(r['Planned due'] for r in rs))} |")
out += ["", "## Summary by workstream", "", "| Workstream | Items | Estimate days |", "|---|---|---|"]
ws_ = defaultdict(lambda: [0, 0])
for r in rows:
    ws_[r["Workstream"]][0] += 1; ws_[r["Workstream"]][1] += r["Estimate (days)"]
for k, (n, e) in ws_.items():
    out.append(f"| {k} | {n} | {e:g} |")
out += ["", "## Summary by owner", "", "| Owner | Items | Planned days |", "|---|---|---|"]
ow = defaultdict(lambda: [0, 0])
for r in rows:
    ow[r["Owner"]][0] += 1; ow[r["Owner"]][1] += planned_days(r)
for k, (n, e) in sorted(ow.items()):
    out.append(f"| {k} | {n} | {e:g} |")
dq = wb["Plan decisions"]
out += ["", "## Decisions: answers to the 63 questions", "", "Status: Decided = settled; Confirm = default chosen, founder confirms before the item starts; External = needs a CA, pharmacist, host or other outside input; Dropped = removed from the plan.", "",
        "| Q | Topic | Items | Status | Answer | What changed in the plan |", "|---|---|---|---|---|---|"]
for q in dq.iter_rows(min_row=5, values_only=True):
    if q[0]:
        out.append("| " + " | ".join(str(x).replace("|", "/") for x in q[:6]) + " |")
out += ["", "### Needs your confirmation before the affected item starts", ""]
for q in dq.iter_rows(min_row=5, values_only=True):
    if q[0] and q[3] in ("Confirm", "External"):
        out.append(f"- Q{q[0]} ({q[3]}): {q[1]}, items {q[2]}")
for w in (1, 2, 3):
    rs = [r for r in rows if r["Wave"] == w]
    out += ["", f"## Wave {w}: {rs[0]['Release gate']} gate", "",
            "| WI | Task | Work item | Workstream | Owner | Est. | Start | Due | Depends on |", "|---|---|---|---|---|---|---|---|---|"]
    for r in rs:
        out.append(f"| {r['WI ID']} | {r['Register task'] or ''} | {r['Work item']} | {r['Workstream']} | {r['Owner']} | {r['Estimate (days)']:g} | {d(r['Planned start'])} | {d(r['Planned due'])} | {r['Depends on'] or ''} |")
    out += ["", f"### Wave {w} item detail", ""]
    for r in rs:
        out += [f"#### {r['WI ID']} {r['Work item']}", "",
                f"- Register task: {r['Register task'] or 'none'}  |  Workstream: {r['Workstream']}  |  Layer: {r['Layer']}  |  Kind: {r['Kind']}  |  Owner: {r['Owner']}" + (f"  |  Cut line rank: {r['Cut line']}" if r['Cut line'] else "") + (f"  |  Decisions: Q{r['Decision refs (Q)']}" if r['Decision refs (Q)'] else ""),
                f"- Estimate: {r['Estimate (days)']:g} days ({planned_days(r):g} with buffer)  |  Planned {d(r['Planned start'])} to {d(r['Planned due'])}  |  Depends on: {r['Depends on'] or 'none'}",
                f"- Scope: {r['Scope (quality-first)']}",
                f"- Acceptance criteria: {r['Acceptance criteria']}",
                f"- Planned test file: `{r['Planned test file']}`" if r["Planned test file"] else "- Planned test file: none (process item)",
                "- Stages:"]
        for s in ("Design", "Build", "Tests", "Review", "Register sync"):
            out.append(f"  - [{'x' if r[s] in ('N/A',) else ' '}] {s}" + (" (not applicable)" if r[s] == "N/A" else ""))
        out += ["- Tracking: actual start ____  actual done ____  PR ____  evidence ____", ""]
open(sys.argv[2], "w", encoding="utf-8").write("\n".join(out) + "\n")
print(len(out), "lines")
