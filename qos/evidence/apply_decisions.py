"""Record the 2026-10-04 delegated founder decisions in the register (sheets + Release gate + duplicate handling).
usage: python qos/evidence/apply_decisions.py <workbook.xlsx>   (recalculate in Excel afterwards)"""
import sys, copy
from datetime import date
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

BI, BP, BL, P = "Build in MVP", "Build in Pilot", "Build in GA", "Accept recommended"
GATE = {BI: "MVP", BP: "Pilot", BL: "Later"}
# task -> (decision, effort days, what quality-first means here)
MVP = {
 # Wave 1: compliance, integrity, security and money paths. Gate MVP.
 "P-0115": (BI, 3, "Audit event on every posting action (journal post, purchase complete with books off included); add IP and before/after values; one test per action; drop the xfail."),
 "X-0432": (BI, 0, "Delivered with P-0115."),
 "P-0507": (BI, 2, "Rule 11(g) coverage map signed off by the CA, edit log on every accounting record, certification export."),
 "C-0616": (BI, 0, "Delivered with P-0507: report certifying logging was never disabled."),
 "S-0561": (BI, 3, "Scheduled worker runs trial-balance and invariants nightly, quarantines unbalanced batches and alerts."),
 "S-0568": (BI, 4, "Soft delete (is_deleted) on masters, PROTECT on line-item foreign keys, restore flow."),
 "S-0570": (BI, 3, "WAL archiving and point-in-time recovery plus the daily encrypted dump; restore drill documented."),
 "S-0572": (BI, 2, "Redis token bucket on public endpoints with per-tenant limits."),
 "SEC-0622": (BI, 2, "Role-based field masking for purchase price and margin in serializers, exports and PDFs."),
 "SEC-0624": (BI, 1.5, "Mandatory TOTP for owner and admin, enrolment on first login, recovery codes."),
 "SEC-0625": (BI, 3, "Redis access-token blocklist and websocket drop on deactivation."),
 "SEC-0626": (BI, 2, "Sanitise notes and descriptions on input and in PDF templates; add injection tests."),
 "SEC-0628": (BI, 3, "AES-256-GCM for stored credentials and encrypted bank account and IFSC fields, with data migration."),
 "X-0715": (BI, 2, "Snapshot customer name and address on the invoice at completion; backfill existing invoices."),
 "C-0609": (BI, 5, "Schedule H/X register with doctor and prescription logging, printable for inspectors."),
 "C-0602": (BI, 4, "Four-condition Section 16(2) checklist on each purchase bill."),
 "C-0603": (BI, 3, "60-day and 7-day alert tiers, TIME_BARRED status and ITC exclusion."),
 "C-0604": (BI, 5, "Section 17(5) auto-classification (motor vehicle, food, gifts) with override and audit."),
 "P-0503": (BI, 3, "Section 17(5) blocked-credit classification in the 2A/2B reconciliation."),
 "C-0607": (BI, 3, "Section 50 interest calculator and delayed-filing worksheet."),
 "X-0527": (BI, 3, "60-day rule, dispatch block and unlock PIN on top of the opt-in credit hold."),
 "X-0530": (BI, 3, "Bank dishonour fee charged to the customer, notice letter, owner alert."),
 "X-0713": (BI, 2, "Automated account action and Section 138 notice generator on a dishonoured cheque."),
 "X-0707": (BI, 3, "ITC quarantine and supplier payment hold when a supplier GSTIN is cancelled."),
 "P-0128": (BI, 5, "General pending-approval queue with approve and reject and an audit trail."),
 "P-0152": (BI, 2, "Dual approval for stock adjustments above a threshold, on the approval queue."),
 "X-0702": (BI, 3, "Time-limited WhatsApp approval token with 15-minute expiry for credit-limit overrides."),
 "R-0636": (BI, 3, "Real POS checkout benchmark and optimisation to the 200 ms target."),
 "X-0514": (BI, 2, "Part B vehicle update action for e-way bills."),
 "X-0515": (BI, 2, "E-way bill validity extension action."),
 "C-0618": (BI, 2, "Part B time-limit alert and multi-vehicle splitting."),
 "C-0606": (BI, 3, "Rule 138 PIN-to-PIN distance lookup and validity computation."),
 # Wave 2: usability and breadth. Gate Pilot.
 "P-0046": (BP, 2, "Native salt/composition and manufacturer fields, searchable at billing."),
 "P-0052": (BP, 4, "Bulk invoice CSV import with validation and per-row errors."),
 "P-0081": (BP, 6, "Dispatch, in-transit and receipt-confirmation stages for stock transfers."),
 "P-0089": (BP, 5, "Per-outlet breakdown and outlet switching on the dashboard."),
 "P-0119": (BP, 3, "Interest exposure and bad-debt risk on overdue customers."),
 "P-0180": (BP, 2, "Contact fields and user-to-location mapping on warehouses."),
 "X-0706": (BP, 3, "Transfer vehicle-breakdown Part B workflow on stock transfers."),
 "UX-0596": (BP, 1, "Typed-confirmation input on destructive dialogs."),
 "I-0651": (BP, 5, "Direct ESC/POS over WebUSB and network printers."),
 "I-0655": (BP, 3, "Wedge timing heuristic and a browser camera scanner."),
 "SaaS-0669": (BP, 5, "GST invoice generation and emailing for subscription charges."),
 # Wave 3: differentiators. Gate Later (GA).
 "D-0684": (BL, 10, "Per-debtor payment-delay model with Monte Carlo cash-flow simulation."),
 "D-0690": (BL, 6, "Fuzzy and Levenshtein invoice-number matching in the GSTR-2B reconciler."),
 "D-0693": (BL, 10, "Tally XML backup parser and multi-year history migration."),
}
P0D = {  # P0 gap decisions sheet: the same full-build decisions, plus the P1 and P2 rows
 "P-0115": BI + " (3 days)", "X-0432": BI + " (with P-0115)", "P-0507": BI + " (2 days, CA sign-off)",
 "S-0561": BI + " (3 days): full worker with quarantine", "S-0579": "Close as built after confirming the nightly seal runs in production, then add a monitoring alert",
 "C-0603": BI + " (3 days): alert tiers and exclusion engine", "SEC-0624": BI + " (1.5 days)",
 "GOS-0721": "Add end-to-end test and validate before Pilot", "GOS-0737": "Add end-to-end test with stock and ledger check, validate before Pilot",
 "GOS-0739": "Add end-to-end test with stock and ledger check, validate before Pilot",
 "C-0614": BL + " (3 days): PAN utility check", "X-0522": BL + " (2 days, with C-0614)",
 "S-0574": BL + " (8 days): anomaly detector", "S-0562": BP + " (2 days): release orphaned allocations on expired sessions",
 "UX-0581": BP + " (3 days)", "UX-0582": BP + " (3 days)", "UX-0586": BP + " (2 days)",
 "UX-0588": BP + " (2 days)", "UX-0590": BP + " (2 days)", "UX-0591": BP + " (1 day)",
}
DEFER = P
NOTE = "Decision 2026-10-04 (delegated, quality-first): "
OWNER = "Founder (delegated to Claude)"
TODAY = date(2026, 10, 4)


def hdr_like(ws, row, col, text, src_col):
    c = ws.cell(row, col, text)
    s = ws.cell(row, src_col)
    c.font, c.fill, c.alignment, c.border = copy.copy(s.font), copy.copy(s.fill), copy.copy(s.alignment), copy.copy(s.border)
    ws.column_dimensions[get_column_letter(col)].width = 50 if text.startswith("Why") else 18


def add_note(tr, head, row, text):
    old = tr.cell(row, head["Notes"]).value
    tr.cell(row, head["Notes"], (str(old) + " | " if old else "") + NOTE + text)


wb = openpyxl.load_workbook(sys.argv[1])
tr = wb["Task Register"]
head = {c.value: c.column for c in tr[1]}
rows = {tr.cell(r, 1).value: r for r in range(2, tr.max_row + 1)}

# --- MVP scope after review
ws = wb["MVP scope after review"]
hdr_like(ws, 3, 10, "Effort (days)", 8)
hdr_like(ws, 3, 11, "Why / what quality means", 7)
n = 0
for r in range(4, ws.max_row + 1):
    tid = ws.cell(r, 1).value
    if tid not in MVP:
        continue
    d, eff, why = MVP[tid]
    ws.cell(r, 8, d); ws.cell(r, 9, OWNER); ws.cell(r, 10, eff); ws.cell(r, 11, why)
    ws.cell(r, 11).alignment = Alignment(wrap_text=True, vertical="top")
    tt = rows[tid]
    add_note(tr, head, tt, f"{d} ({eff} days). {why}")
    tr.cell(tt, head["Release gate"], GATE[d])
    if eff and tr.cell(tt, head["Estimate days"]).value is None:
        tr.cell(tt, head["Estimate days"], eff)
    n += 1
print("MVP scope decisions", n)

# --- P0 gap decisions
ws = wb["P0 gap decisions"]
m = 0
for r in range(4, ws.max_row + 1):
    tid = ws.cell(r, 1).value
    if not tid:
        continue
    ws.cell(r, 9, P0D.get(tid, DEFER)); ws.cell(r, 10, OWNER); ws.cell(r, 11, TODAY)
    m += 1
print("P0 gap decisions", m)

# --- Duplicate proposals
ws = wb["Duplicate proposals"]
d_ = 0
for r in range(4, ws.max_row + 1):
    prop, a, b = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
    if not prop:
        continue
    ra, rb = rows[a], rows[b]
    if prop == "MERGE":
        ws.cell(r, 10, f"Approved: {b} retired as duplicate of {a} (row kept so links survive)")
        tr.cell(rb, head["Duplicate or related"], f"Duplicate of {a} - retired")
        tr.cell(rb, head["Release gate"], "Later")
        add_note(tr, head, rb, f"retired as duplicate of {a}.")
    else:
        ws.cell(r, 10, "Approved: keep both rows, one test linked to both")
        ta, tb = tr.cell(ra, head["Test case ID"]).value, tr.cell(rb, head["Test case ID"]).value
        src_row = ra if ta else (rb if tb else None)
        if src_row:
            dst_row = rb if src_row == ra else ra
            if not tr.cell(dst_row, head["Test case ID"]).value:
                tr.cell(dst_row, head["Test case ID"], tr.cell(src_row, head["Test case ID"]).value)
                for k in ("Verification result", "Last verified", "Verified by"):
                    tr.cell(dst_row, head[k], tr.cell(src_row, head[k]).value)
    d_ += 1
print("duplicate decisions", d_)

# --- Adjacent/Future P0 and amnesty / anti-profiteering rows
for tid, txt in {"SaaS-0666": "Keep in MVP: built. Add an end-to-end test and a load test of provisioning.",
                 "SaaS-0668": "Keep in MVP: built and tested.", "SaaS-0678": "Keep in MVP: built and tested.",
                 "D-0681": "Keep in MVP: built and tested.",
                 "C-0617": "Build in GA (3 days): amnesty tracking and reconciliation for CA clients.",
                 "C-0619": "Build in GA (3 days): anti-profiteering price-cut compliance view."}.items():
    add_note(tr, head, rows[tid], txt)
wb.save(sys.argv[1])
print("saved")
