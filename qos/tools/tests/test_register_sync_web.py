"""REG-WEB: register_sync scores Vitest and Playwright JUnit onto register rows."""

import importlib.util
from pathlib import Path

import openpyxl

_SPEC = importlib.util.spec_from_file_location(
    "register_sync",
    Path(__file__).resolve().parents[1] / "register_sync.py",
)
sync = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(sync)

VITEST = """<?xml version="1.0" encoding="UTF-8"?>
<testsuites>
  <testsuite name="src/__tests__/ux_0596.test.tsx" tests="2" failures="1">
    <testcase classname="src/__tests__/ux_0596.test.tsx" name="enables after the exact name" time="0.01"/>
    <testcase classname="src/__tests__/ux_0596.test.tsx" name="stays disabled on a partial name" time="0.01">
      <failure message="expected disabled">button was enabled</failure>
    </testcase>
  </testsuite>
</testsuites>
"""

PLAYWRIGHT = """<?xml version="1.0" encoding="UTF-8"?>
<testsuites>
  <testsuite name="e2e/pos.spec.ts" tests="1" failures="0" skipped="0">
    <testcase name="checkout from the keyboard" classname="e2e/pos.spec.ts" time="1.2"/>
  </testsuite>
</testsuites>
"""

PYTEST = """<?xml version="1.0" encoding="UTF-8"?>
<testsuite>
  <testcase classname="tests.test_mfa" name="test_enrol" time="0.2"/>
</testsuite>
"""


def _write(tmp: Path, name: str, body: str) -> Path:
    path = tmp / name
    path.write_text(body, encoding="utf-8")
    return path


def _register(tmp: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Register"
    ws.append(["Task ID", "Test case ID", "Verification result", "Last verified", "Verified by"])
    ws.append(["UX-0596", "web/src/__tests__/ux_0596.test.tsx", "", "", ""])
    ws.append(["UX-0581", "web/e2e/pos.spec.ts", "", "", ""])
    ws.append(["SEC-0624", "backend/tests/test_mfa.py", "", "", ""])
    ws.append(["UX-0586", "web/src/__tests__/missing.test.tsx", "", "", ""])
    path = tmp / "register.xlsx"
    wb.save(path)
    return path


def test_parser_reads_pytest_vitest_and_playwright(tmp_path: Path):
    vitest = sync.load_junit(_write(tmp_path, "vitest.xml", VITEST))
    assert ("enables after the exact name", "pass") in vitest["web/src/__tests__/ux_0596.test.tsx"]
    assert ("stays disabled on a partial name", "fail") in vitest["src/__tests__/ux_0596.test.tsx"]

    playwright = sync.load_junit(_write(tmp_path, "playwright.xml", PLAYWRIGHT))
    assert sync.score("web/e2e/pos.spec.ts", playwright) == "Pass"

    pytest_report = sync.load_junit(_write(tmp_path, "pytest.xml", PYTEST))
    assert sync.score("backend/tests/test_mfa.py::test_enrol", pytest_report) == "Pass"
    assert sync.score("backend/tests/test_mfa.py::test_other", pytest_report) is None


def test_failed_web_test_marks_the_row_fail_and_unknown_path_is_unmatched(tmp_path: Path, capsys):
    by_file = sync.merge_junit([
        _write(tmp_path, "vitest.xml", VITEST),
        _write(tmp_path, "playwright.xml", PLAYWRIGHT),
        _write(tmp_path, "pytest.xml", PYTEST),
    ])
    assert sync.score("web/src/__tests__/ux_0596.test.tsx", by_file) == "Fail"
    assert sync.score("web/src/__tests__/missing.test.tsx", by_file) is None

    register = _register(tmp_path)
    code = sync.main.__wrapped__ if hasattr(sync.main, "__wrapped__") else None
    assert code is None
    import sys
    argv = sys.argv
    sys.argv = [
        "register_sync.py",
        "--junit", str(tmp_path / "vitest.xml"),
        "--junit", str(tmp_path / "playwright.xml"),
        "--junit", str(tmp_path / "pytest.xml"),
        "--register", str(register),
        "--by", "unit",
        "--date", "2026-10-04",
    ]
    try:
        rc = sync.main()
    finally:
        sys.argv = argv
    assert rc == 0
    out = capsys.readouterr().out
    assert "unmatched  UX-0586" in out

    wb = openpyxl.load_workbook(register)
    rows = {r[0]: r for r in wb["Task Register"].iter_rows(min_row=2, values_only=True)}
    assert rows["UX-0596"][2] == "Fail"
    assert rows["UX-0581"][2] == "Pass"
    assert rows["SEC-0624"][2] == "Pass"
    assert rows["UX-0586"][2] in (None, "")
