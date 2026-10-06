# -*- coding: utf-8 -*-
"""
BizBoard Master Task & Persona Register Runner
Processes all 494 tasks, fills every single field with authentic information,
and produces both Bizboard_Master_Task_Register_Filled.csv and Bizboard_Master_Task_Register_Filled.xlsx.
"""

import csv
import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath('.'))

from scripts.generate_full_register import RAW_TASKS, CROSS_TASKS
from scripts.build_full_tasks import get_persona_task_details, get_cross_task_details

def main():
    headers = [
        "Task ID", "Task Type", "Persona / Scope", "Category",
        "Task / Business Action", "Primary Actor", "Supporting Persona(s)",
        "Business Outcome", "Downstream Impact", "Priority",
        "Validation Status", "Notes"
    ]

    all_rows = []

    # 1. Process Persona tasks (P-0001 to P-0188)
    for t_id, t_type, persona, cat, action, prio in RAW_TASKS:
        category, primary_actor, supporting, outcome, downstream, priority, status, notes = get_persona_task_details(
            t_id, persona, action
        )
        row = [
            t_id, t_type, persona, category, action,
            primary_actor, supporting, outcome, downstream,
            priority, status, notes
        ]
        all_rows.append(row)

    # 2. Process Cross-Persona tasks (X-0189 to X-0494)
    for t_id, t_type, persona, cat, action, prio in CROSS_TASKS:
        category, primary_actor, supporting, outcome, downstream, priority, status, notes = get_cross_task_details(
            t_id, cat, action
        )
        row = [
            t_id, t_type, persona, category, action,
            primary_actor, supporting, outcome, downstream,
            priority, status, notes
        ]
        all_rows.append(row)

    print(f"Total rows enriched: {len(all_rows)}")
    assert len(all_rows) == 494, f"Expected 494 rows, got {len(all_rows)}"

    # Verification: check for any empty or 'Not Started' values in enriched fields
    empty_cells = 0
    not_started_cells = 0
    for r_idx, r in enumerate(all_rows):
        for c_idx, val in enumerate(r):
            if val is None or str(val).strip() == "":
                print(f"Empty cell at row {r_idx+1}, col {headers[c_idx]}: ID={r[0]}")
                empty_cells += 1
            if str(val).strip().lower() == "not started":
                print(f"'Not Started' at row {r_idx+1}, col {headers[c_idx]}: ID={r[0]}")
                not_started_cells += 1

    print(f"Verification Results -> Empty cells: {empty_cells}, 'Not Started' cells: {not_started_cells}")
    if empty_cells > 0 or not_started_cells > 0:
        raise ValueError("Enrichment incomplete!")

    # 3. Export to CSV
    csv_file = "Bizboard_Master_Task_Register_Filled.csv"
    with open(csv_file, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(all_rows)
    print(f"Saved CSV: {csv_file}")

    # 4. Export to styled Excel (.xlsx)
    xlsx_file = "Bizboard_Master_Task_Register_Filled.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Master Tasks Register
    ws = wb.active
    ws.title = "Master Task Register"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark navy blue
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    data_font_bold = Font(name="Segoe UI", size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    p0_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid") # Light red
    p0_font = Font(name="Segoe UI", size=10, bold=True, color="C62828")
    p1_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid") # Light amber
    p1_font = Font(name="Segoe UI", size=10, bold=True, color="F57F17")
    p2_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid") # Light green
    p2_font = Font(name="Segoe UI", size=10, bold=True, color="2E7D32")
    status_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid") # Soft blue
    status_font = Font(name="Segoe UI", size=10, bold=True, color="1565C0")

    # Write headers
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num in [1, 2, 10, 11] else "left", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    # Write data rows
    for r_idx, r in enumerate(all_rows, start=2):
        ws.append(r)
        ws.row_dimensions[r_idx].height = 22

        for c_idx in range(1, len(r) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)

            # Center alignment for specific columns
            if c_idx in [1, 2]: # Task ID, Type
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = data_font_bold

            # Priority styling
            elif c_idx == 10:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if cell.value == "P0":
                    cell.fill = p0_fill
                    cell.font = p0_font
                elif cell.value == "P1":
                    cell.fill = p1_fill
                    cell.font = p1_font
                elif cell.value == "P2":
                    cell.fill = p2_fill
                    cell.font = p2_font

            # Validation Status styling
            elif c_idx == 11:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill = status_fill
                cell.font = status_font

    # Set column widths
    col_widths = {
        1: 12,  # Task ID
        2: 11,  # Task Type
        3: 28,  # Persona / Scope
        4: 26,  # Category
        5: 42,  # Task / Business Action
        6: 28,  # Primary Actor
        7: 28,  # Supporting Persona(s)
        8: 40,  # Business Outcome
        9: 45,  # Downstream Impact
        10: 12, # Priority
        11: 22, # Validation Status
        12: 50  # Notes
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    # Sheet 2: Executive Summary & Dashboard
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title
    ws_sum["A1"] = "BizBoard Master Task & Persona Register - Summary & Metrics"
    ws_sum["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
    ws_sum.row_dimensions[1].height = 30

    # High-level Metrics Table
    ws_sum["A3"] = "Metric"
    ws_sum["B3"] = "Count"
    ws_sum["A3"].font = header_font
    ws_sum["A3"].fill = header_fill
    ws_sum["B3"].font = header_font
    ws_sum["B3"].fill = header_fill
    ws_sum["B3"].alignment = Alignment(horizontal="center")

    summary_metrics = [
        ("Total Tasks", len(all_rows)),
        ("Persona Tasks (P-Series)", len(RAW_TASKS)),
        ("Cross-Persona Tasks (X-Series)", len(CROSS_TASKS)),
        ("Distinct Personas / Scopes", len(set(r[2] for r in all_rows))),
        ("Distinct Functional Categories", len(set(r[3] for r in all_rows))),
        ("Validated (Supported) Tasks", sum(1 for r in all_rows if "Validated" in r[10])),
        ("P0 Priority Tasks (Blocker/Core)", sum(1 for r in all_rows if r[9] == "P0")),
        ("P1 Priority Tasks (Critical)", sum(1 for r in all_rows if r[9] == "P1")),
        ("P2 Priority Tasks (Standard/Advisory)", sum(1 for r in all_rows if r[9] == "P2")),
    ]

    for idx, (m, v) in enumerate(summary_metrics, start=4):
        ws_sum[f"A{idx}"] = m
        ws_sum[f"B{idx}"] = v
        ws_sum[f"A{idx}"].font = data_font_bold
        ws_sum[f"A{idx}"].border = border_thin
        ws_sum[f"B{idx}"].font = data_font_bold
        ws_sum[f"B{idx}"].border = border_thin
        ws_sum[f"B{idx}"].alignment = Alignment(horizontal="center")

    # Category breakdown table
    cat_counts = {}
    for r in all_rows:
        cat_counts[r[3]] = cat_counts.get(r[3], 0) + 1

    start_cat_row = 15
    ws_sum[f"A{start_cat_row}"] = "Category"
    ws_sum[f"B{start_cat_row}"] = "Task Count"
    ws_sum[f"A{start_cat_row}"].font = header_font
    ws_sum[f"A{start_cat_row}"].fill = header_fill
    ws_sum[f"B{start_cat_row}"].font = header_font
    ws_sum[f"B{start_cat_row}"].fill = header_fill
    ws_sum[f"B{start_cat_row}"].alignment = Alignment(horizontal="center")

    for idx, (cat_name, count) in enumerate(sorted(cat_counts.items(), key=lambda x: -x[1]), start=start_cat_row+1):
        ws_sum[f"A{idx}"] = cat_name
        ws_sum[f"B{idx}"] = count
        ws_sum[f"A{idx}"].font = data_font
        ws_sum[f"A{idx}"].border = border_thin
        ws_sum[f"B{idx}"].font = data_font_bold
        ws_sum[f"B{idx}"].border = border_thin
        ws_sum[f"B{idx}"].alignment = Alignment(horizontal="center")

    ws_sum.column_dimensions["A"].width = 36
    ws_sum.column_dimensions["B"].width = 16

    wb.save(xlsx_file)
    print(f"Saved styled Excel workbook: {xlsx_file}")

if __name__ == "__main__":
    main()
