"""Bring recipes in from a spreadsheet, and send them back out in the same shape (MFG-X1).

One row per material or output. Rows with the same finished good and recipe name make one
revision. Nothing is saved unless every row is good: a bad row names itself and the whole file
waits for a fix.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from core.csv_utils import csv_safe
from core.exceptions import BusinessRuleError
from masters.models import Product, Unit

from manufacturing.models import Bom, BomLine, BomOutput

from . import bom as bom_svc

COLUMNS = [
    "row_type", "finished_sku", "bom_name", "output_qty", "yield_pct", "item_sku", "qty", "unit",
    "scrap_pct", "issue_method", "kind", "cost_share_pct", "recovery_rate",
]

HELP = {
    "row_type": "LINE for a material, OUTPUT for a co-product, by-product or scrap",
    "finished_sku": "SKU of the product the recipe makes",
    "bom_name": "Name of the recipe (rows with the same product and name are one recipe)",
    "output_qty": "How much one batch makes (default 1)",
    "yield_pct": "Good output as a percent of input (default 100)",
    "item_sku": "SKU of the material (LINE) or the output product (OUTPUT)",
    "qty": "Quantity per batch",
    "unit": "Short name of the unit, if not the item's own unit",
    "scrap_pct": "LINE only: extra percent lost in use",
    "issue_method": "LINE only: MANUAL or BACKFLUSH",
    "kind": "OUTPUT only: CO_PRODUCT, BY_PRODUCT or SCRAP",
    "cost_share_pct": "OUTPUT only: share of cost for a co-product",
    "recovery_rate": "OUTPUT only: value per unit of a by-product or scrap",
}


def template_csv() -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(COLUMNS)
    w.writerow(["LINE", "CHILLI-PWD-100", "Chilli powder 100 kg", "100", "98", "CHILLI-DRY", "62", "kg", "2", "MANUAL", "", "", ""])
    w.writerow(["LINE", "CHILLI-PWD-100", "Chilli powder 100 kg", "100", "98", "PACK-BAG", "100", "", "0", "BACKFLUSH", "", "", ""])
    w.writerow(["OUTPUT", "CHILLI-PWD-100", "Chilli powder 100 kg", "100", "98", "CHILLI-STALK", "3", "kg", "", "", "BY_PRODUCT", "0", "4"])
    return out.getvalue()


def parse_file(upload) -> list[dict]:
    """Rows from a .csv or .xlsx upload, keyed by the template's column names."""
    name = (getattr(upload, "name", "") or "").lower()
    data = upload.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        grid = [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig")
        grid = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    grid = [r for r in grid if any(r)]
    if not grid:
        raise BusinessRuleError("The file is empty.", code="MFG_IMPORT_EMPTY")
    header = [h.strip().lower() for h in grid[0]]
    missing = [c for c in ("row_type", "finished_sku", "item_sku", "qty") if c not in header]
    if missing:
        raise BusinessRuleError(
            f"The file needs these columns: {', '.join(missing)}. Download the template.", code="MFG_IMPORT_COLUMNS",
        )
    rows = []
    for i, raw in enumerate(grid[1:], start=2):
        row = {h: (raw[j] if j < len(raw) else "") for j, h in enumerate(header)}
        row["_row"] = i
        rows.append(row)
    return rows


def _dec(value, default=None):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        raise ValueError(f"'{value}' is not a number")


def run_import(company, user, rows, *, commit=False, activate=True) -> dict:
    """Validate every row and, when ``commit``, save each recipe as a new revision."""
    products = {p.sku.upper(): p for p in Product.objects.filter(company=company).exclude(sku="")}
    units = {u.short_name.lower(): u for u in Unit.objects.filter(company=company)}
    errors, groups = [], {}

    def err(row, message):
        errors.append({"row": row["_row"], "message": message})

    for row in rows:
        kind = (row.get("row_type") or "").strip().upper()
        if kind not in ("LINE", "OUTPUT"):
            err(row, "row_type must be LINE or OUTPUT")
            continue
        fin = products.get((row.get("finished_sku") or "").strip().upper())
        item = products.get((row.get("item_sku") or "").strip().upper())
        if fin is None:
            err(row, f"finished_sku '{row.get('finished_sku')}' is not a product")
            continue
        if item is None:
            err(row, f"item_sku '{row.get('item_sku')}' is not a product")
            continue
        try:
            qty = _dec(row.get("qty"))
            if qty is None or qty <= 0:
                raise ValueError("qty must be above zero")
            unit = None
            if (row.get("unit") or "").strip():
                unit = units.get(row["unit"].strip().lower())
                if unit is None:
                    raise ValueError(f"unit '{row['unit']}' does not exist")
            scrap = _dec(row.get("scrap_pct"), Decimal("0"))
            share = _dec(row.get("cost_share_pct"), Decimal("0"))
            recovery = _dec(row.get("recovery_rate"))
            out_qty = _dec(row.get("output_qty"), Decimal("1"))
            yld = _dec(row.get("yield_pct"), Decimal("100"))
        except ValueError as exc:
            err(row, str(exc))
            continue
        key = (fin.pk, (row.get("bom_name") or fin.name).strip())
        grp = groups.setdefault(key, {
            "product": fin, "name": key[1], "output_qty": out_qty, "yield_pct": yld, "lines": [], "outputs": [],
            "rows": [],
        })
        grp["rows"].append(row["_row"])
        if kind == "LINE":
            method = (row.get("issue_method") or "MANUAL").strip().upper()
            if method not in ("MANUAL", "BACKFLUSH"):
                err(row, "issue_method must be MANUAL or BACKFLUSH")
                continue
            grp["lines"].append({
                "component": item, "qty": qty, "unit": unit, "scrap_pct": scrap, "issue_method": method,
                "sequence": len(grp["lines"]),
            })
        else:
            okind = (row.get("kind") or "").strip().upper()
            if okind not in ("CO_PRODUCT", "BY_PRODUCT", "SCRAP"):
                err(row, "kind must be CO_PRODUCT, BY_PRODUCT or SCRAP")
                continue
            grp["outputs"].append({
                "product": item, "kind": okind, "qty": qty, "cost_share_pct": share,
                "recovery_rate": recovery,
            })
    for grp in groups.values():
        if not grp["lines"]:
            errors.append({"row": grp["rows"][0], "message": f"'{grp['name']}' has no material rows"})
            continue
        try:
            bom_svc.validate_content(
                company, grp["product"], grp["lines"], grp["outputs"],
                output_qty=grp["output_qty"], yield_pct=grp["yield_pct"],
            )
        except BusinessRuleError as exc:
            errors.append({"row": grp["rows"][0], "message": f"'{grp['name']}': {exc.detail}"})
    summary = [
        {"product": g["product"].name, "name": g["name"], "materials": len(g["lines"]),
         "outputs": len(g["outputs"]), "first_row": g["rows"][0]}
        for g in groups.values()
    ]
    result = {"recipes": summary, "errors": errors, "created": 0, "committed": False}
    if errors or not commit:
        return result
    with transaction.atomic():
        for grp in groups.values():
            bom = Bom.objects.create(
                company=company, product=grp["product"], name=grp["name"], status=Bom.Status.DRAFT,
                revision_no=bom_svc.next_revision_no(company, grp["product"]),
                output_qty=grp["output_qty"], yield_pct=grp["yield_pct"], change_reason="Imported",
                created_by=user, updated_by=user,
            )
            for ln in grp["lines"]:
                BomLine.objects.create(bom=bom, company=company, **ln)
            for out in grp["outputs"]:
                BomOutput.objects.create(bom=bom, company=company, **out)
            if activate:
                bom_svc.activate_revision(bom, user)
            result["created"] += 1
    result["committed"] = True
    return result


def export_csv(company, *, only_default=True) -> str:
    """Active recipes in the import shape, so a file can go out, be edited and come back in."""
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(COLUMNS)
    qs = Bom.objects.filter(company=company, status=Bom.Status.ACTIVE)
    if only_default:
        qs = qs.filter(is_default=True)
    for bom in qs.select_related("product").prefetch_related("lines__component", "lines__unit", "outputs__product"):
        for ln in bom.lines.all():
            w.writerow([csv_safe(v) for v in (
                "LINE", bom.product.sku, bom.name, bom.output_qty, bom.yield_pct, ln.component.sku, ln.qty,
                ln.unit.short_name if ln.unit_id else "", ln.scrap_pct, ln.issue_method, "", "", "",
            )])
        for o in bom.outputs.all():
            w.writerow([csv_safe(v) for v in (
                "OUTPUT", bom.product.sku, bom.name, bom.output_qty, bom.yield_pct, o.product.sku, o.qty, "", "",
                "", o.kind, o.cost_share_pct, "" if o.recovery_rate is None else o.recovery_rate,
            )])
    return out.getvalue()
