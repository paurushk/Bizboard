from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.db import transaction
from django.db.models import ProtectedError, Sum
from django.http import HttpResponse
from rest_framework import status
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from django.utils import timezone

from billing.permissions import SubscriptionWritesAllowed
from core.exceptions import BusinessRuleError
from core.permissions import (
    CanCreatePayments,
    CanExport,
    CanPostJournals,
    CanViewFinancialReports,
    HasCompany,
    IsOwner,
    get_company_user,
)
from core.viewsets import CompanyScopedViewSet

from .expense_posting import pending_summary, post_pending_expenses, sync_expense_posting
from .journal_sources import source_label
from .models import Account, AccountingPeriod, BankReconSession, CashShiftRegister, CostCenter, Expense, ExpenseCategoryAccount, FixedAsset, JournalEntry, JournalLine, RecurringExpense
from .reports import balance_sheet, cash_flow, close_financial_year, profit_and_loss, trial_balance
from .serializers import (
    AccountSerializer, AccountingPeriodSerializer, AccountingSettingsSerializer,
    BankReconSessionSerializer, CashShiftRegisterSerializer, CostCenterSerializer, ExpenseCategoryAccountSerializer,
    ExpenseSerializer, RecurringExpenseSerializer,
    FixedAssetSerializer, JournalEntrySerializer, UnreconciledGlLineSerializer,
)
from .services import BooksHealthService, PostingService, seed_chart_of_accounts


# CR-089: Enabling books seeds the CoA. Historical documents created while
# books were disabled require running backfill_accounting_postings to generate GL entries.
class AccountingEnabledMixin:
    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not get_company_user(request).company.accounting_enabled:
            raise BusinessRuleError("Accounting is not enabled for this company.")


_MUTATE_ACTIONS = ("create", "update", "partial_update", "destroy")


class AccountViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    queryset = Account.objects.all()
    serializer_class = AccountSerializer

    def get_permissions(self):
        action_name = getattr(self, "action", None)
        if action_name in (*_MUTATE_ACTIONS, "transfer_balance"):
            return [IsAuthenticated(), HasCompany(), CanPostJournals()]
        if action_name == "export":
            return [IsAuthenticated(), HasCompany(), CanViewFinancialReports(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get_queryset(self):
        qs = super().get_queryset().select_related("parent")
        if getattr(self, "action", None) not in ("list", "export"):
            return qs
        params = self.request.query_params
        if params.get("type"):
            qs = qs.filter(type=params["type"].upper())
        active = (params.get("active") or "").lower()
        if active in ("true", "1", "yes"):
            qs = qs.filter(is_active=True)
        elif active in ("false", "0", "no"):
            qs = qs.filter(is_active=False)
        if params.get("q"):
            from django.db.models import Q

            qs = qs.filter(Q(code__icontains=params["q"]) | Q(name__icontains=params["q"]))
        return qs

    def get_serializer_context(self):
        ctx = super().get_serializer_context()
        if getattr(self, "action", None) in ("list", "export"):
            used = set(
                JournalLine.objects.filter(company=self.company).values_list("account_id", flat=True).distinct()
            )
            ctx["used_account_ids"] = used
            if self.request.query_params.get("with_balance") in ("1", "true", "yes"):
                from datetime import date as date_cls

                from .reports import _balances

                raw = self.request.query_params.get("as_of")
                try:
                    as_of = date_cls.fromisoformat(raw[:10]) if raw else None
                except ValueError:
                    raise BusinessRuleError("as_of must be YYYY-MM-DD.")
                ctx["account_balances"] = {
                    row["account_id"]: row["balance"] for row in _balances(self.company, as_of=as_of)
                }
        return ctx

    @action(detail=False, methods=["get"])
    def export(self, request):
        """ACS-50: the chart (with balances when asked) as CSV or XLSX."""
        rows = AccountSerializer(
            self.filter_queryset(self.get_queryset()), many=True, context=self.get_serializer_context()
        ).data
        header = ["Code", "Name", "Type", "Parent", "Active", "Balance"]
        data = [
            [r["code"], r["name"], r["type"], r["parent_name"], "Yes" if r["is_active"] else "No", r["balance"] or ""]
            for r in rows
        ]
        if request.query_params.get("file") == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Chart of accounts"
            ws.append(header)
            for r in data:
                ws.append(r)
            out = BytesIO()
            wb.save(out)
            resp = HttpResponse(
                out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            resp["Content-Disposition"] = 'attachment; filename="chart-of-accounts.xlsx"'
            return resp
        import csv

        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = 'attachment; filename="chart-of-accounts.csv"'
        writer = csv.writer(resp)
        writer.writerow(header)
        for r in data:
            writer.writerow([("'" + c) if isinstance(c, str) and c[:1] in ("=", "+", "-", "@") else c for c in r])
        return resp

    @action(detail=True, methods=["post"], url_path="transfer-balance")
    @transaction.atomic
    def transfer_balance(self, request, pk=None):
        """ACS-50: draft a journal that moves this account's balance to another, so it can then be deactivated."""
        from datetime import date as date_cls

        from core.services.document_numbers import DocumentNumberService, resolve_series_gstin

        source = self.get_object()
        if source.is_system:
            raise BusinessRuleError("System accounts cannot be emptied here.")
        target = Account.objects.filter(company=self.company, pk=request.data.get("to_account"), is_active=True).first()
        if target is None or target.pk == source.pk:
            raise BusinessRuleError("Choose a different active account to move the balance to.")
        agg = source.journal_lines.filter(
            entry__status__in=[JournalEntry.Status.POSTED, JournalEntry.Status.REVERSED]
        ).aggregate(d=Sum("debit"), c=Sum("credit"))
        balance = (agg["d"] or Decimal("0")) - (agg["c"] or Decimal("0"))
        if balance == 0:
            raise BusinessRuleError("This account has no balance to move.")
        raw = request.data.get("entry_date")
        entry_date = date_cls.fromisoformat(str(raw)[:10]) if raw else timezone.localdate()
        number = DocumentNumberService.next_number(
            self.company, "JOURNAL_ENTRY", gstin=resolve_series_gstin(self.company), on_date=entry_date,
        )
        entry = JournalEntry.objects.create(
            company=self.company, status=JournalEntry.Status.DRAFT, number=number, entry_date=entry_date,
            narration=f"Move balance of {source.code} {source.name} to {target.code} {target.name}",
            created_by=request.user, updated_by=request.user,
        )
        amount = abs(balance)
        # Clearing a debit balance credits the source and debits the target, and the reverse.
        if balance > 0:
            lines = [(source, 0, amount), (target, amount, 0)]
        else:
            lines = [(source, amount, 0), (target, 0, amount)]
        JournalLine.objects.bulk_create([
            JournalLine(company=self.company, entry=entry, account=a, debit=Decimal(d), credit=Decimal(c))
            for a, d, c in lines
        ])
        return Response(
            {"journal_id": entry.pk, "number": entry.number, "amount": str(amount)}, status=status.HTTP_201_CREATED,
        )

    def perform_create(self, serializer):
        serializer.save(company=self.company, created_by=self.request.user, updated_by=self.request.user)
        self._audit("CREATE", serializer.instance)

    def perform_update(self, serializer):
        if serializer.instance.is_system:
            raise BusinessRuleError("System accounts cannot be edited.")
        serializer.save(updated_by=self.request.user, company=self.company)
        self._audit("UPDATE", serializer.instance)

    def perform_destroy(self, instance):
        if instance.is_system:
            raise BusinessRuleError("System accounts cannot be deleted.")
        try:
            super().perform_destroy(instance)
        except ProtectedError:
            # B1-004: account is referenced by postings / fixed assets / recon.
            raise BusinessRuleError(
                "This account is used by existing postings and cannot be "
                "deleted. Deactivate it instead."
            )


class PeriodViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    queryset = AccountingPeriod.objects.all()
    serializer_class = AccountingPeriodSerializer

    def get_permissions(self):
        # BB-000453: period CLOSE / mutate is Owner-only (not Accountant CanPostJournals).
        if getattr(self, "action", None) in (*_MUTATE_ACTIONS, "soft_close", "close", "close_period", "reopen", "generate"):
            return [IsAuthenticated(), HasCompany(), IsOwner()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def perform_create(self, serializer):
        serializer.save(company=self.company, created_by=self.request.user, updated_by=self.request.user)

    def perform_update(self, serializer):
        # B1-030: keep updated_by fresh and re-run the serializer overlap guard.
        serializer.save(company=self.company, updated_by=self.request.user)

    def _refuse_if_earlier_open(self, period):
        earlier_open = AccountingPeriod.objects.filter(
            company=self.company,
            end_date__lt=period.start_date,
            status=AccountingPeriod.Status.OPEN,
        ).order_by("start_date").first()
        if earlier_open is not None:
            raise BusinessRuleError(
                f"Close {earlier_open.name} first — periods must be closed in order."
            )

    @action(detail=True, methods=["get"], url_path="close-readiness")
    def close_readiness(self, request, pk=None):
        """ACS-30: why this month can or cannot close, with a fix link for each blocker."""
        from . import period_ops

        return Response(period_ops.close_readiness(self.company, self.get_object()))

    @action(detail=False, methods=["get"], url_path="close-summary")
    def close_summary(self, request):
        """ACS-31: one cheap row per month for the Periods table."""
        from . import period_ops

        return Response({"results": period_ops.close_summary(self.company)})

    @action(detail=True, methods=["post"])
    def reopen(self, request, pk=None):
        """ACS-32: step a closed month back (Closed to Soft-closed, Soft-closed to Open). Owner only, reason required."""
        from . import period_ops

        period = self.get_object()
        reopened = period_ops.reopen_period(self.company, period.pk, request.user, request.data.get("reason", ""))
        return Response(self.get_serializer(reopened).data)

    @action(detail=False, methods=["post"])
    def generate(self, request):
        """ACS-33: create the twelve months of a financial year in one go."""
        from . import period_ops

        try:
            year = int(request.data.get("fy_start_year") or request.data.get("fyStartYear"))
        except (TypeError, ValueError):
            raise BusinessRuleError("Choose the year the financial year starts in.")
        return Response(period_ops.generate_year(self.company, year, request.user), status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"], url_path="soft-close")
    @transaction.atomic
    def soft_close(self, request, pk=None):
        # CR-081: lock period row so concurrent PostingService.post waits.
        from django.shortcuts import get_object_or_404

        period = get_object_or_404(self.get_queryset().select_for_update(), pk=pk)
        if period.status != AccountingPeriod.Status.OPEN:
            raise BusinessRuleError("Only open periods can be soft closed.")
        self._refuse_if_earlier_open(period)
        BooksHealthService.assert_period_close_allowed(
            self.company, period=BooksHealthService._period_label(period),
        )
        period.status = AccountingPeriod.Status.SOFT_CLOSED
        period.updated_by = request.user
        period.save(update_fields=["status", "updated_by", "updated_at"])
        from reporting.gst_periods import gst_period_open_warnings

        data = self.get_serializer(period).data
        data["warnings"] = gst_period_open_warnings(
            self.company, period.start_date, period.end_date,
        )
        return Response(data)

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def close(self, request, pk=None):
        # CR-081: lock period row so concurrent PostingService.post waits.
        from django.shortcuts import get_object_or_404

        period = get_object_or_404(self.get_queryset().select_for_update(), pk=pk)
        if period.status == AccountingPeriod.Status.CLOSED:
            raise BusinessRuleError("Period is already closed.")
        # B1-022: no non-contiguous close — an earlier OPEN period would let
        # back-dated entries land in it after this one is "closed".
        self._refuse_if_earlier_open(period)
        BooksHealthService.assert_period_close_allowed(
            self.company, period=BooksHealthService._period_label(period),
        )
        period.status = AccountingPeriod.Status.CLOSED
        period.updated_by = request.user
        period.save(update_fields=["status", "updated_by", "updated_at"])
        from reporting.gst_periods import gst_period_open_warnings

        data = self.get_serializer(period).data
        data["warnings"] = gst_period_open_warnings(
            self.company, period.start_date, period.end_date,
        )
        return Response(data)

    @action(detail=True, methods=["post"], url_path="close-period")
    @transaction.atomic
    def close_period(self, request, pk=None):
        """Close the books for the month, then the GST period. Two steps, one call."""
        from django.shortcuts import get_object_or_404

        from reporting.gst_periods import soft_close_period
        from reporting.models import GstReturnPeriod

        period = get_object_or_404(self.get_queryset().select_for_update(), pk=pk)
        if period.status != AccountingPeriod.Status.CLOSED:
            self._refuse_if_earlier_open(period)
            BooksHealthService.assert_period_close_allowed(
                self.company, period=BooksHealthService._period_label(period),
            )
            period.status = AccountingPeriod.Status.CLOSED
            period.updated_by = request.user
            period.save(update_fields=["status", "updated_by", "updated_at"])
        month = f"{period.start_date:%Y-%m}"
        gst = GstReturnPeriod.objects.filter(company=self.company, period=month).first()
        if gst is None or gst.status == GstReturnPeriod.Status.OPEN:
            gst = soft_close_period(self.company, month, request.user)
        data = self.get_serializer(period).data
        data["gst_period_status"] = gst.status
        data["books_closed"] = period.status == AccountingPeriod.Status.CLOSED
        data["gst_closed"] = gst.status != GstReturnPeriod.Status.OPEN
        return Response(data)


class CostCenterViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    queryset = CostCenter.objects.all()
    serializer_class = CostCenterSerializer

    def get_permissions(self):
        if getattr(self, "action", None) in _MUTATE_ACTIONS:
            return [IsAuthenticated(), HasCompany(), CanPostJournals()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get_queryset(self):
        qs = super().get_queryset().select_related("parent")
        if getattr(self, "action", None) == "list":
            active = (self.request.query_params.get("active") or "").lower()
            if active in ("true", "1"):
                qs = qs.filter(is_active=True)
            elif active in ("false", "0"):
                qs = qs.filter(is_active=False)
        return qs

    def get_serializer_context(self):
        ctx = super().get_serializer_context()
        if getattr(self, "action", None) == "list":
            from django.db.models import Count

            from purchases.models import PurchaseInvoice
            from sales.models import SalesInvoice

            usage = {}

            def bump(rows, key):
                for row in rows:
                    slot = usage.setdefault(row["cost_center_id"], {"journal_lines": 0, "invoices": 0, "expenses": 0})
                    slot[key] += row["n"]

            bump(JournalLine.objects.filter(company=self.company, cost_center__isnull=False).values("cost_center_id").annotate(n=Count("id")), "journal_lines")
            bump(SalesInvoice.objects.filter(company=self.company, cost_center__isnull=False).values("cost_center_id").annotate(n=Count("id")), "invoices")
            bump(PurchaseInvoice.objects.filter(company=self.company, cost_center__isnull=False).values("cost_center_id").annotate(n=Count("id")), "invoices")
            bump(Expense.objects.filter(company=self.company, cost_center__isnull=False).values("cost_center_id").annotate(n=Count("id")), "expenses")
            ctx["cost_center_usage"] = usage
        return ctx

    def perform_create(self, serializer):
        serializer.save(company=self.company, created_by=self.request.user, updated_by=self.request.user)
        self._audit("CREATE", serializer.instance)

    @action(detail=False, methods=["get"])
    def report(self, request):
        """ACS-51: income, expense and net by cost centre for a date range, with an unallocated bucket."""
        from datetime import date as date_cls

        def parse(name):
            raw = request.query_params.get(name)
            if not raw:
                return None
            try:
                return date_cls.fromisoformat(raw[:10])
            except ValueError:
                raise BusinessRuleError(f"{name} must be YYYY-MM-DD.")

        date_from, date_to = parse("date_from"), parse("date_to")
        qs = JournalLine.objects.filter(
            company=self.company,
            entry__status__in=[JournalEntry.Status.POSTED, JournalEntry.Status.REVERSED],
            account__type__in=(Account.Type.INCOME, Account.Type.EXPENSE),
        ).exclude(entry__purpose="FY_CLOSE").exclude(entry__reversal_of__purpose="FY_CLOSE")
        if date_from:
            qs = qs.filter(entry__entry_date__gte=date_from)
        if date_to:
            qs = qs.filter(entry__entry_date__lte=date_to)
        rows = qs.values("cost_center_id", "account__type").annotate(d=Sum("debit"), c=Sum("credit"))
        buckets = {}
        for row in rows:
            slot = buckets.setdefault(row["cost_center_id"], {"income": Decimal("0"), "expense": Decimal("0")})
            d, c = row["d"] or Decimal("0"), row["c"] or Decimal("0")
            if row["account__type"] == Account.Type.INCOME:
                slot["income"] += c - d
            else:
                slot["expense"] += d - c
        names = {cc.pk: cc for cc in CostCenter.objects.filter(company=self.company)}
        results = []
        total_income = total_expense = Decimal("0")
        for cc_id, vals in sorted(buckets.items(), key=lambda kv: (kv[0] is None, names[kv[0]].code if kv[0] in names else "")):
            total_income += vals["income"]
            total_expense += vals["expense"]
            cc = names.get(cc_id)
            results.append({
                "cost_center": cc_id,
                "code": cc.code if cc else "",
                "name": cc.name if cc else "Unallocated",
                "income": str(vals["income"]),
                "expense": str(vals["expense"]),
                "net": str(vals["income"] - vals["expense"]),
            })
        return Response({
            "results": results,
            "totals": {
                "income": str(total_income), "expense": str(total_expense), "net": str(total_income - total_expense),
            },
        })

    def perform_destroy(self, instance):
        try:
            super().perform_destroy(instance)
        except ProtectedError:
            # B1-004: cost centre is referenced by journal lines.
            raise BusinessRuleError(
                "This cost centre is used by existing postings and cannot be "
                "deleted. Deactivate it instead."
            )


class JournalViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    """BB-000085: company-scoped journals (was bare ModelViewSet)."""

    queryset = JournalEntry.objects.all()
    serializer_class = JournalEntrySerializer
    # ACS-20: PUT/PATCH/DELETE reach drafts only; the guard is in update()/destroy().
    http_method_names = ["get", "post", "put", "patch", "delete", "head", "options"]

    def get_permissions(self):
        # BB-000200/316: create / post / reverse journals requires Owner or
        # can_post_journals.
        action_name = getattr(self, "action", None)
        if action_name in ("create", "post", "reverse", "update", "partial_update", "destroy"):
            return [IsAuthenticated(), HasCompany(), CanPostJournals()]
        if action_name == "export":
            return [IsAuthenticated(), HasCompany(), CanViewFinancialReports(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get_queryset(self):
        qs = (
            super().get_queryset()
            .select_related("posted_by", "reversal_of")
            .prefetch_related("lines__account", "lines__cost_center", "lines__customer", "lines__supplier")
        )
        if getattr(self, "action", None) not in ("list", "export"):
            return qs
        return self._filtered(qs)

    def _filtered(self, qs):
        """ACS-23: server-side register filters. Discarded drafts are hidden unless asked for."""
        from django.db.models import Q

        params = self.request.query_params
        status_param = (params.get("status") or "").strip().upper()
        if status_param:
            qs = qs.filter(status__in=[v for v in status_param.split(",") if v])
        else:
            qs = qs.exclude(status=JournalEntry.Status.DISCARDED)
        if params.get("date_from"):
            qs = qs.filter(entry_date__gte=params["date_from"])
        if params.get("date_to"):
            qs = qs.filter(entry_date__lte=params["date_to"])
        if params.get("number"):
            qs = qs.filter(number__icontains=params["number"].strip())
        source = (params.get("source_type") or "").strip()
        if source:
            qs = qs.filter(source_type=source)
        q = (params.get("q") or "").strip()
        if q:
            qs = qs.filter(
                Q(narration__icontains=q) | Q(reference__icontains=q) | Q(lines__narration__icontains=q)
            ).distinct()
        if params.get("account"):
            qs = qs.filter(lines__account_id=params["account"]).distinct()
        return qs

    @transaction.atomic
    def create(self, request, *args, **kwargs):
        # B1-002: header + lines must be all-or-nothing, and next_number's
        # select_for_update guard needs an open transaction to be effective.
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        # ACS-24: a repeated create with the same token (double click, retry) returns the first draft.
        token = (serializer.validated_data.get("client_token") or "").strip()
        if token:
            existing = JournalEntry.objects.filter(company=self.company, client_token=token).first()
            if existing is not None:
                return Response(self.get_serializer(existing).data, status=status.HTTP_200_OK)
        lines = serializer.validated_data.pop("lines")
        # B1-031: cross-tenant account / cost_center / bank_statement_line checks
        # live in JournalLineSerializer (validate_account / validate_cost_center;
        # bank_statement_line is read-only per B1-003) — no duplicate view loop.
        number = (serializer.validated_data.get("number") or "").strip()
        from core.services.document_numbers import DocumentNumberService, resolve_series_gstin

        if not number:
            number = DocumentNumberService.next_number(
                self.company,
                "JOURNAL_ENTRY",
                gstin=resolve_series_gstin(self.company),
                on_date=serializer.validated_data.get("entry_date"),
            )
        else:
            # B1-013: a client-supplied number is otherwise fully trusted and
            # can collide with a later auto-allocated one -- bump the series
            # counter past it when it looks like one of ours.
            DocumentNumberService.claim_number(
                self.company,
                "JOURNAL_ENTRY",
                number,
                gstin=resolve_series_gstin(self.company),
                on_date=serializer.validated_data.get("entry_date"),
            )
        entry = JournalEntry.objects.create(company=self.company, status=JournalEntry.Status.DRAFT,
            number=number, entry_date=serializer.validated_data.get("entry_date"),
            narration=serializer.validated_data.get("narration", ""),
            reference=serializer.validated_data.get("reference", ""),
            client_token=token,
            created_by=request.user, updated_by=request.user)
        JournalLine.objects.bulk_create([
            JournalLine(company=entry.company, entry=entry, **line) for line in lines
        ])
        return Response(self.get_serializer(entry).data, status=status.HTTP_201_CREATED)

    def _draft_for_change(self, pk):
        from django.shortcuts import get_object_or_404

        entry = get_object_or_404(JournalEntry.objects.select_for_update().filter(company=self.company), pk=pk)
        if entry.status != JournalEntry.Status.DRAFT:
            raise BusinessRuleError(
                "Only a draft journal can be changed. Reverse a posted manual journal instead.",
                code="journal_not_draft",
            )
        return entry

    @transaction.atomic
    def update(self, request, *args, **kwargs):
        """ACS-20: edit a draft. Header and lines are replaced together; the number stays."""
        entry = self._draft_for_change(kwargs.get("pk"))
        serializer = self.get_serializer(entry, data=request.data, partial=bool(kwargs.get("partial")))
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        lines = data.pop("lines", None)
        for field in ("entry_date", "narration", "reference"):
            if field in data:
                setattr(entry, field, data[field])
        entry.updated_by = request.user
        entry.save()
        if lines is not None:
            entry.lines.all().delete()
            JournalLine.objects.bulk_create([JournalLine(company=entry.company, entry=entry, **line) for line in lines])
        from core.services.audit import record_document_event

        record_document_event(document=entry, user=request.user, event="journal.draft_edited", before={"status": "DRAFT"})
        entry = self.get_queryset().get(pk=entry.pk)
        return Response(self.get_serializer(entry).data)

    @transaction.atomic
    def destroy(self, request, *args, **kwargs):
        """ACS-20: discard a draft. The number is kept (status DISCARDED) so the series has no silent gap."""
        entry = self._draft_for_change(kwargs.get("pk"))
        entry.status = JournalEntry.Status.DISCARDED
        entry.updated_by = request.user
        entry.save(update_fields=["status", "updated_by", "updated_at"])
        from core.services.audit import record_document_event

        record_document_event(document=entry, user=request.user, event="journal.draft_discarded", before={"status": "DRAFT"})
        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=False, methods=["get"])
    def export(self, request):
        """ACS-23: the filtered register as CSV (or XLSX with ?file=xlsx), one row per line."""
        qs = self._filtered(self.get_queryset()).order_by("entry_date", "id")[:20000]
        header = ["Number", "Date", "Status", "Source", "Narration", "Reference", "Account code", "Account",
                  "Party", "Cost centre", "Line narration", "Debit", "Credit"]
        rows = []
        for entry in qs:
            for line in entry.lines.all():
                rows.append([
                    entry.number, entry.entry_date.isoformat(), entry.status,
                    source_label(entry.source_type) if entry.source_type else "Manual entry",
                    entry.narration, entry.reference, line.account.code, line.account.name,
                    (line.customer.name if line.customer_id else (line.supplier.name if line.supplier_id else "")),
                    line.cost_center.name if line.cost_center_id else "", line.narration,
                    str(line.debit), str(line.credit),
                ])
        if request.query_params.get("file") == "xlsx":
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Journals"
            sheet.append(header)
            for row in rows:
                sheet.append(row)
            output = BytesIO()
            workbook.save(output)
            response = HttpResponse(
                output.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = 'attachment; filename="journals.xlsx"'
            return response
        import csv

        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="journals.csv"'
        writer = csv.writer(response)
        writer.writerow(header)
        for row in rows:
            # A cell that starts like a formula would run in Excel; prefix it.
            writer.writerow([("'" + c) if isinstance(c, str) and c[:1] in ("=", "+", "-", "@") else c for c in row])
        return response

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def post(self, request, pk=None):
        # CR-156: assert + status flip in one atomic so soft_close cannot race.
        entry = JournalEntry.objects.select_for_update().filter(company=self.company, pk=pk).first()
        if entry is None:
            from django.http import Http404

            raise Http404
        if entry.status != JournalEntry.Status.DRAFT:
            raise BusinessRuleError("Only draft journals can be posted.", code="journal_not_draft")
        totals = entry.lines.aggregate(debit=Sum("debit"), credit=Sum("credit"))
        if not entry.lines.exists() or (totals["debit"] or 0) != (totals["credit"] or 0):
            raise BusinessRuleError("Journal lines must balance before posting.")
        # B1-006: use the shared period guard so the manual post path enforces
        # the GST-period lock and the ACC-04 "must be an open period" opt-in,
        # not just the AccountingPeriod CLOSED/SOFT_CLOSED check it had.
        from reporting.gst_periods import assert_period_allows_money_amend

        assert_period_allows_money_amend(self.company, entry.entry_date)
        # CR-086: manual journals respect books_start_date (PostingService.post parity).
        cutover = getattr(self.company, "books_start_date", None)
        if cutover and entry.entry_date and entry.entry_date < cutover:
            raise BusinessRuleError(
                f"{entry.entry_date} is before the books start date ({cutover}). "
                "Manual journals cannot post before cutover."
            )
        entry.status = JournalEntry.Status.POSTED
        entry.source_type, entry.source_id, entry.purpose = "MANUAL_JOURNAL", entry.id, "POST"
        entry.posted_at, entry.posted_by = timezone.now(), request.user
        entry.save(update_fields=["status", "source_type", "source_id", "purpose", "posted_at", "posted_by", "updated_at"])
        # The rollup checkpoint is keyed on line ids and a draft's lines already existed, so posting
        # one adds no new id. Drop the checkpoint so the next report rebuilds instead of trusting it.
        from .models import AccountBalanceRollup

        AccountBalanceRollup.objects.filter(company=self.company).delete()
        from core.services.audit import record_document_event

        record_document_event(
            document=entry, user=request.user, event="journal.posted",
            before={"status": "DRAFT"},
        )
        return Response(self.get_serializer(entry).data)

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def reverse(self, request, pk=None):
        # CR-079: reverse JE + status flip share one transaction with PostingService.reverse.
        entry = self.get_object()
        # ACS-10: only a hand-posted journal is reversed from here. A system
        # journal belongs to its source document; reversing it alone leaves the
        # document completed with no ledger entry (and `PostingService.reverse`
        # lets non-manual entries through a soft-closed month).
        if entry.source_type != "MANUAL_JOURNAL":
            raise BusinessRuleError(
                "This journal was posted by "
                f"{source_label(entry.source_type)}. Cancel or amend that document to reverse its "
                "ledger entry; only manual journals can be reversed here.",
                code="journal_system_not_reversible",
            )
        return Response(self.get_serializer(PostingService.reverse(entry, request.user)).data)

    @action(detail=False, methods=["get"], url_path="unreconciled-lines")
    def unreconciled_lines(self, request):
        """F2-028: still-unreconciled GL lines for one account, resolved
        server-side. The bank-recon picker used to page the most recent 100
        journals client-side, so a GL line older than that could never be
        matched to a bank line.
        """
        account_id = request.query_params.get("account")
        if not account_id:
            raise BusinessRuleError("account query parameter is required.")
        try:
            account = Account.objects.get(pk=account_id, company=self.company)
        except (Account.DoesNotExist, ValueError, TypeError):
            raise BusinessRuleError("Unknown account.")
        lines = (
            JournalLine.objects.filter(
                company=self.company,
                account=account,
                entry__status=JournalEntry.Status.POSTED,
                bank_statement_line__isnull=True,
                reconciled_at__isnull=True,
            )
            .select_related("entry")
            .order_by("-entry__entry_date", "-entry__id", "-id")
        )
        page = self.paginate_queryset(lines)
        target = page if page is not None else lines
        serializer = UnreconciledGlLineSerializer(target, many=True)
        if page is not None:
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)


class BankReconSessionViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    queryset = BankReconSession.objects.all()
    serializer_class = BankReconSessionSerializer
    http_method_names = ["get", "post", "patch", "head", "options"]

    _WRITE_ACTIONS = (
        "create", "partial_update", "match", "match_many", "unmatch", "accept_suggestions", "create_entry", "finalise",
    )

    def get_permissions(self):
        action_name = getattr(self, "action", None)
        if action_name == "reopen":
            return [IsAuthenticated(), HasCompany(), IsOwner()]
        if action_name in self._WRITE_ACTIONS:
            return [IsAuthenticated(), HasCompany(), CanPostJournals()]
        if action_name == "report":
            return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get_queryset(self):
        qs = super().get_queryset().select_related("account", "statement__bank_account", "finalised_by")
        if getattr(self, "action", None) == "list":
            params = self.request.query_params
            if params.get("status"):
                qs = qs.filter(status__in=[v for v in params["status"].upper().split(",") if v])
            if params.get("account"):
                qs = qs.filter(account_id=params["account"])
        return qs.order_by("-id")

    def get_serializer_context(self):
        ctx = super().get_serializer_context()
        if getattr(self, "action", None) in ("retrieve", "create", "finalise", "reopen"):
            ctx["with_summary"] = True
        return ctx

    def perform_create(self, serializer):
        from . import bank_match

        account = serializer.validated_data["account"]
        statement = serializer.validated_data["statement"]
        if account.company_id != self.company.id or statement.company_id != self.company.id:
            raise BusinessRuleError("Account and statement must belong to this company.")
        gl = JournalLine.objects.filter(
            entry__company=self.company,
            entry__status=JournalEntry.Status.POSTED,
            account=account,
        ).aggregate(d=Sum("debit"), c=Sum("credit"))
        statement_balance = statement.lines.aggregate(total=Sum("amount"))["total"] or Decimal("0")
        gl_balance = (gl.get("d") or Decimal("0")) - (gl.get("c") or Decimal("0"))
        session = serializer.save(
            company=self.company,
            gl_balance=gl_balance,
            statement_balance=statement_balance,
            created_by=self.request.user,
            updated_by=self.request.user,
        )
        if session.statement_date is None:
            session.statement_date = bank_match.statement_date_for(session)
            session.save(update_fields=["statement_date"])
        bank_match.refresh_status(session)

    def _session(self):
        return self.get_object()

    def partial_update(self, request, *args, **kwargs):
        """ACS-41: set the statement date and the bank's closing balance. Nothing else changes here."""
        from . import bank_match

        session = self.get_object()
        if session.status == BankReconSession.Status.FINALISED:
            raise BusinessRuleError("This reconciliation is finalised. Reopen it to change it.", code="recon_finalised")
        data = {k: request.data[k] for k in ("statement_date", "statement_closing_balance") if k in request.data}
        serializer = self.get_serializer(session, data=data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save(updated_by=request.user)
        bank_match.refresh_status(session)
        session.refresh_from_db()
        return Response(self.get_serializer(session).data)

    def _page(self, qs, rows):
        page = self.paginate_queryset(qs)
        target = page if page is not None else list(qs)
        data = [rows(x) for x in target]
        return self.get_paginated_response(data) if page is not None else Response(data)

    @action(detail=True, methods=["post"])
    def match(self, request, pk=None):
        from . import bank_match

        session = self._session()
        jl, bl = request.data.get("journal_line"), request.data.get("bank_statement_line")
        if not jl or not bl:
            raise BusinessRuleError("An unreconciled journal line and bank statement line are required.")
        bank_match.match_gl(self.company, request.user, session, [jl], [bl])
        return Response({"ok": True, "match_kind": "gl_statement_line", "other_recon_path": "/payments/reconciliation"})

    @action(detail=True, methods=["post"], url_path="match-many")
    def match_many(self, request, pk=None):
        from . import bank_match

        session = self._session()
        result = bank_match.match_gl(
            self.company, request.user, session,
            request.data.get("journal_lines") or [], request.data.get("bank_lines") or [],
        )
        return Response({"ok": True, **result})

    @action(detail=True, methods=["post"])
    def unmatch(self, request, pk=None):
        from payments.models import BankStatementLine

        from . import bank_match

        session = self._session()
        line = BankStatementLine.objects.filter(
            company=self.company, pk=request.data.get("bank_line"), statement=session.statement
        ).first()
        if line is None:
            raise BusinessRuleError("Statement line not found on this reconciliation.")
        if session.status == BankReconSession.Status.FINALISED:
            raise BusinessRuleError("This reconciliation is finalised. Reopen it to change matches.", code="recon_finalised")
        return Response({"ok": True, **bank_match.unmatch_line(self.company, request.user, line)})

    @action(detail=True, methods=["get"], url_path="statement-lines")
    def statement_lines(self, request, pk=None):
        """ACS-42: the statement side of the workspace, filterable and paged."""
        from django.db.models import Q

        from payments.models import BankLineMatchStatus

        session = self._session()
        qs = session.statement.lines.all()
        params = request.query_params
        state = (params.get("status") or "open").lower()
        if state == "open":
            qs = qs.filter(match_status__in=[BankLineMatchStatus.UNMATCHED, BankLineMatchStatus.SUGGESTED])
        elif state == "matched":
            qs = qs.filter(match_status=BankLineMatchStatus.MATCHED)
        elif state != "all":
            qs = qs.filter(match_status=state.upper())
        if params.get("q"):
            qs = qs.filter(Q(narration__icontains=params["q"]) | Q(utr__icontains=params["q"]))
        if params.get("amount"):
            qs = qs.filter(amount=params["amount"])
        if params.get("date_from"):
            qs = qs.filter(txn_date__gte=params["date_from"])
        if params.get("date_to"):
            qs = qs.filter(txn_date__lte=params["date_to"])
        return self._page(qs.order_by("txn_date", "id"), lambda b: {
            "id": b.pk, "txn_date": b.txn_date.isoformat(), "narration": b.narration, "utr": b.utr,
            "amount": str(b.amount), "match_status": b.match_status,
        })

    @action(detail=True, methods=["get"], url_path="ledger-lines")
    def ledger_lines(self, request, pk=None):
        """ACS-42: the ledger side of the workspace: this bank's unreconciled lines, filterable and paged."""
        from django.db.models import Q

        from . import bank_match

        session = self._session()
        as_at = bank_match.statement_date_for(session)
        qs = JournalLine.objects.filter(
            company=self.company, account=session.account, entry__status=JournalEntry.Status.POSTED,
            bank_statement_line__isnull=True, reconciled_at__isnull=True,
            entry__entry_date__lte=as_at,
        ).select_related("entry")
        params = request.query_params
        if params.get("q"):
            qs = qs.filter(Q(entry__narration__icontains=params["q"]) | Q(entry__number__icontains=params["q"]))
        if params.get("date_from"):
            qs = qs.filter(entry__entry_date__gte=params["date_from"])
        if params.get("date_to"):
            qs = qs.filter(entry__entry_date__lte=params["date_to"])
        if params.get("amount"):
            amount = Decimal(params["amount"])
            qs = qs.filter(Q(debit=amount) | Q(credit=abs(amount)))
        return self._page(qs.order_by("entry__entry_date", "id"), lambda j: {
            "id": j.pk, "entry_id": j.entry_id, "entry_number": j.entry.number,
            "entry_date": j.entry.entry_date.isoformat(), "narration": j.entry.narration,
            "signed_amount": str((j.debit or Decimal("0")) - (j.credit or Decimal("0"))),
        })

    @action(detail=True, methods=["get"])
    def suggestions(self, request, pk=None):
        from . import bank_match

        session = self._session()
        window = int(request.query_params.get("window") or 5)
        return Response({"results": bank_match.suggestions(session, window_days=max(0, min(window, 30)))})

    @action(detail=True, methods=["post"], url_path="accept-suggestions")
    def accept_suggestions(self, request, pk=None):
        from . import bank_match

        session = self._session()
        pairs = request.data.get("pairs")
        if pairs is None:
            wanted = (request.data.get("confidence") or "high").lower()
            pairs = [
                {"bank_line": r["bank_line"], "journal_line": r["journal_line"]}
                for r in bank_match.suggestions(session)
                if r["confidence"] == wanted
            ]
        if request.data.get("preview") in (True, "true", 1, "1"):
            return Response({"preview": True, "pairs": pairs, "count": len(pairs)})
        return Response(bank_match.accept_suggestions(self.company, request.user, session, pairs))

    @action(detail=True, methods=["post"], url_path="create-entry")
    def create_entry(self, request, pk=None):
        from datetime import date as date_cls

        from . import bank_match

        session = self._session()
        raw = request.data.get("entry_date")
        try:
            entry_date = date_cls.fromisoformat(str(raw)[:10]) if raw else None
        except ValueError as exc:
            raise BusinessRuleError("entry_date must be YYYY-MM-DD.") from exc
        result = bank_match.create_entry_from_line(
            self.company, request.user, session, request.data.get("bank_line"), request.data.get("account"),
            narration=request.data.get("narration", ""), entry_date=entry_date,
        )
        return Response(result, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"])
    def finalise(self, request, pk=None):
        from . import bank_match

        cu = get_company_user(request)
        session = bank_match.finalise(
            self._session().pk, self.company, request.user,
            write_off=request.data.get("write_off") in (True, "true", 1, "1"),
            is_owner=bool(cu and cu.role == "OWNER"),
        )
        return Response(self.get_serializer(session).data)

    @action(detail=True, methods=["post"])
    def reopen(self, request, pk=None):
        from . import bank_match

        session = bank_match.reopen(self._session().pk, self.company, request.user, request.data.get("reason", ""))
        return Response(self.get_serializer(session).data)

    @action(detail=True, methods=["get"])
    def report(self, request, pk=None):
        """ACS-46: the reconciliation statement as JSON (default), XLSX or PDF."""
        from . import bank_match

        session = self._session()
        kind = (request.query_params.get("file") or "json").lower()
        if kind == "json":
            return Response(bank_match.reconciliation_report(session))
        # A file leaves the product, so it needs the export right.
        if not CanExport().has_permission(request, self):
            from rest_framework.exceptions import PermissionDenied

            raise PermissionDenied()
        if kind == "xlsx":
            resp = HttpResponse(
                bank_match.report_xlsx(session),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            resp["Content-Disposition"] = f'attachment; filename="bank-reconciliation-{session.pk}.xlsx"'
            return resp
        if kind == "pdf":
            resp = HttpResponse(bank_match.report_pdf(session), content_type="application/pdf")
            resp["Content-Disposition"] = f'attachment; filename="bank-reconciliation-{session.pk}.pdf"'
            return resp
        raise BusinessRuleError("file must be json, xlsx or pdf.")


class FixedAssetViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    queryset = FixedAsset.objects.all()
    serializer_class = FixedAssetSerializer

    def initial(self, request, *args, **kwargs):
        # Scope revision 2026-09-09b: D6 fixed assets is a KNOWN LIMITATION for
        # the pilot — the whole surface is inaccessible when ENABLE_FIXED_ASSETS
        # is off (backend/.env.pilot.example sets it to 0).
        from django.conf import settings
        from django.http import Http404

        if not getattr(settings, "ENABLE_FIXED_ASSETS", False):
            raise Http404()
        super().initial(request, *args, **kwargs)

    def get_permissions(self):
        action_name = getattr(self, "action", None)
        if action_name in (*_MUTATE_ACTIONS, "dispose", "run_depreciation", "register_check_draft"):
            return [IsAuthenticated(), HasCompany(), CanPostJournals()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get_serializer_context(self):
        ctx = super().get_serializer_context()
        if getattr(self, "action", None) == "list":
            last = {}
            for e in JournalEntry.objects.filter(
                company=self.company, source_type="FIXED_ASSET", status=JournalEntry.Status.POSTED,
                purpose__startswith="DEPRECIATION-",
            ).values("source_id", "entry_date"):
                month = e["entry_date"].strftime("%Y-%m")
                if month > last.get(e["source_id"], ""):
                    last[e["source_id"]] = month
            ctx["last_dep_month"] = last
        return ctx

    @transaction.atomic
    def perform_create(self, serializer):
        from .services import seed_chart_of_accounts

        accounts = seed_chart_of_accounts(self.company, self.request.user)
        validated = serializer.validated_data
        funding = validated.get("funding") or FixedAsset.Funding.ALREADY_BOOKED
        asset = serializer.save(
            company=self.company,
            created_by=self.request.user,
            updated_by=self.request.user,
            asset_account=validated.get("asset_account") or accounts["1600"],
            accumulated_depreciation_account=validated.get("accumulated_depreciation_account") or accounts["1650"],
            depreciation_expense_account=validated.get("depreciation_expense_account") or accounts["5300"],
        )
        if funding == FixedAsset.Funding.ALREADY_BOOKED:
            return
        # ACS-53: book the purchase so the register and the ledger agree from day one.
        if funding == FixedAsset.Funding.CASH:
            credit = {"account": accounts["1100"]}
        elif funding == FixedAsset.Funding.BANK:
            credit = {"account": PostingService._bank_gl_account(self.company, asset.funding_bank_account, asset.acquisition_date)}
        else:
            credit = {"account": accounts["2100"], "supplier": asset.funding_supplier}
        cost = asset.acquisition_cost
        entry = PostingService.post(
            company=self.company, source_type="FIXED_ASSET", source_id=asset.id, purpose="ACQUISITION",
            entry_date=asset.acquisition_date, user=self.request.user,
            narration=f"Acquisition: {asset.name}",
            lines=[{"account": asset.asset_account, "debit": cost}, {**credit, "credit": cost}],
        )
        if entry is None:
            raise BusinessRuleError(
                "The acquisition date is before the books start date, so the purchase cannot be booked. "
                "Register it as already booked, or use a later date.",
                code="asset_acquisition_before_books",
            )
        asset.acquisition_journal = entry
        asset.save(update_fields=["acquisition_journal", "updated_at"])

    @action(detail=False, methods=["get"])
    def presets(self, request):
        """ACS-55: starting points for life and rates. These are suggestions; confirm the rates with your CA."""
        return Response({"results": [
            {"key": "computers", "name": "Computers and laptops", "useful_life_months": 36, "method": "SLM", "wdv_annual_rate": "40.00"},
            {"key": "office_equipment", "name": "Office equipment", "useful_life_months": 60, "method": "SLM", "wdv_annual_rate": "15.00"},
            {"key": "furniture", "name": "Furniture and fittings", "useful_life_months": 120, "method": "SLM", "wdv_annual_rate": "10.00"},
            {"key": "vehicles", "name": "Vehicles", "useful_life_months": 96, "method": "SLM", "wdv_annual_rate": "15.00"},
            {"key": "plant", "name": "Plant and machinery", "useful_life_months": 180, "method": "SLM", "wdv_annual_rate": "15.00"},
            {"key": "buildings", "name": "Buildings", "useful_life_months": 360, "method": "SLM", "wdv_annual_rate": "10.00"},
        ], "note": "Rates are suggestions. Check them with your CA before relying on them."})

    @action(detail=False, methods=["get"])
    def summary(self, request):
        """ACS-52: totals for the cards, and the register against the ledger (ACS-53)."""
        from . import depreciation

        assets = list(FixedAsset.objects.filter(company=self.company, status=FixedAsset.Status.ACTIVE))
        cost = sum((a.acquisition_cost for a in assets), Decimal("0"))
        accum = sum((a.depreciated_amount for a in assets), Decimal("0"))
        month_charge = sum((a.monthly_depreciation for a in assets), Decimal("0"))
        gl = self._gl_totals()
        return Response({
            "count": len(assets), "total_cost": str(cost), "accumulated": str(accum), "net_book_value": str(cost - accum),
            "this_month_charge": str(month_charge),
            "register_vs_gl": {
                "register_cost": str(cost), "gl_cost": str(gl["cost"]), "cost_difference": str(cost - gl["cost"]),
                "register_accumulated": str(accum), "gl_accumulated": str(gl["accum"]),
                "accumulated_difference": str(accum - gl["accum"]),
            },
            "errors": FixedAsset.objects.filter(company=self.company, status=FixedAsset.Status.ACTIVE).exclude(last_depreciation_error="").count(),
            "last_completed_month": depreciation.last_completed_month_end().strftime("%Y-%m"),
        })

    def _gl_totals(self):
        def balance(code):
            agg = JournalLine.objects.filter(
                company=self.company, account__code=code,
                entry__status__in=[JournalEntry.Status.POSTED, JournalEntry.Status.REVERSED],
            ).aggregate(d=Sum("debit"), c=Sum("credit"))
            return (agg["d"] or Decimal("0")) - (agg["c"] or Decimal("0"))

        return {"cost": balance("1600"), "accum": -balance("1650")}

    @action(detail=False, methods=["post"], url_path="register-check/draft")
    @transaction.atomic
    def register_check_draft(self, request):
        """ACS-53: draft (not post) the journal that brings 1600/1650 in line with the register."""
        from core.services.document_numbers import DocumentNumberService, resolve_series_gstin

        assets = list(FixedAsset.objects.filter(company=self.company, status=FixedAsset.Status.ACTIVE))
        cost = sum((a.acquisition_cost for a in assets), Decimal("0"))
        accum = sum((a.depreciated_amount for a in assets), Decimal("0"))
        gl = self._gl_totals()
        cost_diff, accum_diff = cost - gl["cost"], accum - gl["accum"]
        if cost_diff == 0 and accum_diff == 0:
            raise BusinessRuleError("The register already agrees with the ledger.")
        accounts = seed_chart_of_accounts(self.company, request.user)
        equity = accounts["3200"]
        lines = []
        if cost_diff > 0:
            lines += [(accounts["1600"], cost_diff, 0), (equity, 0, cost_diff)]
        elif cost_diff < 0:
            lines += [(equity, -cost_diff, 0), (accounts["1600"], 0, -cost_diff)]
        if accum_diff > 0:
            lines += [(equity, accum_diff, 0), (accounts["1650"], 0, accum_diff)]
        elif accum_diff < 0:
            lines += [(accounts["1650"], -accum_diff, 0), (equity, 0, -accum_diff)]
        entry_date = timezone.localdate()
        number = DocumentNumberService.next_number(
            self.company, "JOURNAL_ENTRY", gstin=resolve_series_gstin(self.company), on_date=entry_date,
        )
        entry = JournalEntry.objects.create(
            company=self.company, status=JournalEntry.Status.DRAFT, number=number, entry_date=entry_date,
            narration="Bring fixed asset ledgers in line with the asset register (review before posting)",
            created_by=request.user, updated_by=request.user,
        )
        JournalLine.objects.bulk_create([
            JournalLine(company=self.company, entry=entry, account=a, debit=Decimal(d), credit=Decimal(c)) for a, d, c in lines
        ])
        return Response({"journal_id": entry.pk, "number": entry.number}, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["get"])
    def schedule(self, request, pk=None):
        """ACS-52: every month of this asset: charged, missing, or still to come."""
        from . import depreciation

        return Response(depreciation.schedule(self.get_object()))

    @action(detail=False, methods=["post"], url_path="run-depreciation")
    def run_depreciation(self, request):
        """ACS-54: preview (dry_run) or post depreciation for every missing month up to a chosen month."""
        from datetime import date as date_cls

        from . import depreciation

        raw = request.data.get("month")
        if raw:
            try:
                year, month = int(str(raw)[:4]), int(str(raw)[5:7])
                through = date_cls(year, month, 1)
            except (TypeError, ValueError):
                raise BusinessRuleError("month must be YYYY-MM.")
            through = (through.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        else:
            through = depreciation.last_completed_month_end()
        dry_run = str(request.data.get("dry_run", "")).lower() in ("1", "true", "yes")
        return Response(depreciation.run_company(self.company, request.user, through, dry_run=dry_run))

    @action(detail=True, methods=["post"])
    def dispose(self, request, pk=None):
        """BB-000459: dispose with optional proceeds; NBV never hits Depreciation (5300)."""
        from django.db import transaction

        confirm = request.data.get("confirm")
        if confirm not in (True, "true", "True", 1, "1"):
            raise BusinessRuleError("confirm must be true to dispose a fixed asset.")
        with transaction.atomic():
            asset = FixedAsset.objects.select_for_update().get(pk=self.get_object().pk, company=self.company)
            if asset.status != FixedAsset.Status.ACTIVE:
                raise BusinessRuleError("Asset has already been disposed.")
            if asset.depreciated_amount < 0 or asset.depreciated_amount > asset.acquisition_cost:
                raise BusinessRuleError("Accumulated depreciation is invalid for disposal.")
            try:
                proceeds = Decimal(str(request.data.get("proceeds", "0") or "0"))
            except Exception as exc:
                raise BusinessRuleError("proceeds must be a number.") from exc
            if proceeds < 0:
                raise BusinessRuleError("proceeds cannot be negative.")
            proceeds = proceeds.quantize(Decimal("0.01"))

            # ACS-12: the disposal date is the user's, not "today". It cannot
            # precede the purchase or fall in the future; period locks are
            # enforced by PostingService on that date.
            raw_date = request.data.get("disposal_date")
            if raw_date:
                from datetime import date as date_cls

                try:
                    disposal_date = date_cls.fromisoformat(str(raw_date)[:10])
                except ValueError as exc:
                    raise BusinessRuleError("disposal_date must be a valid date (YYYY-MM-DD).") from exc
            else:
                disposal_date = timezone.localdate()
            if disposal_date > timezone.localdate():
                raise BusinessRuleError("disposal_date cannot be in the future.")
            if asset.acquisition_date and disposal_date < asset.acquisition_date:
                raise BusinessRuleError("disposal_date cannot be before the acquisition date.")

            # ACS-12 follow-up: the part-month up to the sale is charged first, so book value is right.
            from . import depreciation

            if str(request.data.get("charge_part_month", "")).lower() in ("1", "true", "yes"):
                before = asset.depreciated_amount
                depreciation.post_disposal_month(asset, disposal_date, request.user)
                if asset.depreciated_amount != before:
                    asset.save(update_fields=["depreciated_amount", "updated_at"])
            accounts = seed_chart_of_accounts(self.company, request.user)
            net_book_value = asset.acquisition_cost - asset.depreciated_amount
            lines = []
            if asset.depreciated_amount > 0:
                lines.append({"account": asset.accumulated_depreciation_account, "debit": asset.depreciated_amount})
            if proceeds > 0:
                bank_account_id = request.data.get("bank_account")
                if bank_account_id not in (None, ""):
                    from payments.models import BankAccount

                    bank_account = BankAccount.objects.filter(company=self.company, pk=bank_account_id).first()
                    if bank_account is None:
                        raise BusinessRuleError("Unknown bank account.")
                    cash_account = PostingService._bank_gl_account(self.company, bank_account, disposal_date)
                else:
                    # Default cash; callers may pass cash_account_code=1500 for the bank aggregate.
                    cash_code = str(request.data.get("cash_account_code") or "1100")
                    if cash_code not in ("1100", "1500"):
                        raise BusinessRuleError("cash_account_code must be 1100 (Cash) or 1500 (Bank).")
                    cash_account = accounts[cash_code]
                lines.append({"account": cash_account, "debit": proceeds})
            loss = net_book_value - proceeds
            if loss > 0:
                lines.append({"account": accounts["5600"], "debit": loss})
            elif loss < 0:
                lines.append({"account": accounts["5700"], "credit": abs(loss)})
            lines.append({"account": asset.asset_account, "credit": asset.acquisition_cost})
            PostingService.post(
                company=self.company,
                source_type="FIXED_ASSET",
                source_id=asset.id,
                purpose="DISPOSAL",
                entry_date=disposal_date,
                user=request.user,
                narration=f"Disposal: {asset.name}",
                lines=lines,
            )
            asset.status, asset.disposed_at = FixedAsset.Status.DISPOSED, disposal_date
            asset.save(update_fields=["status", "disposed_at", "updated_at"])
        return Response(self.get_serializer(asset).data)


# Above this, the owner screen refuses to post and an operator runs
# ``backfill_accounting_postings``. Dry-run still works so the preview is visible.
OWNER_BACKFILL_INVOICE_LIMIT = 2000


def _books_switch_allowed(company) -> bool:
    """Settings may turn books on only when there is nothing to back-fill,
    or the posted journals already tie. History goes through the back-fill job.
    """
    from inventory.models import StockMovement
    from purchases.models import PurchaseInvoice
    from sales.models import SalesInvoice

    has_history = (
        SalesInvoice.objects.filter(company=company).exists()
        or PurchaseInvoice.objects.filter(company=company).exists()
        or StockMovement.objects.filter(company=company).exists()
    )
    if not has_history:
        return True
    # Ask the hypothetical question ("would the books tie if they were on?")
    # without touching the shared cache, which holds the real answer.
    was = company.accounting_enabled
    company.accounting_enabled = True
    try:
        return BooksHealthService._gl_basis_ready_uncached(company)
    finally:
        company.accounting_enabled = was


def accounting_backfill_needed(company) -> bool:
    """True when books are on but back-fill is not complete.

    Complete means a tying trial balance, no DOCS_GL_* alerts, posted
    documents, and inventory GL within tolerance of valuation.
    """
    if not getattr(company, "accounting_enabled", False):
        return False
    from accounting.services import BooksHealthService

    return not BooksHealthService.gl_basis_ready(company)


def _settings_payload(company) -> dict:
    return AccountingSettingsSerializer(
        {
            "accounting_enabled": company.accounting_enabled,
            "accounting_backfill_needed": accounting_backfill_needed(company),
            "expense_gl_start_date": company.expense_gl_start_date,
            "expense_gl_pending": pending_summary(company),
        }
    ).data


class AccountingSettingsView(APIView):
    serializer_class = AccountingSettingsSerializer

    def get_permissions(self):
        if self.request.method == "GET":
            return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]
        return [IsAuthenticated(), HasCompany(), IsOwner()]

    def get(self, request):
        return Response(_settings_payload(get_company_user(request).company))

    def post(self, request):
        company = get_company_user(request).company
        has_toggle = "accounting_enabled" in request.data
        has_expense_date = "expense_gl_start_date" in request.data
        if not has_toggle and not has_expense_date:
            raise BusinessRuleError("accounting_enabled is required.")
        if has_expense_date:
            self._set_expense_gl_start(company, request.data.get("expense_gl_start_date"))
        if has_toggle:
            raw = request.data.get("accounting_enabled")
            if isinstance(raw, str):
                enabled = raw.strip().lower() in {"1", "true", "yes", "on"}
            else:
                enabled = bool(raw)
            if enabled and not company.accounting_enabled and not _books_switch_allowed(company):
                raise BusinessRuleError(
                    "Books stay off until back-fill produces a trial balance that ties. "
                    "Preview it on this screen, then post. A company with no invoices "
                    "or stock can turn books on directly."
                )
            company.accounting_enabled = enabled
            company.save(update_fields=["accounting_enabled", "updated_at"])
            BooksHealthService.clear_gl_basis_cache(company)
            if enabled:
                seed_chart_of_accounts(company, request.user)
        return Response(_settings_payload(company))

    @staticmethod
    def _set_expense_gl_start(company, raw):
        """ACS-15: the day expenses start posting to the ledger; empty switches it off."""
        from datetime import date as date_cls

        if raw in (None, ""):
            new_value = None
        else:
            try:
                new_value = date_cls.fromisoformat(str(raw)[:10])
            except ValueError as exc:
                raise BusinessRuleError("expense_gl_start_date must be a valid date (YYYY-MM-DD).") from exc
            if not company.accounting_enabled:
                raise BusinessRuleError("Turn accounting on before posting expenses to the books.")
            cutover = company.books_start_date
            if cutover and new_value < cutover:
                raise BusinessRuleError(
                    f"Expenses cannot post before the books start date ({cutover})."
                )
        company.expense_gl_start_date = new_value
        company.save(update_fields=["expense_gl_start_date", "updated_at"])


class AccountingBackfillView(APIView):
    """Owner dry-run, then an idempotent post of missing journals.

    A closed period blocks the post. Companies above OWNER_BACKFILL_INVOICE_LIMIT
    are told to use the operator command.
    """

    permission_classes = [IsAuthenticated, HasCompany, IsOwner]

    def get(self, request):
        """Status of the last owner back-fill (running / done / failed / idle)."""
        from django.core.cache import cache

        from accounting.tasks import backfill_cache_key

        company = get_company_user(request).company
        return Response(cache.get(backfill_cache_key(company.id)) or {"status": "idle"})

    def post(self, request):
        from core.idempotency import wrap_idempotent
        from core.services.audit import AuditService

        company = get_company_user(request).company

        def _run():
            return self._execute(request, company, AuditService)

        return wrap_idempotent(
            request=request, company=company, scope="accounting_backfill", build=_run,
        )

    def _execute(self, request, company, AuditService):
        from purchases.models import PurchaseInvoice
        from sales.models import SalesInvoice

        dry_run = request.data.get("dry_run") in (True, "true", "True", 1, "1")
        confirm = request.data.get("confirm") in (True, "true", "True", 1, "1")
        if dry_run == confirm:
            raise BusinessRuleError("Send dry_run or confirm, not both and not neither.")
        sales_count = SalesInvoice.objects.filter(company=company).count()
        purchase_count = PurchaseInvoice.objects.filter(company=company).count()
        invoice_count = sales_count + purchase_count
        closed = AccountingPeriod.objects.filter(
            company=company, status=AccountingPeriod.Status.CLOSED,
        ).exists()
        if not dry_run and closed:
            raise BusinessRuleError(
                "Backfill is blocked because an accounting period is already closed."
            )
        if not dry_run and invoice_count > OWNER_BACKFILL_INVOICE_LIMIT:
            raise BusinessRuleError(
                "This company has more than "
                f"{OWNER_BACKFILL_INVOICE_LIMIT} invoices. An operator runs "
                "backfill_accounting_postings."
            )
        if dry_run:
            return Response(
                perform_accounting_backfill(
                    company, request.user, dry_run=True, sales_count=sales_count,
                    purchase_count=purchase_count, closed=closed,
                )
            )
        from django.core.cache import cache

        from accounting.tasks import backfill_cache_key, run_owner_accounting_backfill

        key = backfill_cache_key(company.id)
        # cache.add is atomic: a double click or a second tab cannot start a second
        # concurrent run. A stale "running" marker expires on its own (10 minutes).
        if not cache.add(key, {"status": "running"}, 600):
            current = cache.get(key) or {}
            if current.get("status") == "running":
                return Response(current, status=status.HTTP_202_ACCEPTED)
            # A finished or failed earlier run is history, not a lock.
            cache.set(key, {"status": "running"}, 600)
        run_owner_accounting_backfill.delay(company.id, request.user.id)
        cached = cache.get(key) or {"status": "running"}
        if cached.get("status") == "done":
            return Response(cached)
        if cached.get("status") == "failed":
            raise BusinessRuleError(cached.get("error") or "Back-fill failed. Entries already posted stay; run it again to finish.")
        return Response(cached, status=status.HTTP_202_ACCEPTED)


def _backfill_ready(company, skipped_by_period) -> bool:
    if skipped_by_period:
        return False
    BooksHealthService.clear_gl_basis_cache(company)
    try:
        return bool(BooksHealthService._gl_basis_ready_uncached(company))
    except Exception:  # noqa: BLE001 — a health failure must not leave books on
        return False


def perform_accounting_backfill(
    company, user, *, dry_run: bool, sales_count=None, purchase_count=None, closed=False,
) -> dict:
    """Post missing journals. Idempotent: a document that already has a journal is skipped."""
    from accounting.management.commands.backfill_accounting_postings import Command
    from accounting.reports import balance_sheet, trial_balance
    from core.services.audit import AuditService
    from purchases.models import PurchaseInvoice
    from sales.models import SalesInvoice

    if sales_count is None:
        sales_count = SalesInvoice.objects.filter(company=company).count()
    if purchase_count is None:
        purchase_count = PurchaseInvoice.objects.filter(company=company).count()
    enabled_for_posting = False
    if not dry_run and not company.accounting_enabled:
        # Posting reads the saved flag. Keep books on only when the pass is clean.
        company.accounting_enabled = True
        company.save(update_fields=["accounting_enabled", "updated_at"])
        enabled_for_posting = True
    if company.accounting_enabled:
        seed_chart_of_accounts(company, user)
    command = Command()
    command.stdout = _NullStream()
    command.stderr = _NullStream()
    posted, skipped, would = command._backfill_company(company, dry_run=dry_run)
    skipped_by_period = getattr(command, "skipped_by_period", {}) or {}
    if enabled_for_posting and (skipped_by_period or not _backfill_ready(company, skipped_by_period)):
        company.accounting_enabled = False
        company.save(update_fields=["accounting_enabled", "updated_at"])
        BooksHealthService.clear_gl_basis_cache(company)
    if not dry_run:
        command._retag_party_lines(company)
    BooksHealthService.clear_gl_basis_cache(company)
    tb = trial_balance(company)
    sheet = balance_sheet(company)
    health = BooksHealthService.control_balances(company)
    docs_gl = [
        a for a in health.get("alerts") or []
        if str(a.get("code") or "").startswith("DOCS_GL_")
    ]
    payload = {
        "dry_run": dry_run,
        "posted": posted,
        "skipped": skipped,
        "skipped_by_period": skipped_by_period,
        "would_post": would,
        "sales_invoices": sales_count,
        "purchase_invoices": purchase_count,
        "trial_balance_balanced": bool(tb.get("balanced")),
        "total_debit": str(tb.get("total_debit") or 0),
        "total_credit": str(tb.get("total_credit") or 0),
        "inventory_variance": str(sheet.get("inventory_variance") or 0),
        "docs_gl_alerts": docs_gl,
        "accounting_enabled": company.accounting_enabled,
        "accounting_backfill_needed": (
            accounting_backfill_needed(company) if company.accounting_enabled else would > 0
        ),
        "closed_period": closed,
    }
    if not dry_run:
        AuditService.log(
            company=company,
            user=user,
            action="UPDATE",
            entity_type="Company",
            entity_id=company.id,
            description="Accounting backfill posted from the owner screen",
            metadata={
                "posted": posted,
                "skipped": skipped,
                "trial_balance_balanced": payload["trial_balance_balanced"],
                "inventory_variance": payload["inventory_variance"],
            },
        )
    return payload


class _NullStream:
    def write(self, *args, **kwargs):
        return None


class FinancialYearCloseView(AccountingEnabledMixin, APIView):
    """BB-000664: Owner-only FY close — POST {fyEnd, confirm: true}. GET previews; DELETE-style reverse via POST {reverse: true}."""

    permission_classes = [IsAuthenticated, HasCompany, IsOwner]

    @property
    def company(self):
        return get_company_user(self.request).company

    @staticmethod
    def _parse_end(raw):
        from datetime import date as date_cls

        if not raw:
            raise BusinessRuleError("fyEnd is required (YYYY-MM-DD).")
        try:
            return date_cls.fromisoformat(str(raw)[:10])
        except ValueError as exc:
            raise BusinessRuleError("fyEnd must be YYYY-MM-DD.") from exc

    def get(self, request):
        """ACS-34: preview what closing the year would do."""
        from . import period_ops

        fy_end = self._parse_end(request.query_params.get("fy_end") or request.query_params.get("fyEnd"))
        return Response(period_ops.fy_close_preview(self.company, fy_end))

    def post(self, request):
        from datetime import date as date_cls

        if not self.company.accounting_enabled:
            raise BusinessRuleError("Accounting is not enabled for this company.")
        data = request.data or {}
        raw_end = data.get("fyEnd") or data.get("fy_end")
        if data.get("reverse") in (True, "true", "True", 1, "1"):
            from . import period_ops

            result = period_ops.reverse_fy_close(
                self.company, self._parse_end(raw_end), request.user, data.get("reason", ""),
            )
            return Response({"ok": True, **result})
        confirm = data.get("confirm")
        if confirm not in (True, "true", "True", 1, "1"):
            raise BusinessRuleError("confirm must be true to close the financial year.")
        if not raw_end:
            raise BusinessRuleError("fyEnd is required (YYYY-MM-DD).")
        try:
            fy_end = date_cls.fromisoformat(str(raw_end)[:10])
        except ValueError as exc:
            raise BusinessRuleError("fyEnd must be YYYY-MM-DD.") from exc
        entry = close_financial_year(self.company, fy_end, user=request.user)
        payload = {
            "ok": True,
            "fy_end": fy_end.isoformat(),
            "journal_id": entry.id if entry is not None else None,
        }
        return Response(payload)


class AccountingReportView(AccountingEnabledMixin, APIView):
    permission_classes = [IsAuthenticated, HasCompany, CanViewFinancialReports]

    @property
    def company(self):
        return get_company_user(self.request).company

    def get_permissions(self):
        # BB-000293: XLSX export requires CanExport.
        export_format = (self.request.query_params.get("format") or "json").lower()
        if export_format == "xlsx":
            return [IsAuthenticated(), HasCompany(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request, report):
        # B1-010: parse/validate params in the view — a raw "?as_of=abc" or
        # "?cost_center=x" otherwise flows straight into an ORM date/int lookup
        # and 500s.
        from datetime import date as _date

        def _qp_date(name):
            raw = (request.query_params.get(name) or "").strip()
            if not raw:
                return None
            try:
                return _date.fromisoformat(raw[:10])
            except ValueError as exc:
                raise BusinessRuleError(f"{name} must be YYYY-MM-DD.") from exc

        def _qp_cost_center():
            raw = (request.query_params.get("cost_center") or "").strip()
            if not raw:
                return None
            try:
                return int(raw)
            except (TypeError, ValueError) as exc:
                raise BusinessRuleError("cost_center must be an integer id.") from exc

        as_of = _qp_date("as_of")
        date_from, date_to = _qp_date("from"), _qp_date("to")
        cost_center = _qp_cost_center()
        if report == "trial-balance":
            return self._report_response(report, trial_balance(self.company, as_of), request)
        if report == "profit-and-loss":
            return self._report_response(report, profit_and_loss(self.company, date_from, date_to, cost_center), request)
        if report == "balance-sheet":
            return self._report_response(
                report, balance_sheet(self.company, as_of, cost_center), request,
            )
        if report == "cash-flow":
            return self._report_response(
                report, cash_flow(self.company, date_from, date_to, cost_center), request,
            )
        if report == "books-health":
            # B1-016/B1-026: control_balances() is a heavy fixed-cost call
            # (AR/AP nets, ~8 `_has_missing` .exists() subqueries, depreciation
            # alerts, advance-recon) with no memoisation — a dashboard polling
            # this endpoint hammers the DB every call. Short-cache the GET
            # response only; period_close_blockers (the safety-critical
            # close-gate caller) still calls control_balances() fresh every
            # time so a real AR/AP mismatch is never masked by a stale cache.
            from django.core.cache import cache

            cache_key = f"bizboard:books-health:{self.company.id}"
            cached = cache.get(cache_key)
            if cached is not None:
                return Response(cached)
            payload = BooksHealthService.control_balances(self.company)
            cache.set(cache_key, payload, timeout=20)
            return Response(payload)
        raise BusinessRuleError("Unknown accounting report.")

    def _report_response(self, report, payload, request):
        if request.query_params.get("format") != "xlsx":
            return Response(payload)
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = report.replace("-", " ").title()
        if "rows" not in payload and any(
            k in payload for k in ("operating_activities", "net_cash_flow")
        ):
            # B1-017: cash-flow has no flat "rows" list — flatten the activity
            # buckets so the workbook isn't just a header row.
            sheet.append(["Section", "Metric", "Amount"])
            for section in (
                "operating_activities",
                "investing_activities",
                "financing_activities",
            ):
                bucket = payload.get(section) or {}
                for metric, amount in bucket.items():
                    sheet.append([section.replace("_", " ").title(), metric, amount])
            sheet.append(["", "Net cash flow", payload.get("net_cash_flow", 0)])
        else:
            sheet.append(["Code", "Account", "Type", "Debit", "Credit", "Balance"])
            rows = payload.get("rows", [])
            if isinstance(rows, dict):
                rows = [row for group in rows.values() for row in group]
            for row in rows:
                sheet.append([
                    row.get("account_code", ""), row.get("account_name", ""),
                    row.get("account_type", ""), row.get("debit", 0),
                    row.get("credit", 0), row.get("balance", 0),
                ])
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{report}.xlsx"'
        return response


class CashShiftRegisterViewSet(AccountingEnabledMixin, CompanyScopedViewSet):
    """Daily till: opening float, note count, expected cash, and a day lock."""

    queryset = CashShiftRegister.objects.all()
    serializer_class = CashShiftRegisterSerializer

    def get_permissions(self):
        if getattr(self, "action", None) in (*_MUTATE_ACTIONS, "close_shift", "drop", "today", "history", "summary"):
            return [IsAuthenticated(), HasCompany(), CanCreatePayments()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def perform_create(self, serializer):
        from decimal import ROUND_HALF_UP

        business_date = serializer.validated_data.get("business_date") or timezone.localdate()
        opening = Decimal(str(serializer.validated_data.get("opening_float") or 0)).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP,
        )
        terminal_id = str(serializer.validated_data.get("terminal_id") or f"device-{self.request.user.pk}")[:64]
        terminal_label = str(serializer.validated_data.get("terminal_label") or "")[:64]
        if CashShiftRegister.objects.filter(
            company=self.company, terminal_id=terminal_id, status=CashShiftRegister.Status.OPEN,
        ).exists():
            raise BusinessRuleError("This terminal already has an open shift.", code="pos_shift_open")
        serializer.save(
            company=self.company,
            cashier=self.request.user,
            business_date=business_date,
            opening_float=opening,
            terminal_id=terminal_id,
            terminal_label=terminal_label,
            opened_at=timezone.now(),
            created_by=self.request.user,
            updated_by=self.request.user,
        )

    def perform_update(self, serializer):
        if serializer.instance.locked_at or serializer.instance.status == CashShiftRegister.Status.CLOSED:
            raise BusinessRuleError("A locked cash register cannot be edited.")
        serializer.save(updated_by=self.request.user, company=self.company)

    def perform_destroy(self, instance):
        if instance.locked_at or instance.status == CashShiftRegister.Status.CLOSED:
            raise BusinessRuleError("A locked cash register cannot be edited.")
        super().perform_destroy(instance)

    @action(detail=True, methods=["post"], url_path="close")
    def close_shift(self, request, pk=None):
        from core.permissions import get_company_user

        from .cash_shifts import close_cash_shift

        shift = self.get_object()
        member = get_company_user(request)
        # One cashier's till is closed by that cashier, or by a manager. Any other payment-capable
        # user could otherwise lock someone's drawer with counts that hide a real variance.
        if shift.cashier_id != request.user.pk and getattr(member, "role", "") not in ("OWNER", "MANAGER", "ACCOUNTANT"):
            raise BusinessRuleError("Only the cashier who opened this till, or a manager, can close it.")
        shift = close_cash_shift(shift, request.data.get("denominations") or {}, request.user)
        return Response(self.get_serializer(shift).data)

    @action(detail=False, methods=["get"], url_path="today")
    def today(self, request):
        """The caller's open shift for today. Does not require financial-report access."""
        terminal_id = str(request.query_params.get("terminal_id") or "").strip()
        qs = CashShiftRegister.objects.filter(company=self.company)
        if terminal_id:
            qs = qs.filter(terminal_id=terminal_id)
        else:
            qs = qs.filter(cashier=request.user)
        shift = qs.filter(status=CashShiftRegister.Status.OPEN).order_by("-opened_at", "-id").first()
        if shift is None:
            shift = qs.order_by("-id").first()
        if shift is None:
            return Response({"shift": None})
        from .cash_shifts import expected_system_cash

        data = self.get_serializer(shift).data
        if shift.status == CashShiftRegister.Status.OPEN:
            data["expected_cash"] = str(expected_system_cash(
                shift.company, shift.cashier, shift.business_date, shift.opening_float,
                cash_dropped=shift.cash_dropped, shift=shift,
            ))
        return Response({"shift": data})

    @action(detail=False, methods=["get"], url_path="history")
    def history(self, request):
        rows = CashShiftRegister.objects.filter(
            company=self.company, cashier=request.user,
        ).order_by("-opened_at", "-id")[:30]
        return Response(self.get_serializer(rows, many=True).data)

    @action(detail=True, methods=["get"], url_path="summary")
    def summary(self, request, pk=None):
        from django.http import HttpResponse

        from .shift_summary import render_shift_summary

        shift = self.get_object()
        pdf = render_shift_summary(shift)
        return HttpResponse(pdf, content_type="application/pdf")

    @action(detail=True, methods=["post"], url_path="drop")
    def drop(self, request, pk=None):
        from .cash_shifts import drop_cash

        shift = self.get_object()
        if shift.cashier_id != request.user.pk:
            raise BusinessRuleError("Only the cashier who opened this till can drop cash.")
        shift = drop_cash(shift, request.data.get("amount"), request.user, note=str(request.data.get("note") or ""))
        return Response(self.get_serializer(shift).data)


class ExpenseViewSet(CompanyScopedViewSet):
    queryset = Expense.objects.select_related("category")
    serializer_class = ExpenseSerializer
    audit_entity = "Expense"

    def get_permissions(self):
        action = getattr(self, "action", None)
        if action in ("create", "update", "partial_update", "destroy", "void"):
            return [IsAuthenticated(), HasCompany(), SubscriptionWritesAllowed(), CanCreatePayments()]
        if action == "duplicates":
            return [IsAuthenticated(), HasCompany(), CanCreatePayments()]
        if action == "export":
            return [IsAuthenticated(), HasCompany(), CanViewFinancialReports(), CanExport()]
        if action == "post_pending":
            return [IsAuthenticated(), HasCompany(), SubscriptionWritesAllowed(), CanPostJournals()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get_queryset(self):
        qs = super().get_queryset()
        params = self.request.query_params
        # ACS-14: voided expenses are hidden from the register unless asked for.
        if self.action in ("list", "export") and params.get("include_void") not in ("1", "true", "True"):
            qs = qs.filter(status=Expense.Status.ACTIVE)
        if params.get("category"):
            qs = qs.filter(category_id=params["category"])
        if params.get("date_from"):
            qs = qs.filter(expense_date__gte=params["date_from"])
        if params.get("date_to"):
            qs = qs.filter(expense_date__lte=params["date_to"])
        if params.get("payment_mode"):
            qs = qs.filter(payment_mode=params["payment_mode"].upper())
        if params.get("q"):
            from django.db.models import Q

            term = params["q"].strip()
            qs = qs.filter(Q(party_name__icontains=term) | Q(number__icontains=term) | Q(notes__icontains=term))
        return qs

    def list(self, request, *args, **kwargs):
        """ACS-17: the register also reports the total of everything the filters match, not just this page."""
        response = super().list(request, *args, **kwargs)
        total = self.filter_queryset(self.get_queryset()).aggregate(t=Sum("amount"))["t"] or Decimal("0")
        if isinstance(response.data, dict):
            response.data["filtered_total"] = str(total)
        return response

    @action(detail=False, methods=["get"])
    def export(self, request):
        """ACS-17: the filtered register as CSV or XLSX."""
        qs = self.filter_queryset(self.get_queryset()).select_related("category", "cost_center").order_by("expense_date", "id")[:20000]
        header = ["Number", "Date", "Category", "Party", "Amount", "GST", "TDS", "Mode", "Status", "Notes"]
        rows = [
            [e.number, e.expense_date.isoformat(), e.category.name, e.party_name, str(e.amount),
             str((e.cgst_amount or 0) + (e.sgst_amount or 0) + (e.igst_amount or 0)), str(e.tds_amount or 0),
             e.payment_mode, e.status, e.notes]
            for e in qs
        ]
        if request.query_params.get("file") == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Expenses"
            ws.append(header)
            for r in rows:
                ws.append(r)
            out = BytesIO()
            wb.save(out)
            resp = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            resp["Content-Disposition"] = 'attachment; filename="expenses.xlsx"'
            return resp
        import csv

        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = 'attachment; filename="expenses.csv"'
        writer = csv.writer(resp)
        writer.writerow(header)
        for r in rows:
            writer.writerow([("'" + c) if isinstance(c, str) and c[:1] in ("=", "+", "-", "@") else c for c in r])
        return resp

    def _assert_period_open(self, on_date):
        # ACS-14: an expense is a money record; closed or GST-locked periods are off limits.
        from reporting.gst_periods import assert_period_allows_money_amend

        assert_period_allows_money_amend(self.company, on_date)

    def perform_create(self, serializer):
        from core.services.document_numbers import DocumentNumberService, resolve_series_gstin

        self._assert_period_open(serializer.validated_data.get("expense_date") or timezone.localdate())
        with transaction.atomic():
            instance = serializer.save(
                company=self.company,
                created_by=self.request.user,
                updated_by=self.request.user,
            )
            if not instance.number:
                instance.number = DocumentNumberService.next_number(
                    self.company, "EXPENSE", gstin=resolve_series_gstin(self.company),
                    on_date=instance.expense_date,
                )
                instance.save(update_fields=["number"])
            sync_expense_posting(instance, self.request.user)
            self._audit("CREATE", instance)

    def perform_update(self, serializer):
        instance = serializer.instance
        if instance.status == Expense.Status.VOID:
            raise BusinessRuleError("A voided expense cannot be edited.")
        # Both the old and the new date must be open: moving an expense out of a
        # closed month changes that month just as much as editing it there.
        self._assert_period_open(instance.expense_date)
        new_date = serializer.validated_data.get("expense_date")
        if new_date and new_date != instance.expense_date:
            self._assert_period_open(new_date)
        with transaction.atomic():
            super().perform_update(serializer)
            sync_expense_posting(serializer.instance, self.request.user)

    def perform_destroy(self, instance):
        # DELETE keeps working for old clients but is now a void with a reason.
        self._void(instance, self.request.data.get("reason") or self.request.query_params.get("reason"))

    def _void(self, instance, reason):
        from core.services.audit import AuditService

        reason = (reason or "").strip()
        if not reason:
            raise BusinessRuleError("A reason is required to void an expense.", code="void_reason_required")
        if len(reason) > 300:
            raise BusinessRuleError("The reason must be 300 characters or fewer.")
        if instance.status == Expense.Status.VOID:
            raise BusinessRuleError("This expense is already voided.")
        self._assert_period_open(instance.expense_date)
        instance.status = Expense.Status.VOID
        instance.voided_at = timezone.now()
        instance.voided_by = self.request.user
        instance.void_reason = reason
        instance.updated_by = self.request.user
        instance.save(update_fields=["status", "voided_at", "voided_by", "void_reason", "updated_by", "updated_at"])
        sync_expense_posting(instance, self.request.user)
        AuditService.log(
            company=self.company, user=self.request.user, action="UPDATE", entity_type="Expense",
            entity_id=str(instance.pk), description=f"Expense {instance.number} voided",
            metadata={"event": "expense.voided", "reason": reason, "amount": str(instance.amount)},
        )

    @action(detail=False, methods=["get"])
    def duplicates(self, request):
        """ACS-16: other active expenses with the same date, amount and party -- a likely double entry."""
        raw_date = request.query_params.get("date")
        amount = request.query_params.get("amount")
        party = (request.query_params.get("party") or "").strip()
        if not (raw_date and amount):
            raise BusinessRuleError("date and amount are required.")
        qs = Expense.objects.filter(
            company=self.company, status=Expense.Status.ACTIVE,
            expense_date=raw_date, amount=amount,
        )
        if party:
            qs = qs.filter(party_name__iexact=party)
        exclude = request.query_params.get("exclude")
        if exclude:
            qs = qs.exclude(pk=exclude)
        rows = [{"id": e.id, "number": e.number, "party_name": e.party_name} for e in qs[:5]]
        return Response({"count": len(rows), "results": rows})

    @action(detail=False, methods=["get", "post"], url_path="post-pending")
    def post_pending(self, request):
        """ACS-15: GET previews, POST posts, every active expense on/after the start date not yet in the books."""
        if request.method == "GET":
            summary = pending_summary(self.company)
            return Response({"count": summary["count"], "total": summary["total"]})
        dry_run = str(request.data.get("dry_run", "")).lower() in ("1", "true", "yes")
        return Response(post_pending_expenses(self.company, request.user, dry_run=dry_run))

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def void(self, request, pk=None):
        instance = self.get_object()
        # Lock the row and re-read it: two voids at once must not both see an active expense.
        instance = Expense.objects.select_for_update().get(pk=instance.pk)
        self._void(instance, request.data.get("reason"))
        return Response(self.get_serializer(instance).data)


class RecurringExpenseViewSet(CompanyScopedViewSet):
    """ACS-17: templates for expenses that repeat, and the one-click "create what is due"."""

    queryset = RecurringExpense.objects.select_related("category")
    serializer_class = RecurringExpenseSerializer
    audit_entity = "RecurringExpense"

    def get_permissions(self):
        if getattr(self, "action", None) in ("create", "update", "partial_update", "destroy", "create_due"):
            return [IsAuthenticated(), HasCompany(), SubscriptionWritesAllowed(), CanCreatePayments()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    @action(detail=False, methods=["get"])
    def due(self, request):
        from . import recurring

        rows = recurring.due_templates(self.company)
        return Response({"count": rows.count(), "results": RecurringExpenseSerializer(rows, many=True, context=self.get_serializer_context()).data})

    @action(detail=False, methods=["post"], url_path="create-due")
    def create_due(self, request):
        from . import recurring

        return Response(recurring.create_due(self.company, request.user, ids=request.data.get("ids")))


class ExpenseCategoryAccountViewSet(CompanyScopedViewSet):
    """ACS-15: which expense ledger each category posts to. One mapping per category."""

    queryset = ExpenseCategoryAccount.objects.select_related("account")
    serializer_class = ExpenseCategoryAccountSerializer
    http_method_names = ["get", "post", "delete", "head", "options"]

    def get_permissions(self):
        if getattr(self, "action", None) in ("create", "destroy"):
            return [IsAuthenticated(), HasCompany(), SubscriptionWritesAllowed(), IsOwner()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        mapping, _created = ExpenseCategoryAccount.objects.update_or_create(
            company=self.company,
            category=serializer.validated_data["category"],
            defaults={
                "account": serializer.validated_data["account"],
                "updated_by": request.user,
            },
        )
        if _created:
            mapping.created_by = request.user
            mapping.save(update_fields=["created_by"])
        self._audit("UPDATE", mapping)
        return Response(self.get_serializer(mapping).data, status=status.HTTP_200_OK)
