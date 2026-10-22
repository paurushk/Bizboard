"""WP-F10 / DIM6-19: compare Bizboard's trial balance with the old system's for a parallel month.

The pilot exports its trial balance from Tally or Busy (CSV, or the first sheet of an XLSX) with a
ledger-name column and either Debit/Credit columns or one signed Balance column. A mapping CSV
(``their_ledger,account_code``) maps their ledgers onto our chart; several of theirs may map to one
of ours. The command prints every account where the two disagree by more than the tolerance, and
the ledgers it could not map.

    python manage.py compare_trial_balance --company 42 --as-of 2026-11-30 \\
        --theirs tally_tb.csv --mapping tally_map.csv

Exit code 1 when anything differs, so it can gate the month-end sign-off.
"""

from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

NAME_HEADERS = ("ledger", "particulars", "account", "name", "ledger name", "account name")
DEBIT_HEADERS = ("debit", "dr", "debit amount", "closing debit")
CREDIT_HEADERS = ("credit", "cr", "credit amount", "closing credit")
BALANCE_HEADERS = ("balance", "closing balance", "net")


def _num(raw) -> Decimal:
    text = str(raw if raw is not None else "").replace(",", "").replace("₹", "").strip()
    negative = text.endswith(("Cr", "cr")) or (text.startswith("(") and text.endswith(")"))
    for suffix in ("Dr", "dr", "Cr", "cr"):
        if text.endswith(suffix):
            text = text[: -len(suffix)].strip()
    text = text.removeprefix("(").removesuffix(")").strip()
    if not text:
        return Decimal("0")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise CommandError(f"Not a number: {raw!r}") from exc
    return -value if negative else value


def _rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        values = [list(r) for r in sheet.iter_rows(values_only=True)]
        header = [str(h or "").strip().lower() for h in values[0]]
        return [dict(zip(header, r)) for r in values[1:]]
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [{(k or "").strip().lower(): v for k, v in row.items()} for row in csv.DictReader(fh)]


def _pick(row: dict, names) -> str | None:
    return next((k for k in names if k in row), None)


def read_theirs(path: Path) -> dict[str, Decimal]:
    """{their ledger name: net debit (debit - credit)}."""
    rows = _rows(path)
    if not rows:
        raise CommandError(f"{path} has no rows")
    name_key = _pick(rows[0], NAME_HEADERS)
    if name_key is None:
        raise CommandError(f"{path}: no ledger name column (expected one of {NAME_HEADERS})")
    d_key, c_key, b_key = _pick(rows[0], DEBIT_HEADERS), _pick(rows[0], CREDIT_HEADERS), _pick(rows[0], BALANCE_HEADERS)
    out: dict[str, Decimal] = defaultdict(Decimal)
    for row in rows:
        name = str(row.get(name_key) or "").strip()
        if not name or name.lower() in ("total", "grand total"):
            continue
        if d_key or c_key:
            out[name] += _num(row.get(d_key)) - _num(row.get(c_key))
        elif b_key:
            out[name] += _num(row.get(b_key))
        else:
            raise CommandError(f"{path}: needs Debit/Credit columns or a Balance column")
    return dict(out)


def read_mapping(path: Path) -> dict[str, str]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.reader(fh)
        rows = [r for r in reader if r and r[0].strip()]
    if rows and rows[0][0].strip().lower() in ("their_ledger", "ledger", "their ledger"):
        rows = rows[1:]
    return {r[0].strip(): r[1].strip() for r in rows if len(r) >= 2 and r[1].strip()}


def compare(ours: dict[str, Decimal], theirs: dict[str, Decimal], mapping: dict[str, str], tolerance: Decimal):
    mapped: dict[str, Decimal] = defaultdict(Decimal)
    unmapped = []
    for ledger, amount in theirs.items():
        code = mapping.get(ledger)
        if code is None:
            if amount != 0:
                unmapped.append((ledger, amount))
            continue
        mapped[code] += amount
    diffs = []
    for code in sorted(set(ours) | set(mapped)):
        a, b = ours.get(code, Decimal("0")), mapped.get(code, Decimal("0"))
        if abs(a - b) > tolerance:
            diffs.append((code, a, b, a - b))
    return diffs, unmapped


class Command(BaseCommand):
    help = "Compare this company's trial balance with the old system's export (parallel month check)."

    def add_arguments(self, parser):
        parser.add_argument("--company", type=int, required=True)
        parser.add_argument("--as-of", required=True, help="YYYY-MM-DD, the month end both systems closed.")
        parser.add_argument("--theirs", required=True, help="Their trial balance, CSV or XLSX.")
        parser.add_argument("--mapping", required=True, help="CSV: their_ledger,account_code")
        parser.add_argument("--tolerance", default="1.00", help="Allowed difference per account (rupees).")

    def handle(self, *args, **opts):
        from accounting.reports import trial_balance
        from accounts.models import Company
        from core.rls import rls_bypass

        with rls_bypass():
            company = Company.objects.filter(pk=opts["company"]).first()
            if company is None:
                raise CommandError(f"Company {opts['company']} not found.")
            tb = trial_balance(company, date.fromisoformat(opts["as_of"]))
        ours: dict[str, Decimal] = defaultdict(Decimal)
        names = {}
        for row in tb["rows"]:
            ours[row["account_code"]] += Decimal(row["debit"]) - Decimal(row["credit"])
            names[row["account_code"]] = row["account_name"]
        theirs = read_theirs(Path(opts["theirs"]))
        mapping = read_mapping(Path(opts["mapping"]))
        diffs, unmapped = compare(dict(ours), theirs, mapping, Decimal(opts["tolerance"]))
        for code, a, b, d in diffs:
            self.stdout.write(f"DIFF  {code} {names.get(code, '?')}: ours {a} theirs {b} difference {d}")
        for ledger, amount in unmapped:
            self.stdout.write(f"UNMAPPED  {ledger}: {amount}")
        self.stdout.write(f"accounts compared={len(set(ours) | set(mapping.values()))} differences={len(diffs)} unmapped={len(unmapped)}")
        if diffs or unmapped:
            raise CommandError("Trial balances do not agree; explain or fix each line before signing the month.")
