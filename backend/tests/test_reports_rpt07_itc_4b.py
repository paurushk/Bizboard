"""RPT-07 / D8: GSTR-3B shows numeric 4(B) reversed and ineligible ITC, per tax head."""

from decimal import Decimal

import pytest

from purchases.models import PurchaseInvoice
from tests.conftest import create_draft_purchase, make_product, make_supplier
from tests.test_gst_returns import PERIOD

pytestmark = pytest.mark.django_db


def _purchase(tenant, supplier, product, status, day):
    pur = create_draft_purchase(
        tenant, supplier,
        [{"product": product.id, "quantity": "1", "unit_price": "1000", "gst_rate": "18"}],
    )
    assert tenant.client.post(f"/api/v1/purchases/invoices/{pur['id']}/complete/").status_code == 200
    inv = PurchaseInvoice.objects.get(pk=pur["id"])
    inv.invoice_date = f"{PERIOD}-{day}"
    inv.itc_eligibility = status
    inv.save(update_fields=["invoice_date", "itc_eligibility"])
    return inv


def _gstr3b(tenant):
    tenant.company.gstin = "29ABCDE1234F1ZW"
    tenant.company.state = "Karnataka"
    tenant.company.save(update_fields=["gstin", "state"])
    resp = tenant.client.get("/api/v1/reports/gstr3b/", {"period": PERIOD})
    assert resp.status_code == 200, resp.data
    return resp.data


def test_reversed_and_ineligible_itc_are_summed_per_head_and_left_out_of_claimable(tenant_a):
    supplier = make_supplier(tenant_a.company, name="S", gstin="29AABCU9603R1ZJ")
    product = make_product(tenant_a.company, sku="ITC4B", hsn_code="8517")
    _purchase(tenant_a, supplier, product, PurchaseInvoice.ItcEligibility.CLAIMABLE, "05")
    rev = _purchase(tenant_a, supplier, product, PurchaseInvoice.ItcEligibility.REVERSED, "06")
    inel = _purchase(tenant_a, supplier, product, PurchaseInvoice.ItcEligibility.INELIGIBLE, "07")

    itc = _gstr3b(tenant_a)["itc"]
    block = itc["reversed_ineligible"]
    assert block["needs_ca_confirmation"] is True
    assert block["reversed"]["count"] == 1
    assert block["ineligible"]["count"] == 1
    assert Decimal(block["reversed"]["cgst"]) == rev.cgst_total
    assert Decimal(block["ineligible"]["sgst"]) == inel.sgst_total
    assert Decimal(block["total_tax"]) == (
        rev.cgst_total + rev.sgst_total + rev.igst_total + inel.cgst_total + inel.sgst_total + inel.igst_total
    )
    # only the CLAIMABLE bill counts towards books ITC
    one_bill_cgst = rev.cgst_total
    assert Decimal(itc["books_itc"]["cgst"]) == one_bill_cgst


def test_block_is_zero_when_nothing_is_flagged(tenant_a):
    block = _gstr3b(tenant_a)["itc"]["reversed_ineligible"]
    assert block["reversed"]["count"] == 0 and block["ineligible"]["count"] == 0
    assert Decimal(block["total_tax"]) == 0


def test_gstr3b_excel_carries_the_payable_and_4b_sheets(tenant_a):
    from io import BytesIO

    from openpyxl import load_workbook

    tenant_a.company.gstin = "29ABCDE1234F1ZW"
    tenant_a.company.state = "Karnataka"
    tenant_a.company.save(update_fields=["gstin", "state"])
    resp = tenant_a.client.get("/api/v1/reports/gstr3b/", {"period": PERIOD, "format": "xlsx"})
    assert resp.status_code == 200
    book = load_workbook(BytesIO(resp.content))
    assert {"Payable", "ITC 4B"} <= set(book.sheetnames)
    labels = [r[0] for r in book["ITC 4B"].iter_rows(values_only=True)]
    assert labels[1:3] == ["reversed", "ineligible"]
