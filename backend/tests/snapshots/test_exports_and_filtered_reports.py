"""Golden snapshots for A16 (export file formats) and §H7 (filtered reports).

- Exports: the sales-register CSV header + normalised rows, and the XLSX
  response's content-type + header row — so a change to `EXPORTS` /
  `EXPORT_FIELDS` / the serialiser shows as a diff.
- Filtered reports: trial balance `as_of=` and P&L `from=/to=` — the existing
  report snapshots are all unfiltered.

Baseline: ``SNAPSHOT_UPDATE=1 pytest tests/snapshots/`` — diff MUST be in the PR.
"""

from __future__ import annotations

import csv
import io
import re
from decimal import Decimal

import pytest

from tests.conftest import add_stock, create_draft_invoice, make_customer, make_product

pytestmark = pytest.mark.django_db

_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")


def _books(company):
    company.accounting_enabled = True
    company.gstin = "29AAAAA0000A1ZY"
    company.save(update_fields=["accounting_enabled", "gstin"])
    from accounting.services import seed_chart_of_accounts

    seed_chart_of_accounts(company)


def _sale(tenant, product, customer, qty="2", price="100.00", date=None):
    payload = [{"product": product.id, "quantity": qty, "unit_price": price, "gst_rate": "18"}]
    inv = create_draft_invoice(tenant, customer, payload)
    done = tenant.client.post(f"/api/v1/sales/invoices/{inv['id']}/complete/")
    assert done.status_code == 200, done.data
    return done.data


def test_sales_register_export_csv_snapshot(tenant_a, assert_snapshot):
    company = tenant_a.company
    _books(company)
    product = make_product(company, name="Export Widget", sku="EXW-1", gst_rate="18", selling_price="100")
    add_stock(tenant_a, product, "20")
    cust = make_customer(company, name="Export Customer", state="Karnataka")
    _sale(tenant_a, product, cust, qty="2")
    _sale(tenant_a, product, cust, qty="3")

    resp = tenant_a.client.get("/api/v1/exports/sales-register/", {"format": "csv"})
    assert resp.status_code == 200, resp.content[:400]
    assert resp["Content-Type"].startswith("text/csv")

    body = b"".join(resp.streaming_content).decode("utf-8") if resp.streaming else resp.content.decode("utf-8")
    reader = list(csv.reader(io.StringIO(body)))
    header, rows = reader[0], reader[1:]

    def _norm(row):
        return [
            "<id>" if i == 0 else (_DATE.sub("<date>", cell)) for i, cell in enumerate(row)
        ]

    assert_snapshot(
        "export_sales_register_csv",
        {"header": header, "rows": sorted(_norm(r) for r in rows), "row_count": len(rows)},
    )


def test_sales_register_export_xlsx_contract(tenant_a):
    company = tenant_a.company
    _books(company)
    product = make_product(company, sku="EXW-2", gst_rate="18")
    add_stock(tenant_a, product, "10")
    cust = make_customer(company, state="Karnataka")
    _sale(tenant_a, product, cust)

    resp = tenant_a.client.get("/api/v1/exports/sales-register/", {"format": "xlsx"})
    assert resp.status_code == 200
    assert resp["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header[:3] == ["id", "number", "date"]
    assert ws.max_row >= 2  # header + at least one data row


def test_filtered_report_snapshots(tenant_a, assert_snapshot):
    company = tenant_a.company
    _books(company)
    from accounting.models import Account

    cash = Account.objects.get(company=company, code="1100").id
    sales = Account.objects.get(company=company, code="4100").id

    for d, amt in (("2026-04-10", "1000.00"), ("2026-05-10", "2000.00"), ("2026-06-10", "3000.00")):
        j = tenant_a.client.post(
            "/api/v1/accounting/journals/",
            {"entry_date": d, "narration": "cash sale", "lines": [
                {"account": cash, "debit": amt, "credit": "0"},
                {"account": sales, "debit": "0", "credit": amt},
            ]},
            format="json",
        )
        assert j.status_code == 201, j.data
        assert tenant_a.client.post(f"/api/v1/accounting/journals/{j.data['id']}/post/").status_code in (200, 202)

    def _rows(report, params):
        r = tenant_a.client.get(f"/api/v1/accounting/{report}/", params)
        assert r.status_code == 200, r.data
        return [
            {"code": row["account_code"], "balance": str(row["balance"])}
            for row in r.data["rows"]
            if str(row["balance"]) not in ("0", "0.00")
        ]

    assert_snapshot(
        "report_filtered",
        {
            "trial_balance_as_of_2026_05_31": sorted(
                _rows("trial-balance", {"as_of": "2026-05-31"}), key=lambda x: x["code"]
            ),
            "pnl_apr_to_may": sorted(
                _rows("profit-and-loss", {"from": "2026-04-01", "to": "2026-05-31"}),
                key=lambda x: x["code"],
            ),
            "pnl_full_q1": sorted(
                _rows("profit-and-loss", {"from": "2026-04-01", "to": "2026-06-30"}),
                key=lambda x: x["code"],
            ),
        },
    )


def test_cost_centre_filtered_pnl(tenant_a, assert_snapshot):
    """§H7 — P&L honours a `cost_center` filter: two expense journals tagged to
    different cost centres, and `?cost_center=` returns only that centre's rows."""
    from accounting.models import Account, CostCenter

    company = tenant_a.company
    _books(company)
    ops = CostCenter.objects.create(company=company, code="OPS", name="Operations")
    mktg = CostCenter.objects.create(company=company, code="MKT", name="Marketing")

    cash = Account.objects.get(company=company, code="1100").id
    # 5xxx expense accounts from the seeded COA
    exp = Account.objects.filter(company=company, type=Account.Type.EXPENSE).order_by("code").first()
    assert exp is not None

    for cc, amt in ((ops, "300.00"), (mktg, "700.00")):
        j = tenant_a.client.post(
            "/api/v1/accounting/journals/",
            {"entry_date": "2026-05-10", "narration": f"spend {cc.code}", "lines": [
                {"account": exp.id, "debit": amt, "credit": "0", "cost_center": cc.id},
                {"account": cash, "debit": "0", "credit": amt},
            ]},
            format="json",
        )
        assert j.status_code == 201, j.data
        assert tenant_a.client.post(
            f"/api/v1/accounting/journals/{j.data['id']}/post/"
        ).status_code in (200, 202)

    def _pnl(params):
        r = tenant_a.client.get("/api/v1/accounting/profit-and-loss/", params)
        assert r.status_code == 200, r.data
        return sorted(
            (
                {"code": row["account_code"], "balance": str(row["balance"])}
                for row in r.data["rows"]
                if row["account_type"] == "EXPENSE" and str(row["balance"]) not in ("0", "0.00")
            ),
            key=lambda x: x["code"],
        )

    all_exp = _pnl({"from": "2026-04-01", "to": "2026-06-30"})
    ops_only = _pnl({"from": "2026-04-01", "to": "2026-06-30", "cost_center": ops.id})
    mktg_only = _pnl({"from": "2026-04-01", "to": "2026-06-30", "cost_center": mktg.id})

    ops_total = sum(Decimal(r["balance"]) for r in ops_only)
    mktg_total = sum(Decimal(r["balance"]) for r in mktg_only)
    assert ops_total == Decimal("300.00")
    assert mktg_total == Decimal("700.00")
    assert sum(Decimal(r["balance"]) for r in all_exp) == Decimal("1000.00")

    assert_snapshot(
        "report_cost_centre_filtered",
        {"all_expense": all_exp, "ops_only": ops_only, "marketing_only": mktg_only},
    )


def test_party_and_godown_filtered_reports(tenant_a, assert_snapshot):
    """§H7 — the register/summary reports honour a party filter and a godown
    filter (the plain snapshots are unfiltered)."""
    from inventory.models import MovementType, Warehouse
    from inventory.services import InventoryService

    company = tenant_a.company
    _books(company)

    cust_a = make_customer(company, name="Party A", state="Karnataka")
    cust_b = make_customer(company, name="Party B", state="Karnataka")
    product = make_product(company, name="Filt Widget", sku="FLT-1", gst_rate="18", selling_price="100")

    main = InventoryService.default_warehouse(company)
    r = tenant_a.client.post(
        "/api/v1/inventory/warehouses/", {"name": "Branch", "code": "BR"}, format="json"
    )
    assert r.status_code == 201, r.data
    branch = Warehouse.objects.get(company=company, code="BR")

    for wh, qty in ((main, "30"), (branch, "12")):
        InventoryService.post_movement(
            company=company, product=product, movement_type=MovementType.OPENING_STOCK,
            quantity=Decimal(qty), unit_cost=Decimal("60"), user=tenant_a.owner, warehouse=wh,
        )

    _sale(tenant_a, product, cust_a, qty="2")
    _sale(tenant_a, product, cust_a, qty="1")
    _sale(tenant_a, product, cust_b, qty="4")

    def _register(params):
        resp = tenant_a.client.get("/api/v1/reports/sales-register/", params)
        assert resp.status_code == 200, resp.data
        rows = resp.data["rows"] if isinstance(resp.data, dict) and "rows" in resp.data else resp.data
        return [
            {"customer": row.get("customer"), "grand_total": str(row.get("grand_total"))}
            for row in rows
        ]

    def _stock(params):
        resp = tenant_a.client.get("/api/v1/reports/inventory-summary/", params)
        assert resp.status_code == 200, resp.data
        rows = resp.data.get("rows", resp.data) if isinstance(resp.data, dict) else resp.data
        return [{"sku": row.get("sku"), "available": str(row.get("available"))} for row in rows]

    all_reg = _register({})
    a_reg = _register({"customer": cust_a.id})
    b_reg = _register({"customer": cust_b.id})
    assert len(all_reg) == 3 and len(a_reg) == 2 and len(b_reg) == 1

    all_stock = _stock({})
    branch_stock = _stock({"warehouse": branch.id})

    def _k(rows, *keys):
        return sorted(rows, key=lambda r: tuple(str(r.get(k) or "") for k in keys))

    assert_snapshot(
        "report_party_godown_filtered",
        {
            "sales_register_all": _k(all_reg, "customer", "grand_total"),
            "sales_register_party_a": _k(a_reg, "customer", "grand_total"),
            "sales_register_party_b": _k(b_reg, "customer", "grand_total"),
            "inventory_summary_all": _k(all_stock, "sku", "available"),
            "inventory_summary_branch_only": _k(branch_stock, "sku", "available"),
        },
    )
