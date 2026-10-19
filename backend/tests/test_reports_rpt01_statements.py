"""RPT-01: balance sheet reports when its date was defaulted; exports carry totals and sections."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from accounting.services import seed_chart_of_accounts

pytestmark = pytest.mark.django_db


@pytest.fixture
def books(tenant_a):
    tenant_a.company.accounting_enabled = True
    tenant_a.company.save(update_fields=["accounting_enabled"])
    seed_chart_of_accounts(tenant_a.company, tenant_a.owner)
    return tenant_a


def _payload(resp):
    body = resp.json()
    return body.get("data", body)


def test_balance_sheet_says_when_its_date_was_filled_in(books):
    defaulted = _payload(books.client.get("/api/v1/accounting/balance-sheet/"))
    assert defaulted["asOfDefaulted"] is True
    explicit = _payload(books.client.get("/api/v1/accounting/balance-sheet/", {"as_of": "2026-09-30"}))
    assert explicit["asOfDefaulted"] is False
    assert explicit["asOf"] == "2026-09-30"


def test_balance_sheet_equity_section_already_holds_current_earnings(books):
    data = _payload(books.client.get("/api/v1/accounting/balance-sheet/"))
    equity = next(s for s in data["scheduleIii"]["sections"] if s["key"] == "equity")
    surplus = [l for l in equity["lines"] if l["key"] == "surplusInProfitAndLoss" or l["key"] == "surplus_in_profit_and_loss"]
    assert len(surplus) == 1, "current earnings must appear exactly once, inside Equity"


@pytest.mark.parametrize("report,sheets", [
    ("trial-balance", {"Summary"}),
    ("profit-and-loss", {"Summary", "Schedule III"}),
    ("balance-sheet", {"Summary", "Schedule III"}),
])
def test_accounting_xlsx_has_summary_and_sections(books, report, sheets):
    resp = books.client.get(f"/api/v1/accounting/{report}/", {"format": "xlsx"})
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    assert sheets <= set(wb.sheetnames)


def test_trial_balance_xlsx_ends_with_a_totals_row(books):
    resp = books.client.get("/api/v1/accounting/trial-balance/", {"format": "xlsx"})
    ws = load_workbook(BytesIO(resp.content)).worksheets[0]
    assert ws.cell(row=ws.max_row, column=2).value == "Total"
