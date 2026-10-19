import csv
from planwave.crypto import reveal_bank_account
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.http import Http404, HttpResponse, StreamingHttpResponse
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView


class Echo:
    """A file-like object that returns written data instead of buffering."""
    def write(self, value):
        return value

from core.csv_utils import csv_safe
from core.exceptions import BusinessRuleError
from core.permissions import (
    CanExport,
    CanViewFinancialReports,
    CanViewMargins,
    CanViewTaxWorksheets,
    HasCompany,
    IsOwner,
    get_company_user,
)
from core.services.feature_flags import build_feature_flags
from core.throttles import CompanyRateThrottle

from masters.models import Customer, Product, Supplier

from .annual_filing import note_deadline
from .gst_health import build_gst_health
from .gst_periods import accounting_period_open_warnings, reopen_period, soft_close_period
from .gst_returns import (
    build_ca_pack_zip,
    build_gstr1,
    build_gstr3b,
    build_gstr9,
    gstr_return_to_xlsx,
    parse_period,
    persist_snapshot,
    to_gstn_json,
)
from .models import Gstr2bIngest
from .permissions import assert_gstr_enabled
from .serializers import Gstr2bIngestSerializer
from .services import ReportService, assert_report_date_span


MAX_SUMMARY_INVOICES = 20000


def _tds_worksheets_readable(company) -> bool:
    """Owner may export worksheets when TDS is enabled or TDS/TCS was already posted."""
    if build_feature_flags(company=company).get("ENABLE_TDS"):
        return True
    from payments.models import SupplierPayment
    from purchases.models import PurchaseInvoice
    from sales.models import SalesInvoice

    return (
        PurchaseInvoice.objects.filter(company=company, tds_amount__gt=0).exists()
        or SalesInvoice.objects.filter(company=company, tcs_amount__gt=0).exists()
        or SupplierPayment.objects.filter(company=company, tds_amount__gt=0).exists()
    )


def _parse_date(value):
    if not value:
        return None
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (ValueError, TypeError):
        return None


def _int_or_none(value):
    if value in (None, "", "undefined", "null"):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _maybe_gstn_json(request, payload):
    export_format = (request.query_params.get("format") or "json").lower()
    if export_format not in ("gstn-json", "gstn_json"):
        return None
    from django.conf import settings
    from core.services.feature_flags import flag_enabled

    allowed = bool(getattr(settings, "ENABLE_GSTN_JSON", False))
    if not allowed:
        raise BusinessRuleError("GSTN-shaped JSON export is disabled.")
    company = get_company_user(request).company
    if not flag_enabled(company, "ENABLE_GSTN_JSON"):
        raise BusinessRuleError("GSTN-shaped JSON export is disabled.")
    return Response(to_gstn_json(payload))


class BaseReportView(APIView):
    permission_classes = [IsAuthenticated, HasCompany, CanViewFinancialReports]

    @property
    def company(self):
        return get_company_user(self.request).company


class DashboardView(BaseReportView):
    def get(self, request):
        # BB-000513: KPIs are live aggregates over source documents — no CQRS rollup layer.
        return Response(ReportService.dashboard(self.company))


class SalesRegisterView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Sales register")
        data = ReportService.sales_register(
            self.company,
            date_from=date_from,
            date_to=date_to,
            customer_id=_int_or_none(request.query_params.get("customer")),
            status=request.query_params.get("status"),
            warehouse_id=_int_or_none(request.query_params.get("warehouse")),
            company_gstin_id=_int_or_none(
                request.query_params.get("company_gstin")
                or request.query_params.get("gstin")
            ),
            product_id=_int_or_none(request.query_params.get("product")),
        )
        bucket = (request.query_params.get("time_bucket") or "").strip().lower()
        if bucket in {"day", "week", "month", "fy"}:
            from reporting.time_buckets import group_register_rows

            data = dict(data)
            data["buckets"] = group_register_rows(data.get("rows") or [], bucket=bucket)
        page = request.query_params.get("page")
        page_size = request.query_params.get("page_size")
        if page or page_size or request.query_params.get("paginate"):
            try:
                p = max(1, int(page or 1))
                ps = min(250, max(1, int(page_size or 100)))
            except (ValueError, TypeError):
                p, ps = 1, 100
            rows = data.get("rows", [])
            total_count = len(rows)
            start_idx = (p - 1) * ps
            end_idx = start_idx + ps
            paginated_rows = rows[start_idx:end_idx]
            return Response({
                "count": total_count,
                "page": p,
                "page_size": ps,
                "results": paginated_rows,
                "rows": paginated_rows,
                "totals": data.get("totals", {}),
            })
        return Response(data)


class PurchaseRegisterView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Purchase register")
        data = ReportService.purchase_register(
            self.company,
            date_from=date_from,
            date_to=date_to,
            supplier_id=_int_or_none(request.query_params.get("supplier")),
            status=request.query_params.get("status"),
            warehouse_id=_int_or_none(request.query_params.get("warehouse")),
            company_gstin_id=_int_or_none(
                request.query_params.get("company_gstin")
                or request.query_params.get("gstin")
            ),
            product_id=_int_or_none(request.query_params.get("product")),
        )
        bucket = (request.query_params.get("time_bucket") or "").strip().lower()
        if bucket in {"day", "week", "month", "fy"}:
            from reporting.time_buckets import group_register_rows

            data = dict(data)
            data["buckets"] = group_register_rows(data.get("rows") or [], bucket=bucket)
        page = request.query_params.get("page")
        page_size = request.query_params.get("page_size")
        if page or page_size or request.query_params.get("paginate"):
            try:
                p = max(1, int(page or 1))
                ps = min(250, max(1, int(page_size or 100)))
            except (ValueError, TypeError):
                p, ps = 1, 100
            rows = data.get("rows", [])
            total_count = len(rows)
            start_idx = (p - 1) * ps
            end_idx = start_idx + ps
            paginated_rows = rows[start_idx:end_idx]
            return Response({
                "count": total_count,
                "page": p,
                "page_size": ps,
                "results": paginated_rows,
                "rows": paginated_rows,
                "totals": data.get("totals", {}),
            })
        return Response(data)


class InventorySummaryView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        return Response(ReportService.inventory_summary(
            self.company,
            # CR-068: normalize warehouse like sibling report views.
            warehouse_id=_int_or_none(request.query_params.get("warehouse")),
        ))


class ProductSalesView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        # CR-150: product sales honor the same date-span guard as registers.
        assert_report_date_span(date_from, date_to, kind="Product sales")
        return Response(ReportService.product_sales(
            self.company,
            date_from=date_from,
            date_to=date_to,
        ))


class CustomerSalesView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        # CR-150: customer sales honor the same date-span guard as registers.
        assert_report_date_span(date_from, date_to, kind="Customer sales")
        return Response(ReportService.customer_sales(
            self.company,
            date_from=date_from,
            date_to=date_to,
        ))


class SalesDiscountReportView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Sales discount report")
        return Response(ReportService.sales_discount_report(
            self.company,
            date_from=date_from,
            date_to=date_to,
            customer_id=_int_or_none(request.query_params.get("customer")),
        ))


class PurchaseDiscountReportView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Purchase discount report")
        return Response(ReportService.purchase_discount_report(
            self.company,
            date_from=date_from,
            date_to=date_to,
            supplier_id=_int_or_none(request.query_params.get("supplier")),
        ))


def _certificate_filter(raw):
    value = (raw or "").strip().lower()
    if value and value not in ("pending", "received"):
        raise BusinessRuleError("certificate must be pending or received.")
    return value or None


def _tds_receivable_rows(company, date_from, date_to, customer_id=None, certificate=None):
    """Receipts on which a customer deducted TDS, for matching against Form 26AS.

    One row per receipt: the PAN on the customer master, the section the cashier chose, the TDS
    deducted, and how much of it has been set against invoices so far. The same figure sits in
    ledger account 1270 until the customer's certificate (Form 16A) lets it be claimed.
    """
    from django.db.models import Q, Sum

    from payments.models import CustomerReceipt, ReceiptStatus

    qs = CustomerReceipt.objects.filter(
        company=company, status=ReceiptStatus.POSTED, tds_deducted__gt=0
    ).select_related("customer")
    if date_from:
        qs = qs.filter(receipt_date__gte=date_from)
    if date_to:
        qs = qs.filter(receipt_date__lte=date_to)
    if customer_id:
        qs = qs.filter(customer_id=customer_id)
    if certificate == "pending":
        qs = qs.filter(tds_certificate_received_on__isnull=True)
    elif certificate == "received":
        qs = qs.filter(tds_certificate_received_on__isnull=False)
    qs = qs.annotate(
        applied=Sum("allocations__tds_applied", filter=Q(allocations__reversed_at__isnull=True))
    ).order_by("receipt_date", "id")
    rows = []
    for receipt in qs:
        applied = receipt.applied or Decimal("0")
        rows.append({
            "receipt_id": receipt.id,
            "receipt_number": receipt.number,
            "receipt_date": receipt.receipt_date.isoformat(),
            "customer_id": receipt.customer_id,
            "customer": receipt.customer.name,
            "pan": receipt.customer.pan or "",
            "section": receipt.tds_section or "",
            "amount_received": str(receipt.amount),
            "tds_deducted": str(receipt.tds_deducted),
            "tds_applied": str(applied),
            "tds_unapplied": str(receipt.tds_deducted - applied),
            "certificate_received_on": (
                receipt.tds_certificate_received_on.isoformat() if receipt.tds_certificate_received_on else None
            ),
            "certificate_ref": receipt.tds_certificate_ref or "",
        })
    return rows


class TdsReceivableReportView(BaseReportView):
    """TDS deducted by customers on receipts (Form 26AS reconciliation)."""

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="TDS receivable register")
        rows = _tds_receivable_rows(
            self.company, date_from, date_to, _int_or_none(request.query_params.get("customer")),
            certificate=_certificate_filter(request.query_params.get("certificate")),
        )
        total = sum((Decimal(row["tds_deducted"]) for row in rows), Decimal("0"))
        return Response({"rows": rows, "total_tds": str(total), "count": len(rows)})


class TdsReceivableExportView(BaseReportView):
    """The register as a CSV (cells are neutralised against spreadsheet formulas)."""

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get_permissions(self):
        return [*super().get_permissions(), CanExport()]

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="TDS receivable register")
        rows = _tds_receivable_rows(
            self.company, date_from, date_to, _int_or_none(request.query_params.get("customer")),
            certificate=_certificate_filter(request.query_params.get("certificate")),
        )
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="tds-receivable.csv"'
        writer = csv.writer(response)
        writer.writerow(["Receipt", "Date", "Customer", "PAN", "Section", "Amount received", "TDS deducted", "TDS set against invoices", "TDS not yet set against invoices", "Certificate received on", "Certificate reference"])
        for row in rows:
            writer.writerow([csv_safe(v) for v in (
                row["receipt_number"], row["receipt_date"], row["customer"], row["pan"], row["section"],
                row["amount_received"], row["tds_deducted"], row["tds_applied"], row["tds_unapplied"],
                row["certificate_received_on"] or "", row["certificate_ref"],
            )])
        return response


class InvoiceProfitReportView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"
    # RPT-12 / D4: cost and margin are for the owner and accountants, not everyone with report access.
    permission_classes = [IsAuthenticated, HasCompany, CanViewMargins]

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Invoice profit report")
        return Response(ReportService.invoice_profit_report(
            self.company,
            date_from=date_from,
            date_to=date_to,
            customer_id=_int_or_none(request.query_params.get("customer")),
            cost_center_id=_int_or_none(request.query_params.get("cost_center")),
            cost_basis=request.query_params.get("cost_basis"),
            status=request.query_params.get("status"),
        ))


class InvoiceProfitRollupView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"
    # RPT-12 / D4: cost and margin are for the owner and accountants, not everyone with report access.
    permission_classes = [IsAuthenticated, HasCompany, CanViewMargins]

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Invoice profit rollup")
        group_by = request.query_params.get("group_by") or "customer"
        return Response(ReportService.invoice_profit_rollup(
            self.company,
            group_by,
            date_from=date_from,
            date_to=date_to,
            customer_id=_int_or_none(request.query_params.get("customer")),
            cost_center_id=_int_or_none(request.query_params.get("cost_center")),
        ))


class CashBookView(BaseReportView):
    """Cash book actuals — JSON or XLSX."""

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get_permissions(self):
        # BB-000137: XLSX export requires CanExport (JSON preview keeps financial-report ACL).
        export_format = (
            self.request.query_params.get("export")
            or self.request.query_params.get("format")
            or "json"
        ).lower()
        if export_format == "xlsx":
            return [IsAuthenticated(), HasCompany(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request):
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Cash book")
        payload = ReportService.cash_book(
            self.company,
            date_from=date_from,
            date_to=date_to,
            bank_account_id=_int_or_none(request.query_params.get("bank_account")),
        )
        # Prefer export= over format= — DRF treats ?format= as renderer negotiation.
        export_format = (
            request.query_params.get("export")
            or request.query_params.get("format")
            or "json"
        ).lower()
        if export_format in ("json", "api"):
            return Response(payload)
        if export_format == "xlsx":
            from io import BytesIO

            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Cash book"
            headers = [
                "date",
                "type",
                "number",
                "party",
                "mode",
                "bank_account",
                "source",
                "inflow",
                "outflow",
                "reference",
            ]
            ws.append(headers)
            for row in payload.get("rows") or []:
                ws.append([row.get(h) for h in headers])
            ws2 = wb.create_sheet("Summary")
            ws2.append(["opening", payload.get("opening")])
            ws2.append(["inflow", payload.get("inflow")])
            ws2.append(["outflow", payload.get("outflow")])
            ws2.append(["closing", payload.get("closing")])
            buf = BytesIO()
            wb.save(buf)
            response = HttpResponse(
                buf.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = 'attachment; filename="cash-book.xlsx"'
            return response
        raise BusinessRuleError("Unsupported format. Use export=json or export=xlsx.")


class DayBookView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get_permissions(self):
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request):
        from django.utils import timezone

        from .transactions import day_book

        raw = request.query_params.get("date") or request.query_params.get("on_date")
        if raw and _parse_date(raw) is None:
            # ACS-60: a typo in the date is an error, not a silent switch to today.
            raise BusinessRuleError("date must be a date like 2026-04-01.", code="bad_date")
        on_date = _parse_date(raw) or timezone.localdate()
        return Response(day_book(self.company, on_date))


class LedgerDayBookView(BaseReportView):
    """ACS-60: every ledger line posted in a date range, with a tie-out to the trial balance."""

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get_permissions(self):
        if (self.request.query_params.get("file") or "").lower() in ("csv", "xlsx"):
            return [IsAuthenticated(), HasCompany(), CanViewFinancialReports(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request):
        from django.utils import timezone

        from core.pagination import DefaultPagination

        from . import day_book as ledger

        params = request.query_params
        start, end = ledger.parse_range(
            params.get("date_from") or params.get("date"), params.get("date_to") or params.get("date"), timezone.localdate(),
        )
        qs = ledger.base_queryset(self.company, start, end, params)
        file_kind = (params.get("file") or "").lower()
        if file_kind in ("csv", "xlsx"):
            return self._export(qs[:20000], file_kind)
        everything = ledger.base_queryset(self.company, start, end, {})
        filtered = ledger.totals(qs)
        whole = ledger.totals(everything)
        paginator = DefaultPagination()
        page = paginator.paginate_queryset(qs, request, view=self)
        payload = {
            "date_from": start.isoformat(),
            "date_to": end.isoformat(),
            "results": [ledger.row_of(line) for line in page],
            "count": paginator.page.paginator.count,
            "totals": filtered,
            "day_totals": ledger.per_day_balance(everything),
            "voucher_types": ledger.available_types(self.company, start, end),
            "tie_out": ledger.tie_out(self.company, start, end, Decimal(whole["debit"]), Decimal(whole["credit"])),
            **ledger.neighbours(start),
        }
        return Response(payload)

    def _export(self, lines, kind):
        from . import day_book as ledger

        header = ["Date", "Voucher", "Type", "Account", "Party", "Cost centre", "Debit", "Credit", "Narration"]
        rows = []
        for line in lines:
            r = ledger.row_of(line)
            rows.append([
                r["date"], r["voucher"], r["voucher_type_label"], f"{r['account_code']} {r['account_name']}",
                r["party"], r["cost_center"], r["debit"], r["credit"], r["narration"],
            ])
        if kind == "xlsx":
            from io import BytesIO

            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Day book"
            ws.append(header)
            for r in rows:
                ws.append(r)
            buf = BytesIO()
            wb.save(buf)
            resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            resp["Content-Disposition"] = 'attachment; filename="day-book.xlsx"'
            return resp
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = 'attachment; filename="day-book.csv"'
        writer = csv.writer(resp)
        writer.writerow(header)
        for r in rows:
            writer.writerow([csv_safe(c) if isinstance(c, str) else c for c in r])
        return resp


class SalesSummaryView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get_permissions(self):
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request):

        from sales.models import SalesInvoice

        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        assert_report_date_span(date_from, date_to, kind="Sales summary")
        # RPT-04: every figure below (headline, by-date, tops) uses one definition of sales;
        # see ReportService.sales_totals. The payment split stays on completed invoices only.
        sales = ReportService.sales_totals(self.company, date_from, date_to)
        qs = SalesInvoice.objects.filter(
            company=self.company, status=SalesInvoice.Status.COMPLETED,
        )
        if date_from:
            qs = qs.filter(invoice_date__gte=date_from)
        if date_to:
            qs = qs.filter(invoice_date__lte=date_to)
        if qs.count() > MAX_SUMMARY_INVOICES:
            raise BusinessRuleError(
                f"More than {MAX_SUMMARY_INVOICES} invoices in this range. Choose a shorter date range."
            )
        totals = {
            "invoice_count": sales["invoice_count"],
            "revenue": sales["revenue"],
            "taxable": sales["net_taxable"],
            "gross_invoiced": sales["gross_invoiced"],
            "credit_notes": sales["credit_notes"],
            "debit_notes": sales["debit_notes"],
            "gst_and_other": sales["gst_and_other"],
        }
        # Reuse the existing, CN/DN-netted reports (CR-066/CR-150) rather than
        # a fresh unnetted aggregation, so "top customers/products" here can't
        # disagree with the dedicated Customer Sales / Product Sales reports.
        top_customers = ReportService.customer_sales(self.company, date_from, date_to)["rows"][:10]
        top_products = ReportService.product_sales(self.company, date_from, date_to)["rows"][:10]
        by_date = sales["by_date"]
        from decimal import Decimal

        from ledgers.services import LedgerService

        # Same outstanding calculation Day Book / the invoice list use (and
        # the one that correctly accounts for settlement_discount), so this
        # breakdown can't drift from what those screens show for the same
        # invoices.
        balances = LedgerService.bulk_sales_invoice_outstanding(self.company, list(qs.values_list("id", flat=True)))
        paid = unpaid = partial = 0
        paid_amt = unpaid_amt = partial_amt = Decimal("0")
        for inv in qs.values("id", "grand_total"):
            bal = balances.get(inv["id"]) or Decimal("0")
            grand = inv["grand_total"] or Decimal("0")
            if bal <= Decimal("0.05"):
                paid += 1
                paid_amt += grand
            elif bal + Decimal("0.05") >= grand:
                unpaid += 1
                unpaid_amt += grand
            else:
                partial += 1
                partial_amt += grand
        return Response({
            "date_from": date_from,
            "date_to": date_to,
            "totals": totals,
            "top_customers": top_customers,
            "top_products": top_products,
            "by_date": by_date,
            "payment_breakdown": {
                "paid": {"count": paid, "amount": paid_amt},
                "partial": {"count": partial, "amount": partial_amt},
                "unpaid": {"count": unpaid, "amount": unpaid_amt},
            },
        })


class CustomerLedgerTabsView(BaseReportView):
    def get_permissions(self):
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request, customer_id):
        from decimal import Decimal

        from ledgers.services import LedgerService
        from masters.models import Customer
        from rest_framework.generics import get_object_or_404
        from sales.models import SalesInvoice

        from .transactions import customer_item_wise, customer_overdue_amount, sales_side_transactions

        customer = get_object_or_404(Customer, pk=customer_id, company=self.company)
        date_from = _parse_date(request.query_params.get("date_from"))
        date_to = _parse_date(request.query_params.get("date_to"))
        txn_type = request.query_params.get("txn_type") or request.query_params.get("type")
        types = [txn_type] if txn_type else None
        status_filter = request.query_params.get("status")
        transactions = sales_side_transactions(
            self.company, date_from=date_from, date_to=date_to,
            customer_id=customer.id, txn_types=types, status_filter=status_filter,
        )
        outstanding = LedgerService.customer_outstanding(self.company, customer)
        sales_qs = SalesInvoice.objects.filter(
            company=self.company, customer=customer, status=SalesInvoice.Status.COMPLETED,
        )
        if date_from:
            sales_qs = sales_qs.filter(invoice_date__gte=date_from)
        if date_to:
            sales_qs = sales_qs.filter(invoice_date__lte=date_to)
        from django.db.models import Sum
        from payments.models import CustomerReceipt, ReceiptStatus

        total_sales = sales_qs.aggregate(t=Sum("grand_total"))["t"] or Decimal("0")
        receipts = CustomerReceipt.objects.filter(
            company=self.company, customer=customer, status=ReceiptStatus.POSTED,
        )
        if date_from:
            receipts = receipts.filter(receipt_date__gte=date_from)
        if date_to:
            receipts = receipts.filter(receipt_date__lte=date_to)
        total_received = receipts.aggregate(t=Sum("amount"))["t"] or Decimal("0")
        overdue = customer_overdue_amount(
            self.company, customer, date_from=date_from, date_to=date_to,
        )
        statement = LedgerService.customer_statement(
            self.company, customer, date_from=date_from, date_to=date_to,
        )
        if isinstance(statement, list):
            raw_entries = statement
        else:
            raw_entries = (statement or {}).get("entries") or []
        numbered = []
        for i, entry in enumerate(raw_entries, start=1):
            row = dict(entry)
            row["sr_no"] = i
            row.setdefault("mode", "")
            numbered.append(row)
        statement = {
            "outstanding": outstanding,
            "uncleared_cheques": LedgerService.uncleared_cheque_total(self.company, customer),
            "entries": numbered,
        }
        item_wise = customer_item_wise(
            self.company, customer.id, date_from=date_from, date_to=date_to,
        )
        payload = {
            "customer_id": customer.id,
            "customer_name": customer.name,
            "profile": {
                "name": customer.name,
                "phone": customer.phone,
                "email": customer.email,
                "gstin": customer.gstin,
                "pan": getattr(customer, "pan", ""),
                "billing_address": customer.billing_address,
                "shipping_address": customer.shipping_address,
                "shipping_addresses": [
                    {"id": a.id, "label": a.label, "address": a.address, "is_default": a.is_default}
                    for a in customer.shipping_addresses.all()
                ],
                "credit_limit": customer.credit_limit,
                "credit_days": customer.credit_days,
                "custom_fields": customer.custom_fields,
                "party_bank_name": getattr(customer, "party_bank_name", ""),
                "party_bank_account": reveal_bank_account(getattr(customer, "party_bank_account", "") or ""),
                "party_bank_ifsc": getattr(customer, "party_bank_ifsc", ""),
            },
            "kpis": {
                "total_receivable": outstanding,
                "overdue_amount": overdue,
                "total_sales_amount": total_sales,
                "total_received_amount": total_received,
            },
            "transactions": transactions,
            "statement": statement,
            "item_wise": item_wise,
        }
        export_format = (request.query_params.get("export") or "json").lower()
        if export_format in ("json", "api", ""):
            return Response(payload)
        if export_format == "xlsx":
            from io import BytesIO

            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Transactions"
            txn_headers = [
                "date", "txn_type", "number", "party", "debit", "credit", "balance",
                "mode", "reference",
            ]
            ws.append(txn_headers)
            for row in payload.get("transactions") or []:
                ws.append([row.get(h) for h in txn_headers])
            ws2 = wb.create_sheet("Statement")
            ws2.append(["sr_no", "date", "type", "number", "debit", "credit", "balance", "mode"])
            raw_statement = payload.get("statement")
            if isinstance(raw_statement, list):
                entries = raw_statement
            else:
                entries = (raw_statement or {}).get("entries") or []
            for entry in entries:
                ws2.append([
                    entry.get("sr_no"),
                    entry.get("date"),
                    entry.get("type") or entry.get("particulars") or entry.get("description"),
                    entry.get("number"),
                    entry.get("debit"),
                    entry.get("credit"),
                    entry.get("balance"),
                    entry.get("mode"),
                ])
            ws3 = wb.create_sheet("Item-wise")
            item_headers = ["product", "product_name", "qty", "amount"]
            ws3.append(item_headers)
            for row in payload.get("item_wise") or []:
                ws3.append([
                    row.get("product_id") or row.get("product"),
                    row.get("product_name") or row.get("name"),
                    row.get("qty") or row.get("quantity"),
                    row.get("amount") or row.get("line_total"),
                ])
            ws4 = wb.create_sheet("KPIs")
            for key, val in (payload.get("kpis") or {}).items():
                ws4.append([key, val])
            buf = BytesIO()
            wb.save(buf)
            response = HttpResponse(
                buf.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (customer.name or "customer"))
            response["Content-Disposition"] = f'attachment; filename="{safe_name}-ledger.xlsx"'
            return response
        raise BusinessRuleError("Unsupported format. Use export=json or export=xlsx.")


def _customer_export(company, params):
    qs = Customer.objects.filter(company=company).order_by("name")
    status = params.get("status")
    if status:
        qs = qs.filter(status=status)
    return [
        {
            "id": c.id,
            "name": c.name,
            "phone": c.phone,
            "email": c.email,
            "gstin": c.gstin,
            "state": c.state,
            "taxpayer_type": c.taxpayer_type,
            "status": c.status,
            "billing_address": c.billing_address,
            "shipping_address": c.shipping_address,
            "credit_limit": c.credit_limit,
            "credit_days": c.credit_days,
            "gstin_verification_status": c.gstin_verification_status,
        }
        for c in qs
    ]


def _supplier_export(company, params):
    qs = Supplier.objects.filter(company=company).order_by("name")
    # `status` used to name a column suppliers do not have, which made this export fail outright.
    status = (params.get("status") or "").upper()
    if status in ("ACTIVE", "INACTIVE"):
        qs = qs.filter(is_active=(status == "ACTIVE"))
    return [
        {
            "id": s.id,
            "name": s.name,
            "phone": s.phone,
            "email": s.email,
            "gstin": s.gstin,
            "state": s.state,
            "taxpayer_type": s.taxpayer_type,
            "status": "ACTIVE" if s.is_active else "INACTIVE",
            "address": s.address,
            "gstin_verification_status": s.gstin_verification_status,
        }
        for s in qs
    ]


def _inventory_export(company, params):
    """INV-IDR-10: the stock CSV follows the godown and search the page was showing."""
    rows = ReportService.inventory_summary(company, warehouse_id=params.get("warehouse"))["rows"]
    needle = (params.get("q") or params.get("search") or "").strip().lower()
    if needle:
        rows = [r for r in rows if needle in f"{r.get('product') or ''} {r.get('sku') or ''}".lower()]
    return rows


def _product_export(company, params):
    """The item master. Purchase price is dropped by the view for anyone who may not see margins."""
    qs = (
        Product.objects.filter(company=company, is_deleted=False)
        .select_related("unit", "category", "brand")
        .order_by("name", "id")
    )
    return [
        {
            "id": p.id,
            "sku": p.sku,
            "barcode": p.barcode,
            "name": p.name,
            "category": p.category.name if p.category_id else "",
            "brand": p.brand.name if p.brand_id else "",
            "unit": p.unit.short_name if p.unit_id else "",
            "hsn_code": p.hsn_code,
            "gst_rate": p.gst_rate,
            "selling_price": p.selling_price,
            "mrp": p.mrp,
            "wholesale_price": p.wholesale_price,
            "purchase_price": p.purchase_price,
            "reorder_level": p.reorder_level,
            "track_inventory": p.track_inventory,
            "track_batch": p.track_batch,
            "track_serial": p.track_serial,
            "regulated_category": p.regulated_category,
        }
        for p in qs
    ]


def _price_list_export(company, params):
    """One row per price list slab."""
    from masters.models import PriceListItem

    qs = (
        PriceListItem.objects.filter(company=company, price_list__is_deleted=False)
        .select_related("price_list", "product")
        .order_by("price_list__name", "product__name", "min_qty", "id")
    )
    return [
        {
            "price_list": item.price_list.name,
            "price_list_active": item.price_list.is_active,
            "sku": item.product.sku,
            "product": item.product.name,
            "min_qty": item.min_qty,
            "max_qty": item.max_qty if item.max_qty is not None else "",
            "unit_price": item.unit_price,
            "discount_pct": item.discount_pct,
        }
        for item in qs
    ]


# Columns a user may only see with margin access (owner, or an accountant who may read reports).
MARGIN_ONLY_COLUMNS = {"products": {"purchase_price"}, "items": {"purchase_price"}}

EXPORTS = {
    "products": lambda company, params: _product_export(company, params),
    "items": lambda company, params: _product_export(company, params),
    "price-lists": lambda company, params: _price_list_export(company, params),
    "sales-register": lambda company, params: ReportService.sales_register(
        company,
        date_from=_parse_date(params.get("date_from")),
        date_to=_parse_date(params.get("date_to")),
        customer_id=_int_or_none(params.get("customer")),
        status=params.get("status"),
    )["rows"],
    "sales": lambda company, params: ReportService.sales_register(
        company,
        date_from=_parse_date(params.get("date_from")),
        date_to=_parse_date(params.get("date_to")),
        customer_id=_int_or_none(params.get("customer")),
        status=params.get("status"),
    )["rows"],
    "purchase-register": lambda company, params: ReportService.purchase_register(
        company,
        date_from=_parse_date(params.get("date_from")),
        date_to=_parse_date(params.get("date_to")),
        supplier_id=_int_or_none(params.get("supplier")),
        status=params.get("status"),
    )["rows"],
    "purchases": lambda company, params: ReportService.purchase_register(
        company,
        date_from=_parse_date(params.get("date_from")),
        date_to=_parse_date(params.get("date_to")),
        supplier_id=_int_or_none(params.get("supplier")),
        status=params.get("status"),
    )["rows"],
    "inventory-summary": lambda company, params: _inventory_export(company, params),
    "inventory": lambda company, params: _inventory_export(company, params),
    "customers": lambda company, params: _customer_export(company, params),
    "customers-register": lambda company, params: _customer_export(company, params),
    "customer-list": lambda company, params: _customer_export(company, params),
    "suppliers": lambda company, params: _supplier_export(company, params),
    "suppliers-register": lambda company, params: _supplier_export(company, params),
    "supplier-list": lambda company, params: _supplier_export(company, params),
}

# BUG-323: known column names per export, used as a CSV header fallback
# when a filter narrows the result down to zero rows.
EXPORT_FIELDS = {
    "products": [
        "id", "sku", "barcode", "name", "category", "brand", "unit", "hsn_code", "gst_rate",
        "selling_price", "mrp", "wholesale_price", "purchase_price", "reorder_level",
        "track_inventory", "track_batch", "track_serial", "regulated_category",
    ],
    "items": [
        "id", "sku", "barcode", "name", "category", "brand", "unit", "hsn_code", "gst_rate",
        "selling_price", "mrp", "wholesale_price", "purchase_price", "reorder_level",
        "track_inventory", "track_batch", "track_serial", "regulated_category",
    ],
    "price-lists": [
        "price_list", "price_list_active", "sku", "product", "min_qty", "max_qty", "unit_price", "discount_pct",
    ],
    "sales-register": [
        "id", "number", "date", "customer", "invoice_type", "status",
        "taxable", "cgst", "sgst", "igst", "grand_total",
    ],
    "sales": [
        "id", "number", "date", "customer", "invoice_type", "status",
        "taxable", "cgst", "sgst", "igst", "grand_total",
    ],
    "purchase-register": [
        "id", "number", "date", "supplier", "status",
        "taxable", "cgst", "sgst", "igst", "grand_total",
    ],
    "purchases": [
        "id", "number", "date", "supplier", "status",
        "taxable", "cgst", "sgst", "igst", "grand_total",
    ],
    "inventory-summary": [
        "product_id", "product", "sku", "on_hand", "reserved",
        "available", "reorder_level", "stock_value",
    ],
    "inventory": [
        "product_id", "product", "sku", "on_hand", "reserved",
        "available", "reorder_level", "stock_value",
    ],
    "customers": [
        "id", "name", "phone", "email", "gstin", "state", "taxpayer_type",
        "status", "billing_address", "shipping_address", "credit_limit", "credit_days", "gstin_verification_status",
    ],
    "customers-register": [
        "id", "name", "phone", "email", "gstin", "state", "taxpayer_type",
        "status", "billing_address", "shipping_address", "credit_limit", "credit_days", "gstin_verification_status",
    ],
    "customer-list": [
        "id", "name", "phone", "email", "gstin", "state", "taxpayer_type",
        "status", "billing_address", "shipping_address", "credit_limit", "credit_days", "gstin_verification_status",
    ],
    "suppliers": [
        "id", "name", "phone", "email", "gstin", "state", "taxpayer_type", "status",
        "address", "gstin_verification_status",
    ],
    "suppliers-register": [
        "id", "name", "phone", "email", "gstin", "state", "taxpayer_type", "status",
        "address", "gstin_verification_status",
    ],
    "supplier-list": [
        "id", "name", "phone", "email", "gstin", "state", "taxpayer_type", "status",
        "address", "gstin_verification_status",
    ],
}


class GstReturnView(BaseReportView):
    """Offline GSTR-1 / GSTR-3B JSON preview or XLSX download."""

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"
    build_fn = None

    def perform_content_negotiation(self, request, force=False):
        # DRF treats ?format= as renderer selection; we use it for json|xlsx export.
        from rest_framework.renderers import JSONRenderer

        return (JSONRenderer(), "application/json")

    def get_permissions(self):
        export_format = (self.request.query_params.get("format") or "json").lower()
        # RPT-09: every file format needs export rights, GSTN-shaped JSON included.
        if export_format in ("xlsx", "gstn-json", "gstn_json"):
            return [IsAuthenticated(), HasCompany(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request):
        from planwave.services import assert_books_clear

        assert_books_clear(self.company)
        assert_gstr_enabled(self.company)
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        try:
            parse_period(period)
        except ValueError as exc:
            raise BusinessRuleError(str(exc)) from exc

        export_format = (request.query_params.get("format") or "json").lower()
        company_gstin = request.query_params.get("company_gstin") or request.query_params.get(
            "companyGstin"
        )
        payload = self.build_fn(self.company, period, company_gstin=company_gstin or None)
        persist = request.query_params.get("persist") in ("1", "true", "True")
        # RPT-10: a plain JSON preview is a view, not a download. Only files and saved
        # snapshots are audited, so the trail does not say "downloaded" on every page load.
        if export_format != "json" or persist:
            from core.models import AuditEvent

            AuditEvent.objects.create(
                company=self.company,
                user=request.user,
                action=AuditEvent.Action.CREATE,
                entity_type="GstReturnJson",
                entity_id=period,
                description=(
                    f"{payload.get('return_type', 'GSTR')} snapshot saved"
                    if export_format == "json"
                    else f"{payload.get('return_type', 'GSTR')} downloaded"
                ),
                metadata={"format": export_format, "period": period, "persist": persist},
            )
        if persist:
            persist_snapshot(
                self.company,
                payload.get("return_type", "GSTR-1"),
                period,
                payload,
                user=request.user,
            )

        if export_format == "json":
            return Response(payload)
        gstn = _maybe_gstn_json(request, payload)
        if gstn is not None:
            return gstn
        if export_format == "xlsx":
            content = gstr_return_to_xlsx(payload)
            filename = f"{payload['return_type'].lower().replace('-', '')}-{period}.xlsx"
            response = HttpResponse(
                content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            return response
        raise BusinessRuleError("Unsupported format. Use format=json, format=xlsx, or format=gstn-json.")


class Gstr1View(GstReturnView):
    build_fn = staticmethod(build_gstr1)


class Gstr3bView(GstReturnView):
    build_fn = staticmethod(build_gstr3b)


class GstSnapshotListView(BaseReportView):
    """RPT-07: the saved GSTR-1 / GSTR-3B snapshots, newest first, so a saved month can be found again."""

    def get(self, request):
        from .models import GstReturnSnapshot
        from .permissions import assert_gstr_enabled

        assert_gstr_enabled(self.company)
        qs = GstReturnSnapshot.objects.filter(company=self.company).select_related("generated_by")
        wanted = (request.query_params.get("return_type") or "").upper().replace("-", "")
        if wanted in {"GSTR1", "GSTR3B", "GSTR9"}:
            qs = qs.filter(return_type=wanted)
        period = request.query_params.get("period")
        if period:
            qs = qs.filter(period=period)
        rows = [
            {
                "id": snap.id,
                "return_type": snap.return_type,
                "period": snap.period,
                "content_hash": snap.content_hash[:12],
                "builder_version": snap.builder_version,
                "generated_at": snap.generated_at,
                "generated_by": (getattr(snap.generated_by, "full_name", "") or getattr(snap.generated_by, "email", "")) if snap.generated_by_id else "",
            }
            for snap in qs[:50]
        ]
        return Response({"results": rows, "count": len(rows)})


class Gstr9View(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def perform_content_negotiation(self, request, force=False):
        from rest_framework.renderers import JSONRenderer

        return (JSONRenderer(), "application/json")

    def get_permissions(self):
        export_format = (self.request.query_params.get("format") or "json").lower()
        if export_format == "xlsx":
            return [IsAuthenticated(), HasCompany(), CanExport()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    def get(self, request):
        assert_gstr_enabled(self.company)
        fy = request.query_params.get("fy")
        if not fy:
            raise BusinessRuleError("Query parameter 'fy' is required (e.g. 2025-26).")
        company_gstin = request.query_params.get("company_gstin") or None
        payload = build_gstr9(self.company, fy, company_gstin=company_gstin)
        export_format = (request.query_params.get("format") or "json").lower()
        if export_format == "json":
            return Response(payload)
        gstn = _maybe_gstn_json(request, payload)
        if gstn is not None:
            return gstn
        if export_format == "xlsx":
            content = gstr_return_to_xlsx(payload)
            response = HttpResponse(
                content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = f'attachment; filename="gstr9-{fy}.xlsx"'
            return response
        raise BusinessRuleError("Unsupported format. Use format=json, format=xlsx, or format=gstn-json.")


class Gstr9FilingDateView(BaseReportView):
    """Record the date the annual return for a financial year was filed (Sec 34(2))."""

    def get_permissions(self):
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    @staticmethod
    def _fy_start(value) -> int:
        try:
            start = int(str(value).split("-")[0])
        except (TypeError, ValueError):
            raise BusinessRuleError("Financial year must look like 2025-26.") from None
        if not 2017 <= start <= 2100:
            raise BusinessRuleError("Financial year must look like 2025-26.")
        return start

    @staticmethod
    def _body(company, start, row):
        return {
            "fy": f"{start}-{str(start + 1)[-2:]}",
            "filed_on": row.filed_on.isoformat() if row else None,
            "note_deadline": note_deadline(company, date(start, 6, 1)).isoformat(),
        }

    def get(self, request):
        from .models import AnnualReturnFiling

        start = self._fy_start(request.query_params.get("fy"))
        row = AnnualReturnFiling.objects.filter(company=self.company, fy_start_year=start).first()
        return Response(self._body(self.company, start, row))

    def put(self, request):
        from core.services.audit import AuditService

        from .models import AnnualReturnFiling

        cu = get_company_user(request)
        if not cu or cu.role not in ("OWNER", "MANAGER", "ACCOUNTANT"):
            raise BusinessRuleError("Only an owner, manager or accountant can record the filing date.")
        start = self._fy_start(request.data.get("fy"))
        raw = request.data.get("filed_on")
        if not raw:
            AnnualReturnFiling.objects.filter(company=self.company, fy_start_year=start).delete()
            return Response(self._body(self.company, start, None))
        try:
            filed_on = date.fromisoformat(str(raw))
        except ValueError:
            raise BusinessRuleError("Filing date must be a valid date (YYYY-MM-DD).") from None
        if filed_on < date(start + 1, 4, 1):
            raise BusinessRuleError("An annual return cannot be filed before the financial year ends.")
        row, _ = AnnualReturnFiling.objects.update_or_create(
            company=self.company, fy_start_year=start, defaults={"filed_on": filed_on}
        )
        AuditService.log(
            action="annual_return_filing_date_set",
            company=self.company,
            user=request.user,
            entity_type="AnnualReturnFiling",
            entity_id=row.pk,
            metadata={"fy_start_year": start, "filed_on": filed_on.isoformat()},
        )
        return Response(self._body(self.company, start, row))


class CancelledDocumentNumbersView(BaseReportView):
    """W0-04 / PD-03: FY cancelled-number register for the CA (CSV or JSON)."""

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def perform_content_negotiation(self, request, force=False):
        from rest_framework.renderers import JSONRenderer

        return (JSONRenderer(), "application/json")

    def get(self, request):
        from datetime import date as date_cls

        from django.db.models import Q

        from core.models import StatutoryDocumentEvent
        from purchases.models import PurchaseInvoice
        from sales.models import SalesCreditNote, SalesDebitNote, SalesInvoice

        fy = (request.query_params.get("fy") or "").strip()
        if not fy:
            raise BusinessRuleError("Query parameter 'fy' is required (e.g. 2026-27).")
        import re

        if not re.fullmatch(r"\d{4}-\d{2}", fy):
            raise BusinessRuleError("Query parameter 'fy' must look like 2026-27.")

        # CR-075: GST FY is always Apr–Mar — bound queries in SQL, not Python.
        start_year = int(fy[:4])
        fy_start = date_cls(start_year, 4, 1)
        fy_end = date_cls(start_year + 1, 3, 31)

        rows = []

        def _gstin(doc) -> str:
            stamp = getattr(doc, "company_gstin", None)
            if stamp is not None:
                return (getattr(stamp, "gstin", None) or "").strip().upper()
            return (getattr(self.company, "gstin", None) or "").strip().upper()

        reason_map: dict[tuple[str, int], str] = {}
        for ev in (
            StatutoryDocumentEvent.objects.filter(
                company=self.company,
                event_type=StatutoryDocumentEvent.EventType.CANCEL,
                entity_type__in=(
                    "sales_invoice",
                    "purchase_invoice",
                    "sales_credit_note",
                    "sales_debit_note",
                ),
            )
            .order_by("entity_type", "entity_id", "-id")
            .only("entity_type", "entity_id", "payload")
        ):
            key = (ev.entity_type, ev.entity_id)
            if key in reason_map:
                continue
            reason_map[key] = (
                str(ev.payload.get("reason") or "") if isinstance(ev.payload, dict) else ""
            )

        def _reason(entity_type, entity_id, fallback=""):
            return reason_map.get((entity_type, entity_id), "") or fallback

        for inv in (
            SalesInvoice.objects.filter(company=self.company, status=SalesInvoice.Status.CANCELLED)
            .exclude(number="")
            .filter(
                Q(invoice_date__gte=fy_start, invoice_date__lte=fy_end)
                | Q(
                    invoice_date__isnull=True,
                    cancelled_at__date__gte=fy_start,
                    cancelled_at__date__lte=fy_end,
                )
            )
            .select_related("company_gstin", "updated_by")
        ):
            rows.append({
                "number": inv.number,
                "doc_type": "SALES_INVOICE",
                "gstin": _gstin(inv),
                "date": (inv.invoice_date.isoformat() if inv.invoice_date else ""),
                "reason": _reason("sales_invoice", inv.pk),
                "user": getattr(inv.updated_by, "email", "") or "",
            })
        for inv in (
            PurchaseInvoice.objects.filter(
                company=self.company,
                status=PurchaseInvoice.Status.CANCELLED,
                invoice_date__gte=fy_start,
                invoice_date__lte=fy_end,
            )
            .exclude(number="")
            .select_related("updated_by")
        ):
            rows.append({
                "number": inv.number,
                "doc_type": "PURCHASE_INVOICE",
                "gstin": (getattr(self.company, "gstin", None) or "").strip().upper(),
                "date": (inv.invoice_date.isoformat() if inv.invoice_date else ""),
                "reason": _reason("purchase_invoice", inv.pk),
                "user": getattr(inv.updated_by, "email", "") or "",
            })
        for note in SalesCreditNote.objects.filter(
            company=self.company,
            status=SalesCreditNote.Status.CANCELLED,
            note_date__gte=fy_start,
            note_date__lte=fy_end,
        ).exclude(number=""):
            rows.append({
                "number": note.number,
                "doc_type": "SALES_CREDIT_NOTE",
                "gstin": (getattr(self.company, "gstin", None) or "").strip().upper(),
                "date": (note.note_date.isoformat() if note.note_date else ""),
                "reason": _reason("sales_credit_note", note.pk),
                "user": "",
            })
        for note in SalesDebitNote.objects.filter(
            company=self.company,
            status=SalesDebitNote.Status.CANCELLED,
            note_date__gte=fy_start,
            note_date__lte=fy_end,
        ).exclude(number=""):
            rows.append({
                "number": note.number,
                "doc_type": "SALES_DEBIT_NOTE",
                "gstin": (getattr(self.company, "gstin", None) or "").strip().upper(),
                "date": (note.note_date.isoformat() if note.note_date else ""),
                "reason": _reason("sales_debit_note", note.pk),
                "user": "",
            })
        rows.sort(key=lambda r: (r["date"], r["number"]))
        export_format = (request.query_params.get("format") or "json").lower()
        if export_format == "csv":
            def stream_cancelled():
                writer = csv.DictWriter(
                    Echo(), fieldnames=["number", "doc_type", "gstin", "date", "reason", "user"]
                )
                yield writer.writeheader()
                for row in rows:
                    yield writer.writerow({k: csv_safe(v) for k, v in row.items()})
            response = StreamingHttpResponse(stream_cancelled(), content_type="text/csv")
            response["Content-Disposition"] = f'attachment; filename="cancelled-numbers-{fy}.csv"'
            return response
        return Response({"fy": fy, "columns": ["number", "doc_type", "gstin", "date", "reason", "user"], "rows": rows})


class GstHealthView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        assert_gstr_enabled(self.company)
        period = request.query_params.get("period")
        return Response(build_gst_health(self.company, period))


class GstRateExposureView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        from reporting.gst_rate_scan import backscan_rate_exposure

        assert_gstr_enabled(self.company)
        return Response(backscan_rate_exposure(
            self.company,
            date_from=_parse_date(request.query_params.get("date_from")),
            date_to=_parse_date(request.query_params.get("date_to")),
        ))


class GstCaPackView(BaseReportView):
    permission_classes = [IsAuthenticated, HasCompany, CanExport]
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def perform_content_negotiation(self, request, force=False):
        from rest_framework.renderers import JSONRenderer

        return (JSONRenderer(), "application/json")

    def get(self, request):
        from accounts.models import Company

        assert_gstr_enabled(self.company)
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        parse_period(period)
        company_gstin = request.query_params.get("company_gstin") or request.query_params.get(
            "companyGstin"
        )
        if self.company.registration_type == Company.RegistrationType.COMPOSITION:
            content = build_ca_pack_zip(self.company, period, gstr1=None)
        else:
            g1 = build_gstr1(self.company, period, company_gstin=company_gstin or None)
            g3 = build_gstr3b(self.company, period, gstr1=g1, company_gstin=company_gstin or None)
            persist_snapshot(self.company, "GSTR-1", period, g1, user=request.user)
            persist_snapshot(self.company, "GSTR-3B", period, g3, user=request.user)
            content = build_ca_pack_zip(self.company, period, gstr1=g1)
        response = HttpResponse(content, content_type="application/zip")
        response["Content-Disposition"] = f'attachment; filename="gst-ca-pack-{period}.zip"'
        return response


class GstPeriodView(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def post(self, request):
        assert_gstr_enabled(self.company)
        period = request.data.get("period")
        action = (request.data.get("action") or "").lower()
        if not period:
            raise BusinessRuleError("'period' is required.")
        cu = get_company_user(request)
        if cu.role != "OWNER":
            raise BusinessRuleError("Only Owner can close or reopen GST periods.")
        if action == "soft_close":
            obj = soft_close_period(self.company, period, request.user)
        elif action == "reopen":
            obj = reopen_period(self.company, period, request.user)
        else:
            raise BusinessRuleError("action must be soft_close or reopen.")
        payload = {
            "period": obj.period,
            "status": obj.status,
            "dirty_after_snapshot": obj.dirty_after_snapshot,
        }
        if action == "soft_close":
            payload["warnings"] = accounting_period_open_warnings(self.company, period)
        return Response(payload)


# B5-007: same DoS-prevention ceiling as imports.services.MAX_IMPORT_ROWS --
# not a business-rule cap, just a bound on what one request may materialise.
MAX_2B_UPLOAD_ROWS = 20_000


class Gstr2bIngestViewSet(viewsets.ModelViewSet):
    """Wave 17A: upload/list/match GSTR-2B ingest rows."""

    permission_classes = [IsAuthenticated, HasCompany, CanViewFinancialReports]
    serializer_class = Gstr2bIngestSerializer
    http_method_names = ["get", "post", "patch", "head", "options"]
    # B5-007: unbounded rows + no throttle was a DoS surface (a 50k-row
    # upload = ~100k queries in one request, unthrottled).
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def get_permissions(self):
        action = getattr(self, "action", None)
        if action in (
            "create", "partial_update", "update", "ims_act", "ims_bulk_accept",
            "ims_offline_import", "match", "upload", "ims_gsp_pull",
            "chase_whatsapp", "chase_photo",
        ):
            return [IsAuthenticated(), HasCompany(), IsOwner()]
        return [IsAuthenticated(), HasCompany(), CanViewFinancialReports()]

    @property
    def company(self):
        return get_company_user(self.request).company

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        assert_gstr_enabled(self.company)

    def get_queryset(self):
        qs = Gstr2bIngest.objects.filter(company=self.company)
        period = self.request.query_params.get("period")
        if period:
            qs = qs.filter(period=period)
        source = (self.request.query_params.get("source") or "").strip().upper()
        if source in ("2A", "2B"):
            qs = qs.filter(raw__source=source)
        ims_action = (self.request.query_params.get("ims_action") or "").strip().upper()
        if ims_action in dict(Gstr2bIngest.ImsAction.choices):
            qs = qs.filter(ims_action=ims_action)
        match_class = (self.request.query_params.get("match_class") or "").strip()
        if match_class:
            qs = qs.filter(match_class=match_class)
        return qs

    def perform_create(self, serializer):
        serializer.save(company=self.company)

    def perform_update(self, serializer):
        # B5-005: a manual itc_eligibility override must not leave the books out
        # of step with the 2B row — run the same reclass the IMS accept/reject
        # path uses when the eligibility actually changes.
        old = serializer.instance.itc_eligibility
        obj = serializer.save()
        new = obj.itc_eligibility
        if new != old and obj.purchase_invoice_id:
            from accounting.services import reclass_rejected_itc, reclass_unreviewed_itc

            inv = obj.purchase_invoice
            user = self.request.user
            if new == Gstr2bIngest.ItcEligibility.CLAIMABLE:
                reclass_unreviewed_itc(inv, user=user)
            elif new in (
                Gstr2bIngest.ItcEligibility.INELIGIBLE,
                Gstr2bIngest.ItcEligibility.REVERSED,
            ):
                reclass_rejected_itc(inv, user=user)

    @action(detail=False, methods=["post"], url_path="upload")
    def upload(self, request):
        """Bulk ingest rows: {period, rows: [{supplier_gstin, invoice_number, ...}]}."""
        period = request.data.get("period")
        rows = request.data.get("rows") or []
        if not period:
            raise BusinessRuleError("'period' is required.")
        parse_period(period)
        # B5-007: cap request size (DoS surface) and make the ingest atomic --
        # an error partway through previously left the period half-ingested
        # with a 200-ish partial "created" count and no rollback.
        if len(rows) > MAX_2B_UPLOAD_ROWS:
            raise BusinessRuleError(
                f"2B upload exceeds {MAX_2B_UPLOAD_ROWS} rows. Split the file and retry."
            )
        source = (request.data.get("source") or "2B").strip().upper()
        if source not in ("2A", "2B"):
            source = "2B"
        # B5-007: validate amount fields up front, before any DB write. A
        # DecimalField ValidationError raised mid-loop from inside
        # Model.save() trips Django's low-level rollback-marking machinery in
        # a way that (empirically, on this Django/SQLite combo) survives the
        # savepoint recovery below and poisons the connection for the rest of
        # the request/test -- on top of surfacing as an opaque 500 instead of
        # a clean 400. Rejecting bad rows here means the atomic block below
        # only ever has to unwind ordinary DB state, never a mid-save error.
        from decimal import Decimal, InvalidOperation

        for idx, row in enumerate(rows):
            if not isinstance(row, dict):
                continue
            for field in ("taxable_value", "igst", "cgst", "sgst", "cess"):
                value = row.get(field)
                if value in (None, ""):
                    continue
                try:
                    Decimal(str(value))
                except InvalidOperation:
                    raise BusinessRuleError(
                        f"Row {idx + 1}: '{field}' is not a valid number."
                    )
        created = []
        with transaction.atomic():
            for row in rows:
                raw = dict(row) if isinstance(row, dict) else {}
                raw["source"] = source
                supplier = (row.get("supplier_gstin") or "").strip().upper()[:15]
                invoice_number = (row.get("invoice_number") or "").strip()[:64]
                defaults = {
                    "invoice_date": row.get("invoice_date") or None,
                    "taxable_value": row.get("taxable_value") or 0,
                    "igst": row.get("igst") or 0,
                    "cgst": row.get("cgst") or 0,
                    "sgst": row.get("sgst") or 0,
                    "cess": row.get("cess") or 0,
                    "raw": raw,
                }
                if invoice_number:
                    # B5-023: a re-uploaded 2B commonly reflects a supplier amendment
                    # (taxable value / tax / date revised). update_or_create below
                    # silently replaces those figures with nothing recorded — snapshot
                    # the prior values into the existing append-only IMS history log
                    # first, whenever they actually differ, so an earlier
                    # reconciliation's numbers aren't lost with no trace.
                    existing = Gstr2bIngest.objects.filter(
                        company=self.company,
                        period=period,
                        supplier_gstin=supplier,
                        invoice_number=invoice_number,
                    ).first()
                    if existing is not None:
                        from decimal import Decimal as _Decimal, InvalidOperation as _InvalidOperation

                        _AMOUNT_FIELDS = ("taxable_value", "igst", "cgst", "sgst", "cess")
                        changed = {}
                        for f in _AMOUNT_FIELDS:
                            old_val = _Decimal(str(getattr(existing, f) or 0))
                            try:
                                new_val = _Decimal(str(defaults.get(f) or 0))
                            except _InvalidOperation:
                                new_val = old_val
                            if old_val != new_val:
                                changed[f] = str(old_val)
                        old_date = str(existing.invoice_date) if existing.invoice_date else ""
                        new_date = str(defaults.get("invoice_date") or "")
                        if old_date and new_date and old_date != new_date:
                            changed["invoice_date"] = old_date
                        if changed:
                            from reporting.models import ImsActionHistory

                            ImsActionHistory.objects.create(
                                company=self.company,
                                ingest=existing,
                                action="REUPLOAD_AMENDED",
                                remark=f"2B re-upload for {period} revised {', '.join(sorted(changed))}",
                                acted_by=request.user if request.user.is_authenticated else None,
                                payload={"previous": changed, "raw_previous": existing.raw},
                            )
                    obj, _created = Gstr2bIngest.objects.update_or_create(
                        company=self.company,
                        period=period,
                        supplier_gstin=supplier,
                        invoice_number=invoice_number,
                        defaults=defaults,
                    )
                else:
                    continue
                created.append(obj.pk)
        return Response({"period": period, "created": len(created), "ids": created})

    @action(detail=False, methods=["post"], url_path="match")
    def match(self, request):
        period = request.data.get("period") or request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        from reporting.gstr2b import match_gstr2b_to_purchases
        from reporting.ims import classify_and_match

        match_gstr2b_to_purchases(self.company, period)
        return Response(classify_and_match(self.company, period))

    @action(detail=True, methods=["post"], url_path="ims-act")
    def ims_act(self, request, pk=None):
        from reporting.ims import apply_ims_action

        row = self.get_object()
        action = request.data.get("action") or request.data.get("ims_action")
        remark = request.data.get("remark") or request.data.get("ims_remark") or ""
        apply_ims_action(row, action, remark=remark, user=request.user)
        row.refresh_from_db()
        return Response(self.get_serializer(row).data)

    @action(detail=False, methods=["post"], url_path="ims-bulk-accept")
    def ims_bulk_accept(self, request):
        from reporting.ims import bulk_accept_exact

        period = request.data.get("period") or request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        return Response(bulk_accept_exact(self.company, period, user=request.user))

    @action(detail=False, methods=["get"], url_path="supplier-nudge")
    def supplier_nudge(self, request):
        from masters.models import Supplier
        from reporting.ims import supplier_po_nudge

        raw = request.query_params.get("supplier")
        if not raw:
            raise BusinessRuleError("'supplier' is required.")
        supplier = Supplier.objects.filter(company=self.company, pk=raw).first()
        if supplier is None:
            raise BusinessRuleError("Supplier not found.")
        return Response(supplier_po_nudge(self.company, supplier))

    @action(detail=False, methods=["get"], url_path="ims-summary")
    def ims_summary(self, request):
        from reporting.ims import credit_at_risk

        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        return Response(credit_at_risk(self.company, period))

    @action(detail=False, methods=["get"], url_path="missing-documents")
    def missing_documents(self, request):
        from reporting.chase import list_missing_documents, serialize_chase_row

        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        rows = list_missing_documents(self.company, period)
        return Response({
            "period": period,
            "count": len(rows),
            "items": [serialize_chase_row(r) for r in rows],
        })

    @action(detail=False, methods=["post"], url_path="chase-whatsapp")
    def chase_whatsapp(self, request):
        from reporting.chase import request_whatsapp

        period = request.data.get("period") or request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        return Response(request_whatsapp(
            self.company,
            period,
            user=request.user,
            phone=request.data.get("phone") or "",
        ))

    @action(detail=True, methods=["post"], url_path="chase-photo")
    def chase_photo(self, request, pk=None):
        from reporting.chase import attach_photo_reply

        uploaded = request.FILES.get("file")
        if not uploaded:
            raise BusinessRuleError("A photo or PDF file is required.")
        return Response(attach_photo_reply(
            self.company, int(pk), uploaded, user=request.user,
        ), status=201)

    @action(detail=False, methods=["get"], url_path="ims-scorecard")
    def ims_scorecard(self, request):
        from reporting.ims import supplier_scorecard

        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        return Response({"period": period, "suppliers": supplier_scorecard(self.company, period)})

    @action(detail=True, methods=["post"], url_path="supplier-message")
    def supplier_message(self, request, pk=None):
        from reporting.ims import supplier_defect_message

        row = self.get_object()
        return Response(supplier_defect_message(row))

    @action(detail=False, methods=["post"], url_path="ims-offline-import")
    def ims_offline_import(self, request):
        from reporting.ims_offline import import_offline

        payload = request.data.get("payload") if "payload" in request.data else request.data
        replace = bool(request.data.get("replace"))
        return Response(import_offline(self.company, payload, replace=replace))

    @action(detail=False, methods=["get"], url_path="ims-offline-export")
    def ims_offline_export(self, request):
        from reporting.ims_offline import export_offline

        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        return Response(export_offline(self.company, period))

    @action(detail=False, methods=["post"], url_path="ims-gsp-pull")
    def ims_gsp_pull(self, request):
        """Fail-closed: live GSP only when P0 credentials exist."""
        from core.services.gsp_adapters import get_gstr_filing_adapter
        from reporting.ims_offline import import_offline

        period = request.data.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        adapter = get_gstr_filing_adapter(self.company)
        fetched = adapter.fetch_gstr2b(period)
        rows = fetched.get("rows") or fetched.get("data") or []
        if not isinstance(rows, list) or not rows:
            raise BusinessRuleError("GSP returned no 2B rows — use the GSTN offline-tool file.")
        return Response(import_offline(self.company, {"period": period, "rows": rows}, replace=False))


class Cmp08View(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        from accounts.models import Company
        from reporting.gstr2b import build_cmp08

        assert_gstr_enabled(self.company)
        if self.company.registration_type != Company.RegistrationType.COMPOSITION:
            raise BusinessRuleError("CMP-08 is only for composition dealers.")
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        parse_period(period)
        return Response(build_cmp08(self.company, period))


class Gstr4View(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        from accounts.models import Company
        from reporting.gstr2b import build_gstr4

        assert_gstr_enabled(self.company)
        if self.company.registration_type != Company.RegistrationType.COMPOSITION:
            raise BusinessRuleError("GSTR-4 is only for composition dealers.")
        fy = request.query_params.get("fy")
        if not fy:
            raise BusinessRuleError("Query parameter 'fy' is required (e.g. 2025-26).")
        return Response(build_gstr4(self.company, fy))


class Gstr8View(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        from reporting.gstr2b import build_gstr8

        assert_gstr_enabled(self.company)
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        parse_period(period)
        return Response(build_gstr8(self.company, period))


class Gstr6View(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        from reporting.gstr2b import build_gstr6

        assert_gstr_enabled(self.company)
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        parse_period(period)
        return Response(build_gstr6(self.company, period))


class Gstr7View(BaseReportView):
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def get(self, request):
        from reporting.gstr2b import build_gstr7

        assert_gstr_enabled(self.company)
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        parse_period(period)
        return Response(build_gstr7(self.company, period))


class GstFilingSandboxView(BaseReportView):
    """Wave 17B: sandbox HTTP upload_gstr1 / fetch_gstr2b when GSP_HTTP_SANDBOX=1.

    BB-000742: production/staging without GSP_CERTIFIED must not expose a live-looking
    GSTR filing upload surface — no certified Live filing adapter ships yet.
    """

    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "gst_reports"

    def post(self, request):
        from django.conf import settings

        from core.services.gsp_adapters import get_gstr_filing_adapter

        assert_gstr_enabled(self.company)
        env = (getattr(settings, "DJANGO_ENV", "") or "").lower()
        certified = getattr(settings, "GSP_CERTIFIED", False) is True
        if env in ("production", "staging") and not certified:
            raise Http404()

        action_name = (request.data.get("action") or "").lower()
        period = request.data.get("period")
        if not period:
            raise BusinessRuleError("'period' is required.")
        parse_period(period)
        adapter = get_gstr_filing_adapter(self.company)
        from reporting.filing import record_period_filing
        from reporting.models import GstPeriodFiling

        if action_name not in ("upload_gstr1", "upload_gstr3b", "fetch_gstr2b"):
            raise BusinessRuleError("action must be upload_gstr1, upload_gstr3b, or fetch_gstr2b.")
        result = None
        payload = None
        return_type = "GSTR-3B" if action_name == "upload_gstr3b" else "GSTR-1"
        try:
            if action_name == "upload_gstr1":
                payload = build_gstr1(self.company, period)
                return_type = (payload or {}).get("return_type", "GSTR-1")
                result = adapter.upload_gstr1(payload)
            elif action_name == "upload_gstr3b":
                payload = build_gstr3b(self.company, period)
                result = adapter.upload_gstr3b(payload)
            else:
                result = adapter.fetch_gstr2b(period)
                return Response({"action": action_name, "result": result})
        except BusinessRuleError as exc:
            if action_name in ("upload_gstr1", "upload_gstr3b"):
                record_period_filing(
                    company=self.company,
                    return_type=return_type,
                    period=period,
                    status=GstPeriodFiling.Status.REJECTED,
                    error_message=str(exc),
                    user=request.user,
                )
            raise
        reference = ""
        if isinstance(result, dict):
            reference = str(result.get("reference_id") or result.get("ack") or "")
        ack = result if isinstance(result, dict) else {"ok": True, "detail": str(result)}
        return_type = (
            "GSTR-3B" if action_name == "upload_gstr3b" else (payload or {}).get("return_type", "GSTR-1")
        )
        # The GSP call already happened above; keep the filing record and its
        # snapshot in one transaction so a filing is never durably recorded as
        # ACCEPTED without the snapshot that should back it (or vice versa).
        with transaction.atomic():
            record_period_filing(
                company=self.company,
                return_type=return_type,
                period=period,
                status=GstPeriodFiling.Status.ACCEPTED,
                reference=reference,
                user=request.user,
            )
            snap = persist_snapshot(
                self.company,
                return_type,
                period,
                payload or {},
                user=request.user,
            )
            body = dict(snap.payload or {})
            body["gsp_upload"] = ack
            snap.payload = body
            snap.save(update_fields=["payload"])
        return Response({
            "action": action_name,
            "result": ack,
            "job_id": ack.get("reference_id") or ack.get("job_id") or ack.get("ack") or "",
        })


class _TaxWorksheetView(APIView):
    """26Q / 27EQ worksheet aids. ``?preview=1`` returns JSON; otherwise a CSV download.

    RPT-03: the owner or an accountant who reads financial reports may preview; the CSV file
    additionally needs export rights. Not a live IT portal upload.
    """

    kind = ""  # "tds" | "tcs"
    banner = ""
    filename = ""
    fieldnames: list[str] = []
    party_key = ""
    amount_key = ""
    rows_fn_name = ""

    def get_permissions(self):
        if self.request.query_params.get("preview") in ("1", "true", "True"):
            return [IsAuthenticated(), HasCompany(), CanViewTaxWorksheets()]
        return [IsAuthenticated(), HasCompany(), CanViewTaxWorksheets(), CanExport()]

    def _rows(self, company, period):
        from . import tds_worksheets

        return getattr(tds_worksheets, self.rows_fn_name)(company, period)

    def get(self, request):
        company = get_company_user(request).company
        if not _tds_worksheets_readable(company):
            raise BusinessRuleError(
                "TDS/TCS is not switched on for this company and no TDS/TCS has been posted yet.",
                code="tds_not_enabled",
            )
        period = request.query_params.get("period")
        if not period:
            raise BusinessRuleError("Query parameter 'period' is required (YYYY-MM).")
        from .gst_returns import parse_period

        try:
            parse_period(period)
        except ValueError as exc:
            raise BusinessRuleError(str(exc)) from exc
        rows = self._rows(company, period)
        if request.query_params.get("preview") in ("1", "true", "True"):
            by_section: dict[str, dict] = {}
            for row in rows:
                bucket = by_section.setdefault(
                    row.get("section") or "", {"section": row.get("section") or "", "count": 0,
                                               "taxable": Decimal("0"), "amount": Decimal("0")},
                )
                bucket["count"] += 1
                bucket["taxable"] += Decimal(str(row.get("taxable") or "0"))
                bucket["amount"] += Decimal(str(row.get(self.amount_key) or "0"))
            return Response({
                "kind": self.kind,
                "period": period,
                "disclaimer": self.banner,
                "count": len(rows),
                "rows": rows,
                "by_section": list(by_section.values()),
                "totals": {
                    "taxable": sum((b["taxable"] for b in by_section.values()), Decimal("0")),
                    "amount": sum((b["amount"] for b in by_section.values()), Decimal("0")),
                },
            })
        fieldnames = self.fieldnames
        banner = self.banner

        def stream():
            writer = csv.DictWriter(Echo(), fieldnames=fieldnames, extrasaction="ignore")
            yield writer.writeheader()
            yield writer.writerow({"date": f"# {banner}", "invoice": ""})
            for row in rows:
                yield writer.writerow({k: csv_safe(v) for k, v in row.items()})

        response = StreamingHttpResponse(stream(), content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{self.filename.format(period=period)}"'
        return response


class TdsWorksheetView(_TaxWorksheetView):
    """BB-000670: 26Q worksheet aid."""

    kind = "tds"
    banner = "26Q worksheet aid — not a live IT portal upload"
    filename = "tds-26q-{period}.csv"
    fieldnames = [
        "date", "invoice", "supplier", "supplier_gstin", "section", "rate",
        "taxable", "grand_total", "tds_amount", "net_payable", "source",
    ]
    amount_key = "tds_amount"
    rows_fn_name = "tds_worksheet_rows"


class TcsWorksheetView(_TaxWorksheetView):
    """BB-000670: 27EQ worksheet aid."""

    kind = "tcs"
    banner = "27EQ worksheet aid — not a live IT portal upload"
    filename = "tcs-27eq-{period}.csv"
    fieldnames = [
        "date", "invoice", "customer", "customer_gstin", "section", "rate",
        "taxable", "grand_total", "tcs_amount", "receivable", "source",
    ]
    amount_key = "tcs_amount"
    rows_fn_name = "tcs_worksheet_rows"


# Backwards-compatible alias — prefer core.csv_utils.csv_safe.
_csv_safe = csv_safe


class ExportView(BaseReportView):
    """CSV and XLSX export of registers via Report Service (E5.8 / FR-019)."""

    permission_classes = [IsAuthenticated, HasCompany, CanExport]
    throttle_classes = [CompanyRateThrottle]
    throttle_scope = "heavy_reports"

    def perform_content_negotiation(self, request, force=False):
        renderers = self.get_renderers()
        return (renderers[0], renderers[0].media_type)

    def get(self, request, report):
        if report not in EXPORTS:
            raise BusinessRuleError(f"Unknown export '{report}'. Available: {', '.join(EXPORTS)}.")
        # CR-074: cap dated register exports (inventory/party lists have no date span).
        if report in (
            "sales-register", "sales", "purchase-register", "purchases",
        ):
            assert_report_date_span(
                _parse_date(request.query_params.get("date_from")),
                _parse_date(request.query_params.get("date_to")),
                kind=f"Export '{report}'",
            )
        rows = EXPORTS[report](self.company, request.query_params)
        hidden = MARGIN_ONLY_COLUMNS.get(report, set())
        if hidden and not CanViewMargins().has_permission(request, self):
            rows = [{k: v for k, v in row.items() if k not in hidden} for row in rows]
        # BUG-323: a filtered-to-zero-rows export used to produce a
        # completely empty file with no header, which some spreadsheet
        # tools mis-render as a corrupt/blank CSV.
        fieldnames = EXPORT_FIELDS.get(report) or (list(rows[0].keys()) if rows else [])
        if hidden:
            fieldnames = [f for f in fieldnames if f not in hidden or CanViewMargins().has_permission(request, self)]
        fmt = (request.query_params.get("format") or "csv").strip().lower()
        if fmt in ("xlsx", "excel"):
            import io
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from django.http import HttpResponse

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = report.replace("-", " ").title()[:31]

            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")

            if fieldnames:
                ws.append(list(fieldnames))
                for col_num in range(1, len(fieldnames) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment

                for row in rows:
                    ws.append([_csv_safe(row.get(k)) for k in fieldnames])

                for col in ws.columns:
                    max_len = max(len(str(cell.value or "")) for cell in col)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            response = HttpResponse(
                buf.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = f'attachment; filename="{report}.xlsx"'
            return response

        def stream_export():
            if fieldnames:
                writer = csv.DictWriter(Echo(), fieldnames=fieldnames, extrasaction="ignore")
                yield writer.writeheader()
                for row in rows:
                    yield writer.writerow({k: _csv_safe(row.get(k)) for k in fieldnames})
        response = StreamingHttpResponse(stream_export(), content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{report}.csv"'
        return response
